# -*- coding: utf-8 -*-
"""Đọc hai file Excel xuất từ Misa.

File 1 — "Tổng hợp công nợ phải thu khách hàng": danh sách khách và số tiền đúng.
File 2 — "Phân tích công nợ phải thu theo tuổi nợ": tuổi của các khoản nợ.

Không tin tên file (Windows tự thêm hậu tố (1), (2)...). Luôn đọc ngày chốt
ghi bên trong file.
"""
from __future__ import annotations

import datetime as dt
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List

import openpyxl

NGAY = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")


class LoiDocFile(Exception):
    """File không đúng khuôn Misa, hoặc thiếu thông tin bắt buộc."""


def chuan_hoa_ten(s) -> str:
    """Khoá so khớp khách hàng. Mã khách không thống nhất giữa hai file nên
    so khớp theo tên đã chuẩn hoá."""
    s = unicodedata.normalize("NFC", str(s or "")).upper().strip()
    s = s.replace("“", '"').replace("”", '"').replace("’", "'")
    s = re.sub(r"[.,]", " ", s)
    return re.sub(r"\s+", " ", s)


def _so(x) -> float:
    return float(x) if isinstance(x, (int, float)) else 0.0


def _cac_ngay(text: str) -> List[dt.date]:
    return [dt.date(int(y), int(m), int(d)) for d, m, y in NGAY.findall(str(text or ""))]


@dataclass
class DongTongHop:
    ma: str
    ten: str
    ten_chuan: str
    du_no: float          # số dư cuối kỳ bên Nợ — khách còn nợ mình
    du_co: float          # số dư cuối kỳ bên Có — khách trả trước
    da_tra_trong_ky: float
    dk_no: float = 0.0    # số dư đầu kỳ bên Nợ
    dk_co: float = 0.0    # số dư đầu kỳ bên Có
    ps_no: float = 0.0    # phát sinh Nợ trong kỳ (TK131: nợ mới; TK331: đã trả vendor)
    ps_co: float = 0.0    # phát sinh Có trong kỳ (TK131: đã thu; TK331: nợ vendor mới)

    @property
    def dau_ky(self) -> float:
        return self.dk_no - self.dk_co

    @property
    def cuoi_ky(self) -> float:
        return self.du_no - self.du_co


@dataclass
class TongHop:
    tu_ngay: dt.date
    den_ngay: dt.date
    dong: Dict[str, DongTongHop] = field(default_factory=dict)

    @property
    def tong_du_no(self) -> float:
        return sum(d.du_no for d in self.dong.values() if d.du_no > 0)

    def khach_con_no(self) -> List[DongTongHop]:
        return [d for d in self.dong.values() if d.du_no > 0]


@dataclass
class DongTuoiNo:
    ma: str
    ten: str
    ten_chuan: str
    tong: float
    khong_han: float       # hoá đơn không ghi hạn thanh toán
    truoc_han: float       # chưa tới hạn
    qua_han: List[float]   # 1-30, 31-60, 61-90, 91-120, trên 120


@dataclass
class TuoiNo:
    den_ngay: dt.date
    dong: Dict[str, DongTuoiNo] = field(default_factory=dict)


def _mo(path) -> openpyxl.worksheet.worksheet.Worksheet:
    wb = openpyxl.load_workbook(Path(path), data_only=True, read_only=False)
    return wb.worksheets[0]


def nhan_dang(path) -> str:
    """Trả về 'tong-hop' | 'tuoi-no' | 'chi-tiet' | 'phai-tra' | 'ban-hang' | 'so-chi-tiet' |
    'khong-ro' dựa vào
    tiêu đề bên trong file. Đọc chế độ read_only vì file Chi tiết từ 2022 rất dài."""
    try:
        wb = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        dau = " ".join(
            str(v or "")
            for hang in ws.iter_rows(min_row=1, max_row=5, max_col=12, values_only=True)
            for v in hang
        ).upper()
        wb.close()
    except Exception:
        return "khong-ro"
    dau = unicodedata.normalize("NFC", dau)
    if "TỔNG HỢP CÔNG NỢ PHẢI THU" in dau:
        return "tong-hop"
    if "TỔNG HỢP CÔNG NỢ PHẢI TRẢ" in dau:
        return "phai-tra"
    if "CHI TIẾT CÔNG NỢ PHẢI THU" in dau:
        return "chi-tiet"
    if "PHÂN TÍCH CÔNG NỢ PHẢI THU" in dau:
        return "tuoi-no"
    # v0.5 (WR-06): hai file FIN thêm 09/10
    if "SỔ CHI TIẾT TÀI KHOẢN" in dau and "131" in dau:
        return "so-chi-tiet"
    if "TT THANH TOÁN" in dau and "SỐ HÓA ĐƠN" in dau:
        return "ban-hang"
    return "khong-ro"


def doc_tong_hop(path) -> TongHop:
    """Tổng hợp công nợ phải thu (TK131). File phải trả (TK331) cùng khuôn cột,
    đọc bằng doc_phai_tra."""
    ws = _mo(path)
    ngay = _cac_ngay(ws["A2"].value)
    if len(ngay) < 2:
        raise LoiDocFile(
            f"{Path(path).name}: không đọc được kỳ báo cáo ở ô A2 "
            f"(cần dạng 'Từ ngày dd/mm/yyyy đến ngày dd/mm/yyyy')"
        )
    th = TongHop(tu_ngay=ngay[0], den_ngay=ngay[1])
    for hang in ws.iter_rows(min_row=6, values_only=True):
        ma, ten = hang[0], hang[1]
        if not ma or not ten:
            continue
        if chuan_hoa_ten(ma).startswith(("TỔNG", "CỘNG")):
            break
        tc = chuan_hoa_ten(ten)
        th.dong[tc] = DongTongHop(
            ma=str(ma).strip(), ten=str(ten).strip(), ten_chuan=tc,
            du_no=_so(hang[8]), du_co=_so(hang[9]), da_tra_trong_ky=_so(hang[7]),
            dk_no=_so(hang[4]), dk_co=_so(hang[5]), ps_no=_so(hang[6]), ps_co=_so(hang[7]),
        )
    if not th.dong:
        raise LoiDocFile(f"{Path(path).name}: không có dòng khách hàng nào")
    return th


def doc_tuoi_no(path) -> TuoiNo:
    ws = _mo(path)
    ngay = _cac_ngay(ws["A4"].value)
    if not ngay:
        raise LoiDocFile(
            f"{Path(path).name}: không đọc được ngày chốt ở ô A4 (cần 'Đến ngày dd/mm/yyyy')"
        )
    tn = TuoiNo(den_ngay=ngay[-1])
    for hang in ws.iter_rows(min_row=9, values_only=True):
        ma, ten = hang[0], hang[1]
        if not ma or not ten:
            continue
        nhan = chuan_hoa_ten(ma)
        if nhan.startswith(("MÃ NHÓM", "CỘNG NHÓM", "TỔNG CỘNG")):
            continue
        if not isinstance(hang[5], (int, float)):
            continue
        tc = chuan_hoa_ten(ten)
        tn.dong[tc] = DongTuoiNo(
            ma=str(ma).strip(), ten=str(ten).strip(), ten_chuan=tc,
            tong=_so(hang[5]), khong_han=_so(hang[7]), truoc_han=_so(hang[13]),
            qua_han=[_so(hang[15]), _so(hang[16]), _so(hang[17]), _so(hang[18]), _so(hang[20])],
        )
    if not tn.dong:
        raise LoiDocFile(f"{Path(path).name}: không có dòng khách hàng nào")
    return tn


def kiem_cung_ky(th: TongHop, tn: TuoiNo) -> None:
    """Hai file phải cùng kỳ: ngày chốt của file Tuổi nợ = ngày cuối kỳ của
    file Tổng hợp + 1 ngày. Lệch thì dừng, không được ra báo cáo sai."""
    mong_doi = th.den_ngay + dt.timedelta(days=1)
    if tn.den_ngay != mong_doi:
        raise LoiDocFile(
            "Hai file lệch kỳ, dừng lại để khỏi ra số sai.\n"
            f"  Tổng hợp: {th.tu_ngay:%d/%m/%Y} đến {th.den_ngay:%d/%m/%Y}\n"
            f"  Tuổi nợ : đến ngày {tn.den_ngay:%d/%m/%Y} (đáng ra phải là {mong_doi:%d/%m/%Y})"
        )


# --- v0.4: phải trả nhà cung cấp (TK331), cùng khuôn cột với file phải thu ----

@dataclass
class PhaiTra:
    tu_ngay: dt.date
    den_ngay: dt.date
    dong: List[DongTongHop] = field(default_factory=list)   # giữ thứ tự Misa, khoá theo mã

    @property
    def tong_con_no(self) -> float:
        return sum(d.du_co for d in self.dong if d.du_co > 0)

    @property
    def tong_da_tra(self) -> float:
        return sum(d.ps_no for d in self.dong)

    @property
    def tong_no_moi(self) -> float:
        return sum(d.ps_co for d in self.dong)


def doc_phai_tra(path) -> PhaiTra:
    ws = _mo(path)
    ngay = _cac_ngay(ws["A2"].value)
    if len(ngay) < 2:
        raise LoiDocFile(f"{Path(path).name}: không đọc được kỳ báo cáo ở ô A2")
    pt = PhaiTra(tu_ngay=ngay[0], den_ngay=ngay[1])
    for hang in ws.iter_rows(min_row=6, values_only=True):
        ma, ten = hang[0], hang[1]
        if not ma:
            continue
        if chuan_hoa_ten(ma).startswith(("TỔNG", "CỘNG")):
            break
        pt.dong.append(DongTongHop(
            ma=str(ma).strip(), ten=str(ten or ma).strip(), ten_chuan=chuan_hoa_ten(ten or ma),
            du_no=_so(hang[8]), du_co=_so(hang[9]), da_tra_trong_ky=_so(hang[6]),
            dk_no=_so(hang[4]), dk_co=_so(hang[5]), ps_no=_so(hang[6]), ps_co=_so(hang[7]),
        ))
    return pt

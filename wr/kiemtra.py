# -*- coding: utf-8 -*-
"""Nút 2_Kiem_tra_ghep_khach: đối chiếu file ghép khách FIN điền với file Misa
và file SMS mới nhất trong thư mục input, rồi sinh một file KIEM_TRA cho FIN xem.

Không sửa file của FIN. Khách mới chưa ghép được liệt kê đúng khuôn cột của file
ghép khách để FIN chép dòng sang.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Callable, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import banhang, chitiet, ghep, misa, tindung
from . import soghi as sg
from . import thuonghieu as th
from . import tinhtoan

VANG = PatternFill("solid", fgColor="FFFF00")
TIEN = "#,##0;[Red]-#,##0"
COT_GHEP = ["Mã SMS", "Tên khách SMS", "Còn nợ theo SMS (VND)", "Số job còn nợ", "Gợi ý: mã Misa",
            "Gợi ý: tên Misa", "Số dư Misa (VND)", "Mức tin cậy", "Xác nhận", "Mã Misa đúng (nếu Sai)",
            "Tên Misa (tự hiện)", "Ghi chú"]
COT_MISA = ["Mã khách Misa", "Tên khách Misa", "Số dư Nợ (VND)", "Mã SMS tương ứng", "Hoặc chọn", "Ghi chú"]


def _moi_nhat(ds: List[Path]) -> Optional[Path]:
    return max(ds, key=lambda f: f.stat().st_mtime) if ds else None


def _tim_ban_hang(thu_muc: Path):
    """v0.5: file Bán hàng và Sổ chi tiết TK131 mới nhất (kể cả file đã [DONE])."""
    ds = {"ban-hang": [], "so-chi-tiet": []}
    for f in Path(thu_muc).glob("*.xlsx"):
        if not f.name.startswith("~$"):
            loai = misa.nhan_dang(f)
            if loai in ds:
                ds[loai].append(f)
    return _moi_nhat(ds["ban-hang"]), _moi_nhat(ds["so-chi-tiet"])


def _tim_dau_vao(thu_muc: Path):
    """File Misa cùng kỳ mới nhất (Tổng hợp + Chi tiết, hoặc Tuổi nợ cũ) và file SMS mới nhất, kể cả file đã [DONE]."""
    tat_ca = [f for f in Path(thu_muc).glob("*.xlsx") if not f.name.startswith("~$")]
    ds_th, ds_tn, ds_sms, ds_ct = [], [], [], []
    for f in tat_ca:
        if ghep.la_file_sms(f):
            ds_sms.append(f)
            continue
        loai = misa.nhan_dang(f)
        if loai == "tong-hop":
            ds_th.append(f)
        elif loai == "tuoi-no":
            ds_tn.append(f)
        elif loai == "chi-tiet":
            ds_ct.append(f)
    th_doc = tn_doc = None
    for f in sorted(ds_th, key=lambda f: f.stat().st_mtime, reverse=True):
        try:
            t = misa.doc_tong_hop(f)
        except Exception:
            continue
        for g in sorted(ds_tn, key=lambda f: f.stat().st_mtime, reverse=True):
            try:
                n = misa.doc_tuoi_no(g)
            except Exception:
                continue
            if n.den_ngay == t.den_ngay + dt.timedelta(days=1):
                th_doc, tn_doc = (f, t), (g, n)
                break
        if th_doc:
            break
    if th_doc is None and ds_th:
        f = _moi_nhat(ds_th)
        th_doc = (f, misa.doc_tong_hop(f))
    # v0.4: file Chi tiết 131 cùng ngày cuối kỳ với file Tổng hợp
    ct = None
    if th_doc:
        for f in sorted(ds_ct, key=lambda f: f.stat().st_mtime, reverse=True):
            try:
                if chitiet.ky_cua(f)[1] == th_doc[1].den_ngay:
                    ct = f
                    break
            except Exception:
                continue
    return th_doc, tn_doc, _moi_nhat(ds_sms), ct


def _tieu_de(ws, cot, rong):
    for i, (t, w) in enumerate(zip(cot, rong), start=1):
        o = ws.cell(row=1, column=i, value=t)
        o.font = Font(name=th.PHONG_TIEU_DE, bold=True, color=th.TRANG)
        o.fill = PatternFill("solid", fgColor=th.TIM)
        o.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[o.column_letter].width = w
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"


def _ghi(ws, hang, gia_tri, cot_tien=(), cot_vang=()):
    for i, v in enumerate(gia_tri, start=1):
        o = ws.cell(row=hang, column=i, value=v)
        o.font = Font(name=th.PHONG_NOI_DUNG, size=10)
        if i in cot_tien:
            o.number_format = TIEN
        if i in cot_vang:
            o.fill = VANG


def kiem_tra(cfg, so: sg.SoGhi, ghi_log: Callable[[str], None]) -> Optional[Path]:
    f_ghep = ghep.tim_file_ghep(cfg.ghep)
    if f_ghep is None:
        ghi_log(f"Chưa có file ghép khách trong {cfg.ghep}. Bỏ file FIN đã điền vào đó rồi bấm lại.")
        return None
    bang = ghep.doc_bang_ghep(f_ghep)
    th_doc, tn_doc, f_sms, f_ct = _tim_dau_vao(cfg.vao)
    tong_hop = th_doc[1] if th_doc else None
    dong_sms = ghep.doc_sms(f_sms) if f_sms else []

    wb = openpyxl.Workbook()
    tt = wb.active
    tt.title = "Tóm tắt"
    tt.column_dimensions["A"].width = 46
    tt.column_dimensions["B"].width = 90
    dong_tt = [
        ("KIỂM TRA FILE GHÉP KHÁCH", f"{dt.datetime.now():%d/%m/%Y %H:%M}"),
        ("File ghép khách đang dùng", f_ghep.name),
        ("Misa Tổng hợp dùng để đối chiếu", th_doc[0].name if th_doc else "(chưa có trong input)"),
        ("Misa Chi tiết 131", f_ct.name if f_ct else "(chưa có cùng kỳ)"),
        ("File SMS", f_sms.name if f_sms else "(chưa có trong input)"),
        ("Báo cáo tuần trước tool đã đọc", ", ".join(sorted(so.da_nap)) or "(chưa có — lần đầu bỏ bản "
         "final gần nhất của kế toán trưởng vào thư mục gốc)"),
        ("", ""),
        ("Mã SMS đã ghép với khách Misa", len(bang.sms_sang_misa)),
        ("Mã SMS FIN khai không có trên Misa", len(bang.khong_misa)),
        ("Khách khó đòi (luôn 120+)", ", ".join(sorted(bang.kho_doi)) or "0"),
    ]
    for i, (a, b) in enumerate(dong_tt, start=1):
        tt.cell(row=i, column=1, value=a).font = Font(name=th.PHONG_NOI_DUNG, bold=True)
        tt.cell(row=i, column=2, value=b).font = Font(name=th.PHONG_NOI_DUNG)
    tt["A1"].font = Font(name=th.PHONG_TIEU_DE, size=14, bold=True, color=th.TIM)

    # --- Cần sửa
    cs = wb.create_sheet("Cần sửa")
    _tieu_de(cs, ["Sheet", "Dòng", "Mã", "Vấn đề", "Tool đang xử lý thế nào"], [16, 8, 22, 60, 60])
    h = 2
    for sheet, dong, ma, van_de in bang.loi:
        _ghi(cs, h, [sheet, dong, ma, van_de, "Chưa ghép mã SMS này — số SMS không vào báo cáo"])
        h += 1
    for ma_sms, cu, moi in bang.xung_dot:
        _ghi(cs, h, ["Cả hai sheet", "", ma_sms, f"Mã SMS ghép với cả {cu} và {moi}",
                     f"Theo phía Misa: {moi} (FIN 06/10). Sửa một bên cho khớp để hết cảnh báo"])
        h += 1
    if tong_hop:
        co_ma = {k.ma.upper() for k in tong_hop.dong.values()}
        con_no = {k.ma.upper() for k in tong_hop.khach_con_no()}
        for ma_sms, dich in sorted(bang.sms_sang_misa.items()):
            if dich.upper() not in co_ma:
                _ghi(cs, h, ["Ghép khách", "", ma_sms, f"Mã Misa '{dich}' không có trong file Tổng hợp",
                             "Coi như khách đã trả hết, SMS chưa gạch paid. Gõ sai mã thì sửa lại"])
                h += 1
            elif dich.upper() not in con_no:
                pass   # có trên Misa nhưng không còn nợ: bình thường
    if h == 2:
        _ghi(cs, 2, ["", "", "", "Không có dòng nào cần sửa", ""])

    # --- Khách mới
    sm = wb.create_sheet("SMS mới chưa ghép")
    _tieu_de(sm, COT_GHEP, [16, 40, 18, 10, 16, 40, 18, 22, 20, 20, 36, 30])
    mm = wb.create_sheet("Misa mới chưa ghép")
    _tieu_de(mm, COT_MISA, [18, 46, 18, 20, 24, 30])
    dv = DataValidation(type="list", formula1='"Đúng,Sai,Không có trong Misa"', allow_blank=True)
    sm.add_data_validation(dv)
    so_sms_moi = so_misa_moi = 0
    if tong_hop:
        theo_ma = {k.ma.upper(): k.ma for k in tong_hop.dong.values()}
        theo_ten = {k.ten_chuan: k.ma for k in tong_hop.dong.values()}
        ds_misa = [(k.ma, k.ten, k.du_no) for k in tong_hop.khach_con_no()]
        moi = {}
        da_co_sms = set()
        for x in dong_sms:
            ma, cach = ghep.ma_misa_cua(x, bang, theo_ma, theo_ten)
            if ma:
                da_co_sms.add(theo_ma.get(ma.upper(), ma))
            if cach == "chưa ghép":
                m = moi.setdefault(x.ma_sms, [x.ten_sms, 0.0, set()])
                m[1] += x.so_tien
                m[2].add(x.job)
        h = 2
        for ma_sms, (ten, tien, jobs) in sorted(moi.items(), key=lambda kv: -kv[1][1]):
            g = ghep.goi_y(ten, tien, ds_misa)
            _ghi(sm, h, [ma_sms, ten, tien, len(jobs), g[0], g[1], g[2], g[3]], (3, 7), (9, 10, 12))
            for c in (9, 10, 12):
                sm.cell(row=h, column=c).fill = VANG
            dv.add(f"I{h}")
            h += 1
        so_sms_moi = h - 2
        da_khai = set(bang.sms_sang_misa.values()) | bang.da_khai_misa | set(bang.kho_doi)
        h = 2
        for k in sorted(tong_hop.khach_con_no(), key=lambda k: -k.du_no):
            if k.ma in da_khai or k.ma in da_co_sms:
                continue
            _ghi(mm, h, [k.ma, k.ten, k.du_no], (3,))
            for c in (4, 5, 6):
                mm.cell(row=h, column=c).fill = VANG
            h += 1
        so_misa_moi = h - 2
    if so_sms_moi == 0:
        _ghi(sm, 2, ["", "Không có khách SMS mới"])
    if so_misa_moi == 0:
        _ghi(mm, 2, ["", "Không có khách Misa mới"])

    # --- Đối chiếu từng khách: chạy thử luật tính, không ghi gì vào sổ
    dc = wb.create_sheet("Đối chiếu từng khách")
    _tieu_de(dc, ["Mã Misa", "Khách hàng", "Misa (VND)", "SMS ghép được", "Cách tính tuổi"]
             + sg.TEN_NHOM + ["Cần xem lại"], [14, 40, 16, 16, 14] + [14] * 6 + [70])
    job_misa = None
    if tong_hop and (f_ct or tn_doc):
        if f_ct:
            f_bh, f_sct = _tim_ban_hang(cfg.vao)
            job_misa = (banhang.doc_job_misa(f_bh, f_sct, [(k.ma, k.ten) for k in tong_hop.dong.values()])
                        if f_bh else None)
            kq = tindung.tinh_full(tong_hop, chitiet.doc_chi_tiet(f_ct), so, dong_sms, bang,
                                   job_no_cu=ghep.doc_job_no_cu(cfg.ghep), job_misa=job_misa)
        else:
            kq = tinhtoan.tinh_bang_sms(tong_hop, tn_doc[1], so, dong_sms, bang, tn_doc[1].den_ngay)
        theo_ma = {k.ma.upper(): k.ma for k in tong_hop.dong.values()}
        theo_ten = {k.ten_chuan: k.ma for k in tong_hop.dong.values()}
        sms_khach = {}
        for x in dong_sms:
            ma, _ = ghep.ma_misa_cua(x, bang, theo_ma, theo_ten)
            if ma:
                ma = theo_ma.get(ma.upper(), ma)
                sms_khach[ma] = sms_khach.get(ma, 0) + x.so_tien
        CACH = {"hoa-don": "Hoá đơn Misa", "sms": "SMS", "so-ghi": "Bản final", "sms+so-ghi": "SMS + bản final",
                "kho-doi": "Khó đòi", "uoc-tinh": "ƯỚC TÍNH", "ngoai-le": "Ngoại lệ"}
        for h, d in enumerate(sorted(kq.dong, key=lambda d: -d.total), start=2):
            _ghi(dc, h, [d.ma, d.ten, d.total, sms_khach.get(d.ma, 0), CACH.get(d.nguon, d.nguon)]
                 + d.nhom + [" | ".join(d.cho_xem_lai)], (3, 4, 6, 7, 8, 9, 10, 11))
            if d.nguon == "uoc-tinh":
                for c in range(1, 13):
                    dc.cell(row=h, column=c).fill = PatternFill("solid", fgColor=th.DO_CANH_BAO)
        h = len(kq.dong) + 3
        _ghi(dc, h, ["", "TỔNG", sum(d.total for d in kq.dong), sum(sms_khach.values()), ""]
             + [sum(d.nhom[i] for d in kq.dong) for i in range(6)], (3, 4, 6, 7, 8, 9, 10, 11))
        tt.cell(row=12, column=1, value="Đối chiếu chạy thử").font = Font(name=th.PHONG_NOI_DUNG, bold=True)
        uoc = sum(d.phan_nguon.get("uoc-tinh", 0) for d in kq.dong)
        tt.cell(row=12, column=2, value=(
            f"{len(kq.dong)} khách, tổng {sum(d.total for d in kq.dong):,.0f}. "
            f"Chưa rõ tuổi (đang ước tính): {uoc:,.0f}. SMS cần gạch paid: "
            f"{sum(x[2] for x in kq.sms_cat) + sum(x[3] for x in kq.sms_khong_misa):,.0f}"))
    else:
        _ghi(dc, 2, ["", "Chưa có file Tổng hợp + Chi tiết 131 cùng kỳ trong input nên chưa chạy thử được"])

    # --- Nợ cũ chưa có job trên SMS: FIN tìm job/Sales cho từng hoá đơn, khai một lần
    nc = wb.create_sheet("Nợ chưa có job")
    _tieu_de(nc, ghep.COT_JOB_NO_CU, [16, 40, 16, 14, 18, 22, 24, 30])
    da_khai = ghep.doc_job_no_cu(cfg.ghep)
    dong_nc = []
    if tong_hop and f_ct and job_misa is not None:
        # v0.5: job năm nay đã có trên Misa — chỉ còn hoá đơn TRƯỚC năm nay cần FIN tìm job
        for d in kq.dong:
            if d.nguon == "kho-doi":
                continue
            co = {so_hd for so_hd, _, _ in da_khai.get(d.ma.upper(), [])}
            for ngay, so_hd, han, tien in sorted(d.hoa_don, key=lambda x: (x[0] or dt.date.min)):
                if ngay is None or ngay >= job_misa.tu_ngay:
                    continue
                n = int(so_hd) if str(so_hd).isdigit() else None
                if n not in co:
                    dong_nc.append([d.ma, d.ten, so_hd, ngay.strftime("%d/%m/%Y"), tien, "", "", ""])
    elif tong_hop and (f_ct or tn_doc) and f_ct:
        for d in kq.dong:
            thieu = d.can_doi if d.dong_con else d.total     # khách không có job SMS nào: thiếu cả
            if thieu < 1 or not d.hoa_don:
                continue
            co = {so_hd for so_hd, _, _ in da_khai.get(d.ma.upper(), [])}
            con = thieu
            for ngay, so_hd, han, tien in sorted(d.hoa_don, key=lambda x: (x[0] or dt.date.min)):
                if con < 1 or ngay is None:
                    continue
                n = int(so_hd) if str(so_hd).isdigit() else None
                if n in co:
                    continue
                dong_nc.append([d.ma, d.ten, so_hd, ngay.strftime("%d/%m/%Y"), min(tien, con), "", "", ""])
                con -= tien
    for h, r in enumerate(dong_nc, start=2):
        _ghi(nc, h, r, (5,), (6, 7, 8))
        for c in (6, 7, 8):
            nc.cell(row=h, column=c).fill = VANG
    if not dong_nc:
        _ghi(nc, 2, ["", "Không có nợ cũ nào thiếu job"])
    tep_tc = Path(cfg.ghep) / "Tham_chieu_job_no_cu.xlsx"
    if dong_nc and not tep_tc.exists():
        tc = openpyxl.Workbook()
        w = tc.active
        w.title = ghep.SHEET_JOB_NO_CU
        _tieu_de(w, ghep.COT_JOB_NO_CU, [16, 40, 16, 14, 18, 22, 24, 30])
        for h, r in enumerate(dong_nc, start=2):
            _ghi(w, h, r, (5,), (6, 7, 8))
            for c in (6, 7, 8):
                w.cell(row=h, column=c).fill = VANG
        tc.save(tep_tc)
        ghi_log(f"Đã tạo {tep_tc.name} trong sample: FIN điền cột Job và Sales (ô vàng) một lần")
    tt.cell(row=15, column=1, value="Nợ cũ chưa có job").font = Font(name=th.PHONG_NOI_DUNG, bold=True)
    tt.cell(row=15, column=2, value=f"{len(dong_nc)} hoá đơn — điền Job, Sales ở file Tham_chieu_job_no_cu.xlsx "
            "trong sample (dòng mới: chép từ sheet 'Nợ chưa có job')")

    tt.cell(row=13, column=1, value="Khách mới cần FIN ghép").font = Font(name=th.PHONG_NOI_DUNG, bold=True)
    tt.cell(row=13, column=2, value=f"{so_sms_moi} mã SMS, {so_misa_moi} khách Misa — xem 2 sheet "
            "'... mới chưa ghép', chép dòng sang file ghép khách rồi điền ô vàng")
    tt.cell(row=14, column=1, value="Dòng cần sửa").font = Font(name=th.PHONG_NOI_DUNG, bold=True)
    tt.cell(row=14, column=2, value=f"{len(bang.loi) + len(bang.xung_dot)} — xem sheet 'Cần sửa'")

    dich = Path(cfg.ghep) / f"KIEM_TRA_ghep_khach_{dt.datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(dich)
    ghi_log(f"Kiểm tra ghép khách: {len(bang.sms_sang_misa)} mã đã ghép, {so_sms_moi} SMS mới, "
            f"{so_misa_moi} Misa mới, {len(bang.loi) + len(bang.xung_dot)} dòng cần sửa -> {dich.name}")
    return dich

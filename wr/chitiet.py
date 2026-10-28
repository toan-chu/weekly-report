# -*- coding: utf-8 -*-
"""v0.4 — File "Chi tiết công nợ phải thu khách hàng" (TK131) của Misa, xuất từ
01/01/2022 đến hết kỳ báo cáo.

File này có từng hoá đơn (ngày hoá đơn, số hoá đơn) và từng phiếu thu (TK đối ứng),
nên tính được tuổi nợ thật cho mọi khách, không cần đoán từ tuần trước.

Khuôn file (dòng 4-5 là tiêu đề, dữ liệu từ dòng 6):
    Tên khách hàng: <tên>                   ← dòng mở đầu một khách
    Ngày hạch toán | Ngày chứng từ | Số chứng từ | Ngày hoá đơn | Số hoá đơn |
    Diễn giải | TK công nợ | TK đối ứng | PS Nợ | PS Có | Dư Nợ | Dư Có
    ... | Cộng | ...                          ← dòng cộng của khách, bỏ qua
    Tổng cộng                                ← hết file
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl

from .misa import LoiDocFile, _cac_ngay, chuan_hoa_ten

# Phiếu thu tách theo TK đối ứng: tiền thật về tài khoản hay chỉ là bút toán
NHOM_THU = [
    ("ngan-hang", "Bank", ("112",)),
    ("tien-mat", "Cash", ("111",)),
    ("can-tru", "Offset with payables", ("331",)),
    ("ty-gia-phi", "FX & bank fees", ("515", "635", "642", "811")),
]
TEN_NHOM_THU = {k: ten for k, ten, _ in NHOM_THU} | {"khac": "Other"}


_HD = re.compile(r"(?:hóa đơn|hoá đơn|hđ|hd|invoice|inv)\.?\s*(?:số)?\s*[:#]?\s*"
                 r"((?:\d{3,8}(?:\s*(?:,|;|&|-|và|and|/)\s*)?)+)", re.IGNORECASE)


def so_hd_nhac_toi(dien_giai: str) -> List[int]:
    """Số hoá đơn được nhắc trong diễn giải phiếu thu: 'Thu tiền ... theo HD 00001574,
    00001419' → [1574, 1419]. Mã job dạng EXSANA24100118 không bị bắt nhầm."""
    ra: List[int] = []
    for m in _HD.finditer(dien_giai or ""):
        for so in re.findall(r"\d{3,8}", m.group(1)):
            if int(so) not in ra:
                ra.append(int(so))
    return ra


def _so_hd(so: str) -> Optional[int]:
    return int(so) if so and so.strip().isdigit() else None


def nhom_thu(tk_doi_ung: str) -> str:
    tk = str(tk_doi_ung or "").strip()
    for k, _, dau in NHOM_THU:
        if tk.startswith(dau):
            return k
    return "khac"


@dataclass
class ButToan:
    ngay_ht: dt.date            # ngày hạch toán
    ngay_hd: Optional[dt.date]  # ngày hoá đơn (dòng thu tiền không có)
    so_hd: str
    so_ct: str
    dien_giai: str
    tk_doi_ung: str
    no: float
    co: float


@dataclass
class KhachChiTiet:
    ten: str
    ten_chuan: str
    but_toan: List[ButToan] = field(default_factory=list)

    def so_du(self, den: dt.date) -> float:
        return round(sum(b.no - b.co for b in self.but_toan if b.ngay_ht <= den), 2)

    def thu_trong_ky(self, tu: dt.date, den: dt.date) -> Dict[str, float]:
        """Tiền khách trả trong kỳ (phát sinh Có), tách theo TK đối ứng."""
        ra: Dict[str, float] = {}
        for b in self.but_toan:
            if tu <= b.ngay_ht <= den and b.co:
                k = nhom_thu(b.tk_doi_ung)
                ra[k] = round(ra.get(k, 0) + b.co, 2)
        return ra

    def no_moi_trong_ky(self, tu: dt.date, den: dt.date) -> float:
        return round(sum(b.no for b in self.but_toan if tu <= b.ngay_ht <= den), 2)

    def hoa_don_con_no(self, den: dt.date, da_tt=None) -> Tuple[List[Tuple[dt.date, str, float]], float]:
        """Tiền thu trừ vào hoá đơn thế nào:
          1. Phiếu thu có ghi số hoá đơn trong diễn giải ("theo HD 00001574, 00001419")
             → trừ thẳng vào đúng các hoá đơn đó (cũ trước), như kế toán đã gạch nợ.
          2. Phần còn lại → trước hết trừ vào các hoá đơn Misa (Bán hàng) ghi "Đã thanh toán" tại ngày
             cuối kỳ — da_tt = {(số HĐ, năm)} — vì kế toán đã đối trừ chúng trên Misa (FIN 17/10, STA
             CENTER: HĐ 1560 Misa ghi đã TT); sau đó mới tới hoá đơn CŨ NHẤT trước.
        Trả về ([(ngày hoá đơn, số hoá đơn, còn nợ)], tiền thu chưa trừ hết)."""
        da_tt = da_tt or set()
        hd: Dict[Tuple[dt.date, str], float] = {}
        thu: List[Tuple[float, List[int], dt.date]] = []
        ty_gia: set = set()     # khoá của bút toán chênh lệch tỷ giá (NVK) làm tăng nợ
        for b in self.but_toan:
            if b.ngay_ht > den:
                continue
            ngay = b.ngay_hd or b.ngay_ht
            khoa = (ngay, b.so_hd or b.so_ct)
            # dòng âm là bút toán đảo: ghi Nợ âm = giảm nợ, ghi Có âm = tăng nợ
            tang = max(b.no, 0) + max(-b.co, 0)
            giam = max(b.co, 0) + max(-b.no, 0)
            if tang:
                hd[khoa] = hd.get(khoa, 0) + tang
                if _so_hd(khoa[1]) is None and "tỷ giá" in (b.dien_giai or "").lower():
                    ty_gia.add(khoa)
            if giam:
                thu.append((giam, so_hd_nhac_toi(b.dien_giai), b.ngay_ht))
        con = {k: v for k, v in hd.items()}
        theo_so: Dict[int, List[Tuple[dt.date, str]]] = {}
        for k in sorted(con):
            n = _so_hd(k[1])
            if n is not None:
                theo_so.setdefault(n, []).append(k)
        # Chênh lệch tỷ giá khi đối trừ (NVK) cùng ngày với phiếu thu ghi số hoá đơn → thuộc hoá đơn đó
        # (W42 OPTIMALOG: NVK.2026.04.0208 ngày 04/06 đi với phiếu thu "theo HD 579")
        for k in sorted(ty_gia):
            for _tien, nhac, ngay_thu in thu:
                if ngay_thu != k[0] or not nhac:
                    continue
                dich = next((max((x for x in theo_so.get(n, []) if x[0] <= ngay_thu), default=None)
                             for n in nhac if any(x[0] <= ngay_thu for x in theo_so.get(n, []))), None)
                if dich:
                    con[dich] = round(con[dich] + con.pop(k), 2)
                    break
        du = 0.0
        for tien, nhac, ngay_thu in thu:
            for n in nhac:
                # số hoá đơn Misa đánh lại mỗi năm: chọn hoá đơn mang số đó GẦN NHẤT trước ngày thu
                ung_vien = sorted((k for k in theo_so.get(n, []) if k[0] <= ngay_thu),
                                  key=lambda k: k[0], reverse=True)
                for k in ung_vien[:1]:
                    if tien <= 0.5:
                        break
                    tru = min(tien, con[k])
                    con[k] = round(con[k] - tru, 2)
                    tien = round(tien - tru, 2)
            du += tien
        du = round(du, 2)
        uu_tien = lambda k: 0 if (_so_hd(k[1]), k[0].year) in da_tt else 1
        for k in sorted(con, key=lambda k: (uu_tien(k), k)):
            tru = min(du, con[k])
            du = round(du - tru, 2)
            con[k] = round(con[k] - tru, 2)
        ra = [(k[0], k[1], con[k]) for k in sorted(con) if con[k] > 0.5]
        return ra, du

    def hoa_don_tra_tre(self, den: dt.date, han_cua, da_tt=None, so_ngay: int = 15,
                        so_thang: int = 6) -> List[Tuple[str, dt.date, dt.date, int]]:
        """WR-11 (CEO): hoá đơn đã TẤT TOÁN trong `so_thang` tháng tới `den` mà ngày tất toán trễ hơn hạn
        quá `so_ngay` ngày. han_cua(ngày HĐ) → ngày đến hạn. Ngày tất toán = ngày có phiếu thu đầu tiên mà
        sau đó hoá đơn hết nợ, theo đúng cách tool phân bổ tiền thu (hoa_don_con_no). Hoá đơn còn nợ không
        tính ở đây (đã nằm trong tuổi nợ). Trả [(số HĐ, ngày HĐ, ngày tất toán, số ngày trễ)]."""
        tu = den - dt.timedelta(days=round(so_thang * 30.5))

        def mo(ngay):          # khoá hoá đơn (số HĐ thật) còn nợ tại `ngay`
            con, _ = self.hoa_don_con_no(ngay, da_tt)
            return {(n, s) for n, s, _t in con if _so_hd(s) is not None}

        def co(ngay):          # khoá hoá đơn đã phát sinh tới `ngay`
            return {((b.ngay_hd or b.ngay_ht), b.so_hd) for b in self.but_toan
                    if b.ngay_ht <= ngay and b.no > 0 and _so_hd(b.so_hd) is not None}

        ngay_thu = sorted({b.ngay_ht for b in self.but_toan
                           if tu < b.ngay_ht <= den and (b.co > 0 or b.no < 0)})
        ra: List[Tuple[str, dt.date, dt.date, int]] = []
        truoc_mo, truoc_co = mo(tu), co(tu)
        for d in ngay_thu:
            nay_mo, nay_co = mo(d), co(d)
            for k in sorted((truoc_mo | (nay_co - truoc_co)) - nay_mo):
                tre = (d - han_cua(k[0])).days
                if tre > so_ngay:
                    ra.append((k[1], k[0], d, tre))
            truoc_mo, truoc_co = nay_mo, nay_co
        return ra


@dataclass
class ChiTiet:
    tu_ngay: dt.date
    den_ngay: dt.date
    khach: Dict[str, KhachChiTiet] = field(default_factory=dict)

    def cat_den(self, den: dt.date) -> int:
        """FIN W41 xuất Chi tiết đến cuối tháng (31/10) còn Tổng hợp đến thứ Sáu (09/10): bỏ mọi bút
        toán sau ngày cuối kỳ để file dùng được như xuất đúng kỳ. Trả số bút toán đã bỏ."""
        bo = 0
        for k in self.khach.values():
            giu = [b for b in k.but_toan if b.ngay_ht <= den]
            bo += len(k.but_toan) - len(giu)
            k.but_toan = giu
        self.den_ngay = min(self.den_ngay, den)
        return bo


def _ngay(x) -> Optional[dt.date]:
    if isinstance(x, dt.datetime):
        return x.date()
    if isinstance(x, dt.date):
        return x
    return None


def _so(x) -> float:
    return float(x) if isinstance(x, (int, float)) else 0.0


def doc_chi_tiet(path) -> ChiTiet:
    wb = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    dau = [hang for hang in ws.iter_rows(min_row=1, max_row=3, values_only=True)]
    ngay = []
    for hang in dau:
        for v in hang:
            ngay += _cac_ngay(v)
    if len(ngay) < 2:
        wb.close()
        raise LoiDocFile(f"{Path(path).name}: không đọc được kỳ ở dòng 2 "
                         "('Từ ngày dd/mm/yyyy đến ngày dd/mm/yyyy')")
    ct = ChiTiet(tu_ngay=ngay[0], den_ngay=ngay[1])
    hien_tai: Optional[KhachChiTiet] = None
    for hang in ws.iter_rows(min_row=6, values_only=True):
        if not hang:
            continue
        a = hang[0]
        if isinstance(a, str):
            t = a.strip()
            if t.lower().startswith("tên khách hàng:"):
                ten = t.split(":", 1)[1].strip()
                tc = chuan_hoa_ten(ten)
                hien_tai = ct.khach.setdefault(tc, KhachChiTiet(ten=ten, ten_chuan=tc))
                continue
            if chuan_hoa_ten(t).startswith("TỔNG CỘNG"):
                break
        ngay_ht = _ngay(a)
        if hien_tai is None or ngay_ht is None or len(hang) < 10:
            continue                       # dòng "Số dư đầu kỳ", "Cộng", dòng trống
        hien_tai.but_toan.append(ButToan(
            ngay_ht=ngay_ht, ngay_hd=_ngay(hang[3]), so_hd=str(hang[4] or "").strip(),
            so_ct=str(hang[2] or "").strip(), dien_giai=str(hang[5] or "").strip(),
            tk_doi_ung=str(hang[7] or "").strip(), no=_so(hang[8]), co=_so(hang[9]),
        ))
    wb.close()
    if not ct.khach:
        raise LoiDocFile(f"{Path(path).name}: không có khách hàng nào")
    return ct


def ky_cua(path) -> Tuple[dt.date, dt.date]:
    """Chỉ đọc dòng tiêu đề để biết kỳ của file — file từ 2022 dài, đọc hết mất vài giây."""
    wb = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
    try:
        ngay = []
        for hang in wb.worksheets[0].iter_rows(min_row=1, max_row=3, values_only=True):
            for v in hang:
                ngay += _cac_ngay(v)
    finally:
        wb.close()
    if len(ngay) < 2:
        raise LoiDocFile(f"{Path(path).name}: không đọc được kỳ ở dòng 2")
    return ngay[0], ngay[1]

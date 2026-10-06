# -*- coding: utf-8 -*-
"""Nguồn thứ ba: file SMS (AR-AP) và bảng ghép khách SMS ↔ Misa.

SMS cho biết NGÀY của từng job còn nợ. Misa cho biết TIỀN. Hai hệ thống đặt tên
khách khác nhau (SMS tên tiếng Anh để xuất hoá đơn nước ngoài, Misa tên trong
nước), nên cần bảng ghép do FIN xác nhận một lần — xem
handoff/docs/NOTE-ke-toan-tra-loi-20261006.md.
"""
from __future__ import annotations

import datetime as dt
import difflib
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Set, Tuple

import openpyxl

from .misa import chuan_hoa_ten

# --- file SMS ------------------------------------------------------------

COT_BAT_BUOC = ("partner code", "a/r remain")


@dataclass
class DongSMS:
    ma_sms: str
    ten_sms: str
    salesman: str
    job: str
    huong: str               # Inbound | Outbound | Logistics | ""
    ngay: Optional[dt.date]  # ngày dùng để tính tuổi
    nguon_ngay: str          # ETD | ETA | job | ""
    so_tien: float           # còn nợ, đã quy VND


def _o(x) -> str:
    return str(x or "").strip()


def _ngay(x) -> Optional[dt.date]:
    if isinstance(x, dt.datetime):
        return x.date()
    if isinstance(x, dt.date):
        return x
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", _o(x))
    if not m:
        return None
    try:
        return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def _ngay_tu_job(job: str) -> Optional[dt.date]:
    """Mã job dạng EXSANA26090132: 26 = năm, 09 = tháng. Lấy ngày 15 của tháng."""
    m = re.search(r"[A-Z]{2,}ANA(\d\d)(\d\d)\d", job.upper())
    if not m:
        return None
    nam, thang = 2000 + int(m.group(1)), int(m.group(2))
    return dt.date(nam, thang, 15) if 1 <= thang <= 12 else None


def _dong_tieu_de(ws, toi_da: int = 10) -> Tuple[Optional[int], Dict[str, int]]:
    for r, hang in enumerate(ws.iter_rows(min_row=1, max_row=toi_da, values_only=True), start=1):
        ten = [_o(v).lower() for v in hang]
        if all(c in ten for c in COT_BAT_BUOC):
            return r, {t: i for i, t in enumerate(ten) if t}
    return None, {}


def la_file_sms(path) -> bool:
    try:
        wb = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
        r, _ = _dong_tieu_de(wb.worksheets[0])
        wb.close()
        return r is not None
    except Exception:
        return False


def doc_sms(path) -> List[DongSMS]:
    """Chỉ lấy các dòng phí còn nợ (A/R remain > 0). Mã khách và số job chỉ ghi ở
    dòng đầu của mỗi job, các dòng phí bên dưới thừa hưởng."""
    wb = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    r0, cot = _dong_tieu_de(ws)
    if r0 is None:
        raise ValueError(f"{Path(path).name}: không thấy dòng tiêu đề có Partner Code và A/R remain")

    def c(hang, ten, mac_dinh=None):
        i = cot.get(ten)
        return hang[i] if i is not None and i < len(hang) else mac_dinh

    ra: List[DongSMS] = []
    huong = ""
    hien_tai = None
    for hang in ws.iter_rows(min_row=r0 + 1, values_only=True):
        a, b = _o(hang[0] if hang else ""), _o(c(hang, "partner code"))
        if a.startswith("+"):                      # dòng nhóm: + Inbound / + Outbound
            huong = b.lstrip("+ ").strip()
            continue
        if b.lower().startswith(("sub-total", "grand-total", "balance")):
            continue
        if b or _o(c(hang, "file/job no.")):
            hien_tai = {
                "ma": b, "ten": _o(c(hang, "partner name")) or b,
                "sale": _o(c(hang, "salesman")), "job": _o(c(hang, "file/job no.")),
                "etd": _ngay(c(hang, "etd")), "eta": _ngay(c(hang, "eta")), "huong": huong,
            }
        con = c(hang, "a/r remain")
        if hien_tai is None or not isinstance(con, (int, float)) or con <= 0:
            continue
        roe = c(hang, "roe local")
        roe = roe if isinstance(roe, (int, float)) and roe > 0 else 1.0
        h = hien_tai
        thu_tu = (h["eta"], "ETA", h["etd"], "ETD") if h["huong"].lower() == "inbound" \
            else (h["etd"], "ETD", h["eta"], "ETA")
        ngay, nguon = (thu_tu[0], thu_tu[1]) if thu_tu[0] else (thu_tu[2], thu_tu[3])
        if ngay is None:
            ngay, nguon = _ngay_tu_job(h["job"]), "job"
            if ngay is None:
                nguon = ""
        ra.append(DongSMS(h["ma"], h["ten"], h["sale"], h["job"], h["huong"],
                          ngay, nguon, round(float(con) * roe, 2)))
    wb.close()
    return ra


# --- bảng ghép khách -----------------------------------------------------

@dataclass
class BangGhep:
    nguon: str = ""
    sms_sang_misa: Dict[str, str] = field(default_factory=dict)   # MÃ SMS (hoa) -> mã Misa
    khong_misa: Set[str] = field(default_factory=set)            # mã SMS FIN nói không có trên Misa
    kho_doi: Dict[str, str] = field(default_factory=dict)        # mã Misa -> ghi chú
    da_khai_misa: Set[str] = field(default_factory=set)          # mã Misa có mặt ở sheet 2
    da_khai_sms: Set[str] = field(default_factory=set)           # mã SMS có mặt ở sheet 1
    loi: List[Tuple[str, int, str, str]] = field(default_factory=list)       # sheet, dòng, mã, vấn đề
    xung_dot: List[Tuple[str, str, str]] = field(default_factory=list)       # mã SMS, mã cũ, mã thắng


def _k(ma) -> str:
    return _o(ma).upper()


def tim_file_ghep(thu_muc) -> Optional[Path]:
    """File ghép khách mới nhất trong thư mục sample (bỏ qua file KIEM_TRA tool sinh)."""
    if not thu_muc or not Path(thu_muc).is_dir():
        return None
    ung_vien = [f for f in Path(thu_muc).glob("*.xlsx")
                if not f.name.startswith(("~$", "KIEM_TRA"))]
    return max(ung_vien, key=lambda f: f.stat().st_mtime) if ung_vien else None


def _cot(hang, *tu_khoa) -> Optional[int]:
    for i, v in enumerate(hang):
        t = _o(v).lower()
        if all(k in t for k in tu_khoa):
            return i
    return None


def doc_bang_ghep(path) -> BangGhep:
    bg = BangGhep(nguon=Path(path).name)
    wb = openpyxl.load_workbook(Path(path), data_only=True)

    if "Ghép khách" in wb.sheetnames:
        ws = wb["Ghép khách"]
        tieu = [c.value for c in ws[1]]
        i_sms, i_goi_y = _cot(tieu, "mã sms"), _cot(tieu, "gợi ý", "mã")
        i_xn, i_dung = _cot(tieu, "xác nhận"), _cot(tieu, "mã misa đúng")
        i_muc = _cot(tieu, "tin cậy")
        for r, hang in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            sms = _k(hang[i_sms]) if i_sms is not None else ""
            if not sms:
                continue
            bg.da_khai_sms.add(sms)
            xn = _o(hang[i_xn]).lower() if i_xn is not None else ""
            goi_y = _o(hang[i_goi_y]) if i_goi_y is not None else ""
            dung = _o(hang[i_dung]) if i_dung is not None else ""
            muc = _o(hang[i_muc]).lower() if i_muc is not None else ""
            if xn.startswith("đúng") or xn.startswith("dung"):
                dich = goi_y
            elif xn.startswith("sai"):
                dich = dung
                if not dich:
                    bg.loi.append(("Ghép khách", r, sms, "Chọn Sai nhưng chưa điền Mã Misa đúng"))
                    continue
            elif "không có" in xn or "khong co" in xn:
                bg.khong_misa.add(sms)
                continue
            elif not xn and muc.startswith("khớp"):
                dich = goi_y                      # khớp tên/mã sẵn, FIN không cần chọn
            else:
                bg.loi.append(("Ghép khách", r, sms, "Chưa chọn Đúng / Sai / Không có trong Misa"))
                continue
            bg.sms_sang_misa[sms] = dich

    if "Misa chưa ghép" in wb.sheetnames:
        ws = wb["Misa chưa ghép"]
        tieu = [c.value for c in ws[1]]
        i_ma, i_sms = _cot(tieu, "mã khách misa"), _cot(tieu, "mã sms")
        i_chon, i_gc = _cot(tieu, "hoặc chọn"), _cot(tieu, "ghi chú")
        for r, hang in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            ma = _o(hang[i_ma]) if i_ma is not None else ""
            if not ma:
                continue
            bg.da_khai_misa.add(ma)
            sms = _k(hang[i_sms]) if i_sms is not None else ""
            chon = _o(hang[i_chon]).lower() if i_chon is not None else ""
            gc = _o(hang[i_gc]) if i_gc is not None else ""
            if "thu hồi" in gc.lower() or "khó đòi" in chon or "khó đòi" in gc.lower():
                bg.kho_doi[ma] = gc or "Khó đòi"
            if sms:
                cu = bg.sms_sang_misa.get(sms)
                if cu and cu != ma:
                    # NOTE-ke-toan-tra-loi-20261006 mục 1: phía Misa thắng
                    bg.xung_dot.append((sms, cu, ma))
                bg.sms_sang_misa[sms] = ma
                bg.khong_misa.discard(sms)
    wb.close()
    return bg


# --- gợi ý ghép theo tên --------------------------------------------------

_BO = set("CO LTD LIMITED LLC OOO JSC COMPANY CONG TY TNHH CO PHAN CP INC CORP CORPORATION "
          "PTE GROUP SP Z O THE AND JOINT STOCK MTV".split())


def _tu_khoa(s) -> List[str]:
    s = unicodedata.normalize("NFD", str(s or "")).replace("Đ", "D").replace("đ", "d")
    s = s.encode("ascii", "ignore").decode().upper()
    return [t for t in re.sub(r"[^A-Z0-9 ]", " ", s).split() if t not in _BO]


def diem_giong(a, b) -> float:
    A, B = _tu_khoa(a), _tu_khoa(b)
    if not A or not B:
        return 0.0
    j = len(set(A) & set(B)) / len(set(A) | set(B))
    return max(j, difflib.SequenceMatcher(None, " ".join(A), " ".join(B)).ratio())


def goi_y(ten: str, so_tien: float, ds: Iterable[Tuple[str, str, float]]) -> Tuple[str, str, float, str]:
    """ds: (mã Misa, tên Misa, số dư). Trả (mã, tên, số dư, mức tin cậy)."""
    ds = list(ds)
    if not ds:
        return "", "", 0.0, "Không có khách Misa nào"
    tot = max(ds, key=lambda x: diem_giong(ten, x[1]))
    d = diem_giong(ten, tot[1])
    if d >= 0.9:
        return tot[0], tot[1], tot[2], "Cao"
    cung_tien = [x for x in ds if abs(x[2] - so_tien) < 1000]
    if len(cung_tien) == 1:
        x = cung_tien[0]
        return x[0], x[1], x[2], "Trung bình (trùng số tiền)"
    return tot[0], tot[1], tot[2], "Trung bình" if d >= 0.6 else "Thấp — nhiều khả năng sai"


def ma_misa_cua(dong: DongSMS, bang: Optional[BangGhep], theo_ma: Dict[str, str],
                theo_ten: Dict[str, str]) -> Tuple[Optional[str], str]:
    """Trả (mã Misa hoặc None, cách ghép). theo_ma: MÃ (hoa) -> mã Misa;
    theo_ten: tên chuẩn hoá -> mã Misa."""
    k = _k(dong.ma_sms)
    if bang:
        if k in bang.khong_misa:
            return None, "FIN khai không có trên Misa"
        if k in bang.sms_sang_misa:
            return bang.sms_sang_misa[k], "bảng ghép"
    if k in theo_ma:
        return theo_ma[k], "trùng mã"
    tc = chuan_hoa_ten(dong.ten_sms)
    if tc in theo_ten:
        return theo_ten[tc], "trùng tên"
    return None, "chưa ghép"

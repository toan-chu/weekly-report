# -*- coding: utf-8 -*-
"""v0.4 — ghi các sheet của full credit report: Summary, AR Risk, Payable, Cash Flow,
Methodology. Sheet Receivables và Review vẫn do baocao.ghi_bao_cao ghi.

Ghi SỐ, không ghi công thức: dashboard (WR-05) đọc file bằng trình duyệt nên không
tính lại được công thức. Cách tính viết bằng lời ở sheet Methodology.
"""
from __future__ import annotations

import datetime as dt
from typing import List, Optional, Sequence

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: F401
from openpyxl.utils import get_column_letter

from . import chitiet as ctm
from . import thuonghieu as th
from .tindung import KetQuaFull

TIEN = "#,##0;[Red]-#,##0"
PHAN_TRAM = "0.0%"
THANG_EN = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_VIEN = Side(style="thin", color="D9D9D9")
VIEN = Border(left=_VIEN, right=_VIEN, top=_VIEN, bottom=_VIEN)


def _ngay_en(d: dt.date) -> str:
    return f"At {d.day:02d} {THANG_EN[d.month - 1]} {d.year}"


def _dau_trang(ws, ten_bao_cao: str, ngay_chot: dt.date, ten_bang: str, ky: str = "") -> None:
    ws["A1"], ws["A2"], ws["A3"], ws["A4"] = th.CONG_TY, ten_bao_cao, _ngay_en(ngay_chot), ten_bang
    if ky:
        ws["A3"] = f"{_ngay_en(ngay_chot)} · Reporting period {ky}"
    ws["A1"].font = Font(name=th.PHONG_TIEU_DE, size=12, bold=True, color=th.TIM)
    ws["A2"].font = Font(name=th.PHONG_TIEU_DE, size=11, bold=True)
    ws["A3"].font = Font(name=th.PHONG_NOI_DUNG, size=10, italic=True)
    ws["A4"].font = Font(name=th.PHONG_TIEU_DE, size=11, bold=True, color=th.CAM)


def bang(ws, hang: int, cot: Sequence[tuple], dong: List[list], tong: bool = True,
         mau_dong: Optional[List[Optional[str]]] = None) -> int:
    """cot: [(tên, loại 'ten'|'tien'|'so'|'pt', độ rộng)]. Trả về dòng kế tiếp còn trống."""
    for i, (ten, loai, rong) in enumerate(cot, start=1):
        o = ws.cell(row=hang, column=i, value=ten)
        o.fill = PatternFill("solid", fgColor=th.TIM if loai in ("tien", "pt") else th.CAM)
        o.font = Font(name=th.PHONG_TIEU_DE, size=10, bold=True, color=th.TRANG)
        o.alignment = Alignment(wrap_text=True, vertical="center",
                                horizontal="center" if loai != "ten" else "left")
        if rong:
            ws.column_dimensions[get_column_letter(i)].width = rong
    dau = hang + 1
    hang += 1
    for j, gia_tri in enumerate(dong):
        for i, v in enumerate(gia_tri, start=1):
            loai = cot[i - 1][1]
            if loai in ("tien", "so") and isinstance(v, (int, float)) and abs(v) < 0.4:
                v = None
            o = ws.cell(row=hang, column=i, value=v)
            o.font = Font(name=th.PHONG_NOI_DUNG, size=10)
            o.border = VIEN
            if loai == "tien":
                o.number_format = TIEN
            elif loai == "pt":
                o.number_format = PHAN_TRAM
            else:
                o.alignment = Alignment(vertical="top", wrap_text=loai == "ten" and cot[i - 1][2] >= 30)
            if mau_dong and mau_dong[j]:
                o.fill = PatternFill("solid", fgColor=mau_dong[j])
        hang += 1
    if tong and dong:
        for i, (ten, loai, _) in enumerate(cot, start=1):
            o = ws.cell(row=hang, column=i)
            o.fill = PatternFill("solid", fgColor=th.TIM)
            o.font = Font(name=th.PHONG_TIEU_DE, size=10, bold=True, color=th.TRANG)
            if i == 2 or (i == 1 and cot[1][1] != "ten"):
                o.value = "TOTAL"
            if loai == "tien":
                o.value = round(sum(r[i - 1] for r in dong if isinstance(r[i - 1], (int, float))), 2)
                o.number_format = TIEN
        hang += 1
    ws.freeze_panes = ws.cell(row=dau, column=4 if len(cot) > 6 else 1)
    return hang


# --- Summary ---------------------------------------------------------------

CHI_SO = [
    ("tong_phai_thu", "Total receivables", "tien"),
    ("chua_den_han", "  Not yet due", "tien"),
    ("qua_han", "  Total overdue", "tien"),
    ("qua_han_1_30", "    Overdue 1 - 30 days", "tien"),
    ("qua_han_31_60", "    Overdue 31 - 60 days", "tien"),
    ("qua_han_61_90", "    Overdue 61 - 90 days", "tien"),
    ("qua_han_91_120", "    Overdue 91 - 120 days", "tien"),
    ("qua_han_120", "    Overdue over 120 days", "tien"),
    ("so_khach", "Customers with balance", "so"),
    ("no_moi", "New invoices this week", "tien"),
    ("da_thu", "Collected this week (all credits)", "tien"),
    ("thu_ngan_hang", "  of which: bank", "tien"),
    ("thu_tien_mat", "  of which: cash", "tien"),
    ("thu_can_tru", "  of which: offset with payables", "tien"),
    ("thu_ty_gia_phi", "  of which: FX & bank fees", "tien"),
    ("thu_khac", "  of which: other", "tien"),
    ("no_xau", "Bad debt (overdue of N3 + N4)", "tien"),
    ("so_khach_no_xau", "  Bad-debt customers", "so"),
    ("no_xau_phap_ly", "  of which legal event (N4)", "tien"),
    ("no_xau_120", "  of which over 120 days", "tien"),
    ("ceo_can_xem", "Customers CEO must review (overdue ≥ 100M, no reason)", "so"),
    ("tong_phai_tra", "Total payables", "tien"),
    ("no_vendor_moi", "New vendor bills this week", "tien"),
    ("da_tra_vendor", "Paid to vendors this week", "tien"),
    ("so_vendor", "Vendors with balance", "so"),
    ("dong_tien_rong", "Net cash (bank + cash collected − paid to vendors)", "tien"),
]

LICH_SU = [
    ("tuan", "Week", "ten", 8), ("den_ngay", "Period end", "ten", 12),
    ("tong_phai_thu", "Total AR", "tien", 16), ("chua_den_han", "Not yet due", "tien", 15),
    ("qua_han", "Overdue", "tien", 15), ("qua_han_61_90", "61 - 90", "tien", 14),
    ("qua_han_91_120", "91 - 120", "tien", 14), ("qua_han_120", "120+", "tien", 14),
    ("da_thu", "Collected", "tien", 15), ("da_thu_tien", "Collected (bank + cash)", "tien", 15),
    ("no_moi", "New invoices", "tien", 15), ("no_xau", "Bad debt", "tien", 15),
    ("tong_phai_tra", "Total AP", "tien", 15), ("da_tra_vendor", "Paid to vendors", "tien", 15),
    ("no_vendor_moi", "New vendor bills", "tien", 15), ("dong_tien_rong", "Net cash", "tien", 15),
]


def _thay_doi(cs: dict, tr: dict, khoa: str) -> str:
    nay, cu = cs.get(khoa), tr.get(khoa)
    if not isinstance(cu, (int, float)) or not cu or nay is None:
        return "no last-week figure"
    return f"{'▲' if nay >= cu else '▼'} {abs(nay - cu):,.0f} vs last week ({(nay - cu) / cu:+.1%})"


def _o_so(ws, hang: int, a: str, b: str, nhan: str, gia_tri, phu: str, mau: str) -> None:
    """Một ô số lớn kiểu thẻ: nhãn nhỏ · số to · dòng phụ."""
    for r, v, f in ((hang, nhan, Font(name=th.PHONG_NOI_DUNG, size=9, bold=True, color=mau)),
                    (hang + 1, gia_tri, Font(name=th.PHONG_TIEU_DE, size=16, bold=True, color=mau)),
                    (hang + 2, phu, Font(name=th.PHONG_NOI_DUNG, size=8, color="7F7F7F"))):
        if a != b:
            ws.merge_cells(f"{a}{r}:{b}{r}")
        o = ws[f"{a}{r}"]
        o.value = v
        o.font = f
        o.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if isinstance(v, (int, float)):
            o.number_format = "#,##0"
        for c in range(ord(a), ord(b) + 1):
            ws[f"{chr(c)}{r}"].fill = PatternFill("solid", fgColor="F7F5FB")
    ws[f"{a}{hang}"].border = Border(top=Side(style="medium", color=mau))
    ws.row_dimensions[hang + 1].height = 24


def ghi_summary(wb, kq: KetQuaFull, ngay_chot: dt.date, lich_su: List[dict]) -> None:
    ws = wb.create_sheet(th.SHEET_SUMMARY)
    cs, tr = kq.chi_so, kq.chi_so_truoc
    ky = f"{dt.date.fromisoformat(cs['tu_ngay']):%d/%m/%Y} - {dt.date.fromisoformat(cs['den_ngay']):%d/%m/%Y}"
    _dau_trang(ws, "Weekly Credit Summary", ngay_chot, "Key figures", ky)
    tong = cs.get("tong_phai_thu") or 0
    o_so = [   # 3 hàng thẻ: phải thu · phải trả · dòng tiền. Cột F hẹp nên chỉ để số đếm.
        ("RECEIVABLES · Total", cs.get("tong_phai_thu"), _thay_doi(cs, tr, "tong_phai_thu"), th.TIM),
        ("Overdue", cs.get("qua_han"), f"{(cs.get('qua_han') or 0) / tong:.0%} of receivables" if tong else "",
         th.CAM),
        ("Bad debt (N3 + N4)", cs.get("no_xau"), f"{cs.get('so_khach_no_xau', 0)} customers", "C00000"),
        ("CEO must review", cs.get("ceo_can_xem"), "overdue ≥ 100M, no reason", "C00000"),
        ("PAYABLES · Total", cs.get("tong_phai_tra"), _thay_doi(cs, tr, "tong_phai_tra"), th.TIM),
        ("Paid to vendors this week", cs.get("da_tra_vendor"), "Misa TK331 debit", th.CAM),
        ("New vendor bills this week", cs.get("no_vendor_moi"), "Misa TK331 credit", th.CAM),
        ("Vendors", cs.get("so_vendor"), "with balance", th.TIM),
        ("CASH · Collected this week", cs.get("da_thu"), "all credits to TK131", "2E7D32"),
        ("Collected to bank + cash", cs.get("da_thu_tien"), "real money in", "2E7D32"),
        ("Net cash this week", cs.get("dong_tien_rong"), "bank + cash collected − paid to vendors", "2E7D32"),
        ("Customers", cs.get("so_khach"), "with balance", th.TIM),
    ]
    vi_tri = [("A", "A"), ("B", "C"), ("D", "E"), ("F", "F")]
    for i, (nhan, gia_tri, phu, mau) in enumerate(o_so):
        r = 6 + 4 * (i // 4)
        a, b = vi_tri[i % 4]
        _o_so(ws, r, a, b, nhan, gia_tri, phu, mau)
    ws.cell(row=18, column=1, value="This week vs last week").font = Font(
        name=th.PHONG_TIEU_DE, size=11, bold=True, color=th.CAM)
    dong = []
    for khoa, nhan, loai in CHI_SO:
        if khoa not in cs:
            continue
        nay, cu = cs.get(khoa), tr.get(khoa)
        chenh = (nay - cu) if isinstance(cu, (int, float)) else None
        ty_le = (chenh / cu) if chenh is not None and cu else None
        dong.append([nhan, cu, nay, chenh, ty_le, khoa])
    hang = bang(ws, 19, [("Metric", "ten", 52), ("Last week", "tien", 18), ("This week", "tien", 18),
                        ("Change", "tien", 18), ("Change %", "pt", 11), ("key", "ten", 18)],
                dong, tong=False)
    for r in range(20, hang):
        ws.cell(row=r, column=6).font = Font(name=th.PHONG_NOI_DUNG, size=8, color="A6A6A6")
    hang += 1
    ws.cell(row=hang, column=1, value="Weekly history").font = Font(
        name=th.PHONG_TIEU_DE, size=11, bold=True, color=th.CAM)
    hang += 1
    cot = [(nhan, loai, rong) for _, nhan, loai, rong in LICH_SU]
    dl = [[x.get(k) for k, *_ in LICH_SU] for x in lich_su]
    bang(ws, hang, cot, dl, tong=False)
    # bảng lịch sử dùng chung cột với bảng trên: đặt lại độ rộng cho bảng trên dễ đọc
    for chu, rong in zip("ABCDEF", (52, 18, 18, 18, 15, 18)):
        ws.column_dimensions[chu].width = rong
    ws.freeze_panes = None
    ws.sheet_view.showGridLines = False


# --- AR Risk ---------------------------------------------------------------

COT_RISK = [
    ("No.", "so", 6), ("Code", "ten", 14), ("Customer's Name", "ten", 40), ("Credit Term", "ten", 14),
    ("Salesman", "ten", 12), ("Opening (Misa)", "tien", 15), ("Increase", "tien", 15),
    ("Decrease", "tien", 15), ("Ending = Total", "tien", 15), ("Backdated vs last report", "tien", 15),
    ("Not yet due", "tien", 14), ("1 - 30", "tien", 14), ("31 - 60", "tien", 14), ("61 - 90", "tien", 14),
    ("91 - 120", "tien", 14), ("120+", "tien", 14), ("Total overdue", "tien", 15),
    ("Last week total", "tien", 15), ("Last week overdue", "tien", 15), ("Last week 61+", "tien", 15),
    ("S1", "so", 5), ("S2", "so", 5), ("S3", "so", 5), ("S4", "so", 5), ("S5", "so", 5), ("S6", "so", 5),
    ("Signals (S1-S3)", "so", 9), ("Risk Code", "ten", 14), ("Risk Tier", "ten", 26),
    ("CEO review", "ten", 9), ("Trend (auto)", "ten", 50), ("Reason for late payment", "ten", 36),
    ("Required action", "ten", 40), ("Job Number", "ten", 30),
]


def ghi_risk(wb, kq: KetQuaFull, ngay_chot: dt.date) -> None:
    ws = wb.create_sheet(th.SHEET_RISK)
    _dau_trang(ws, "Receivables Movement & Risk", ngay_chot, "Risk tier per customer")
    dong, mau = [], []
    for i, r in enumerate(kq.rui_ro, start=1):
        dong.append([i, r.ma, r.ten, r.credit_term, r.salesman, r.dau_ky, r.tang, r.giam, r.cuoi_ky,
                     r.nhap_bo_sung, *r.nhom, r.qua_han, r.tt_total, r.tt_qua_han, r.tt_61, *r.s,
                     r.so_dau_hieu, r.ma_dau_hieu, r.nhom_rui_ro, "Yes" if r.ceo else "",
                     r.xu_huong, r.ly_do or "(Sales chưa ghi — xem Trend)", r.viec_can_lam, r.job])
        mau.append(th.DO_CANH_BAO if r.nhom_rui_ro[:2] in ("N3", "N4") else
                   th.VANG_CANH_BAO if r.nhom_rui_ro[:2] == "N2" else None)
    hang = bang(ws, 5, COT_RISK, dong, mau_dong=mau)
    for c in range(21, 28):     # cột dấu hiệu: số 0 vẫn hiện để đọc được
        for r in range(6, 6 + len(dong)):
            o = ws.cell(row=r, column=c)
            if o.value is None:
                o.value = 0
            o.alignment = Alignment(horizontal="center")
    ws.auto_filter.ref = f"A5:{get_column_letter(len(COT_RISK))}{5 + len(dong)}"
    if kq.khach_tra_xong:
        hang += 1
        ws.cell(row=hang, column=2, value="Customers fully paid since last report").font = Font(
            name=th.PHONG_TIEU_DE, size=11, bold=True, color=th.CAM)
        bang(ws, hang + 1, [("No.", "so", 6), ("Customer's Name", "ten", 14), ("Last report balance", "tien", 40)],
             [[i, t, v] for i, (t, v) in enumerate(kq.khach_tra_xong, start=1)])
    ws.freeze_panes = "D6"


# --- Payable ---------------------------------------------------------------

def ghi_payable(wb, kq: KetQuaFull, ngay_chot: dt.date, ky: str) -> None:
    ws = wb.create_sheet(th.SHEET_PAYABLE)
    _dau_trang(ws, th.TEN_BAO_CAO_TRA, ngay_chot, th.TEN_BANG_TRA, ky)
    if kq.phai_tra is None:
        ws["A6"] = "Tuần này chưa có file Tổng hợp công nợ phải trả (TK331) trong input — sheet để trống."
        ws["A6"].font = Font(name=th.PHONG_NOI_DUNG, size=10, bold=True, color="C00000")
        return
    # FIN 09/10: "làm y chang mẫu của em... chỉ cần lấy số tổng, không phân tích gì cả" —
    # 5 cột như sheet Payable của FIN, chỉ vendor còn số dư, lớn trước. Tiền đã trả nằm ở Cash Flow.
    cot = [("No.", "so", 6), ("Code", "ten", 20), ("Vendor's Name", "ten", 54), ("Current", "tien", 20),
           ("Total", "tien", 20)]
    con = sorted((v for v in kq.phai_tra if v.cuoi_ky > 0.5), key=lambda v: -v.cuoi_ky)
    dong = [[i, v.ma, v.ten, v.cuoi_ky, v.cuoi_ky] for i, v in enumerate(con, start=1)]
    bang(ws, 5, cot, dong)
    ws.auto_filter.ref = f"A5:E{5 + len(dong)}"
    ws.freeze_panes = "D6"


# --- Cash Flow -------------------------------------------------------------

def ghi_cash(wb, kq: KetQuaFull, ngay_chot: dt.date, ky: str) -> None:
    ws = wb.create_sheet(th.SHEET_CASH)
    _dau_trang(ws, "Cash Movement", ngay_chot, "A. Collections from customers", ky)
    nhom = ["ngan-hang", "tien-mat", "can-tru", "ty-gia-phi", "khac"]
    cot = [("No.", "so", 6), ("Code", "ten", 16), ("Customer's Name", "ten", 46)] + \
          [(ctm.TEN_NHOM_THU[k], "tien", 16) for k in nhom] + [("Total", "tien", 16)]
    dong = [[i, x.ma, x.ten] + [x.theo_nhom.get(k, 0.0) for k in nhom] + [x.tong]
            for i, x in enumerate(kq.thu, start=1)]
    hang = bang(ws, 5, cot, dong)
    hang += 1
    ws.cell(row=hang, column=1, value="B. Payments to vendors (Misa TK331 debit — may include offsets)"
            ).font = Font(name=th.PHONG_TIEU_DE, size=11, bold=True, color=th.CAM)
    hang += 1
    if kq.phai_tra is None:
        ws.cell(row=hang, column=1, value="Chưa có file Tổng hợp công nợ phải trả tuần này.")
        hang += 2
    else:
        dong_b = [[i, v.ma, v.ten, v.da_tra, v.cuoi_ky] for i, v in enumerate(kq.tra_vendor, start=1)]
        hang = bang(ws, hang, [("No.", "so", 6), ("Code", "ten", 16), ("Vendor's Name", "ten", 46),
                               ("Paid this week", "tien", 16), ("Closing", "tien", 16)], dong_b)
    hang += 1
    ws.cell(row=hang, column=1, value="C. Net").font = Font(
        name=th.PHONG_TIEU_DE, size=11, bold=True, color=th.CAM)
    cs = kq.chi_so
    tom = [["Cash in (bank + cash collected)", cs.get("da_thu_tien")],
           ["Cash out (paid to vendors)", cs.get("da_tra_vendor")],
           ["Net cash", cs.get("dong_tien_rong")]]
    bang(ws, hang + 1, [("Item", "ten", 46), ("Amount", "tien", 18)], tom, tong=False)


# --- Invoices -------------------------------------------------------------

def ghi_invoices(wb, kq: KetQuaFull, ngay_chot: dt.date) -> None:
    """Mỗi dòng một hoá đơn còn nợ. Lọc theo một khách, cộng cột Amount = Total ở Receivables."""
    from . import soghi as sg
    ws = wb.create_sheet(th.SHEET_INVOICES)
    den = dt.date.fromisoformat(kq.chi_so["den_ngay"])
    _dau_trang(ws, "Open invoices behind every receivable", ngay_chot,
               f"Filter one customer: the Amount column adds up to its Total in {th.SHEET_BANG}")
    cot = [("No.", "so", 6), ("Code", "ten", 14), ("Customer's Name", "ten", 40), ("Invoice No.", "ten", 20),
           ("Invoice Date", "ten", 12), ("Credit Term", "ten", 14), ("Due Date", "ten", 12),
           ("Days Overdue", "so", 9), ("Bucket", "ten", 12), ("Amount", "tien", 16), ("Salesman", "ten", 24),
           ("Note", "ten", 30)]
    if kq.job_misa:                       # v0.5 (WR-06): trạng thái + job theo Misa (Bán hàng, Sổ chi tiết)
        cot += [("Misa status", "ten", 18), ("Job (Misa)", "ten", 30)]
    dong, i = [], 0
    for d in kq.dong:
        if d.hoa_don:
            for ngay, so_hd, han, tien in d.hoa_don:
                qua = (den - han).days
                i += 1
                dong.append([i, d.ma, d.ten, so_hd, ngay.strftime("%d/%m/%Y") if ngay else "",
                             d.credit_term, han.strftime("%d/%m/%Y"), max(qua, 0),
                             sg.TEN_NHOM[sg.nhom_tu_so_ngay(qua)], tien, d.salesman,
                             "Tổng hợp 131 cao hơn Chi tiết 131 — tạm coi là nợ mới" if ngay is None else ""])
                if kq.job_misa:
                    n = ctm._so_hd(str(so_hd)) if ngay else None
                    tt, jobs = d.hd_misa.get((n, ngay.year), ("", "")) if n is not None else ("", "")
                    if not tt and ngay:
                        tt = "(trước năm nay)" if ngay.year < den.year else "(không thấy trên Bán hàng)"
                    dong[-1] += [tt, jobs]
        else:
            i += 1
            nhom = sg.TEN_NHOM[d.nhom_gia_nhat]
            dong.append([i, d.ma, d.ten, "(whole balance)", "", d.credit_term, "", None, nhom, d.total,
                         d.salesman, {"kho-doi": "Hard-to-collect: always 120+",
                                      "ngoai-le": "Exception declared by accounting"}.get(d.nguon, d.nguon)])
            if kq.job_misa:
                dong[-1] += ["", ""]
    bang(ws, 5, cot, dong)
    ws.auto_filter.ref = f"A5:{get_column_letter(len(cot))}{5 + len(dong)}"
    ws.freeze_panes = "D6"


# --- Methodology -----------------------------------------------------------

CACH_TINH = [
    ("All sheets", "Money", "Every amount is VND as recorded in Misa. Excel holds values, not formulas.",
     "—"),
    ("Receivables", "Total", "Ending debit balance of the customer.",
     "Misa — Tổng hợp công nợ phải thu (TK131), column Số dư cuối kỳ · Nợ"),
    ("Receivables", "Customer list", "Customers whose ending debit balance > 0. Prepaid customers "
     "(credit balance) are not listed.", "Misa — Tổng hợp công nợ phải thu"),
    ("Receivables", "Ageing buckets", "Each invoice: due date = invoice date + credit term ('14 ngày của "
     "tháng tiếp theo' = 14th of the next month). A receipt that names invoice numbers in its "
     "description ('theo HD 00001574, 00001419') pays those invoices; the rest of the money pays the "
     "OLDEST open invoice first. Days overdue = period end − due date; the due date itself is still "
     "Current, overdue starts the next day. Buckets: Current, 1-30, 31-60, "
     "61-90, 91-120, 120+.", "Misa — Chi tiết công nợ phải thu (TK131) from 01/01/2022: Ngày hoá đơn, "
     "Số hoá đơn, PS Nợ, PS Có"),
    ("Receivables", "Hard-to-collect", "Whole balance in 120+ when FIN marked the customer "
     "'không có khả năng thu hồi' or the reason says 'khó đòi'.", "sample/ ghép khách · last report"),
    ("Receivables", "Credit Term, Reason", "Carried from last week's report (as edited by "
     "accounting).", "Report of last week in the workspace root"),
    ("Receivables", "Salesman, Job Number", "Open jobs of the customer in SMS (A/R remain > 0), "
     "matched through the customer mapping file.", "SMS AR-AP · sample/ ghép khách"),
    ("AR Risk", "Opening / Increase / Decrease", "Opening and movements of the week from Misa. "
     "Opening + Increase − Decrease = Ending.", "Misa — Tổng hợp TK131: Số dư đầu kỳ, Phát sinh Nợ, "
     "Phát sinh Có"),
    ("AR Risk", "Backdated vs last report", "Misa opening − last report's Total. Not zero means "
     "invoices or receipts were posted later with an earlier date.", "Misa + last report"),
    ("AR Risk", "S1", "Balance grew vs last week AND overdue did not decrease (new lots, old debt "
     "not paid).", "this week vs last week"),
    ("AR Risk", "S2", "Overdue > 0 and exactly equal to last week (nothing collected).",
     "this week vs last week"),
    ("AR Risk", "S3", "Overdue over 60 days (61-90 + 91-120 + 120+) larger than last week "
     "(debt is ageing).", "this week vs last week"),
    ("AR Risk", "S4 / S5 / S6", "S4 no reason from sales · S5 no credit term · S6 whole balance "
     "overdue. Informational, not counted.", "Receivables"),
    ("Receivables", "Reason for late payment", "Written by sales/accounting. When empty, the tool "
     "writes '[TỰ ĐỘNG] …' (grey italic) from the trend rules below; that text is ignored next week.",
     "Last week's report · AR Risk Trend"),
    ("Receivables", "Salesman", "Full name as in SMS (CEO 08/10). Short names in old reports are "
     "converted (Hiếu → Trần Văn Hiếu, Hiếu 1 → Phạm Trần Hiếu …).", "SMS AR-AP · thuonghieu.py"),
    ("Receivables", "Rows under a customer (+ / 1 | 2)", "Click + (or the 2 button top-left) to open "
     "the jobs behind a customer: each SMS job still owed, its salesman and amount. The last line "
     "balances to the Total (old debt without an SMS job, or SMS over-stated). Jobs carry no invoice "
     "date, so the customer's ageing buckets are shared out: old debt without a job and the oldest "
     "jobs take the oldest buckets first; rows FIN mapped to a known invoice (sample/"
     "Tham_chieu_job_no_cu.xlsx) take that invoice's own bucket. Columns of the rows always add up "
     "to the customer row. (Used only when the Bán hàng file is missing.)", "SMS AR-AP"),
    ("Receivables", "Rows under a customer — from Misa (v0.5)", "When FIN drops the Bán hàng and Sổ chi "
     "tiết TK131 files (from 1 January this year): one row per job of every invoice NOT marked 'Đã thanh "
     "toán' — 'Chưa thanh toán' and 'Thanh toán một phần' are both listed at the full invoice amount "
     "(FIN 09/10). Job = column 'Mã đối tượng THCP'; amount at the invoice exchange rate, so it matches "
     "the books. Each row sits in the bucket of its own invoice due date. Salesman of the job from SMS. "
     "Invoices before this year: job from sample/Tham_chieu_job_no_cu.xlsx.",
     "Misa — Bán hàng · Sổ chi tiết tài khoản 131 · SMS (salesman)"),
    ("Receivables", "Đối chiếu job (FIN)", "Customer Total − sum of its job rows. Khớp: within 1,000. "
     "VÀNG: gap ≤ 2% (exchange-rate revaluation forgotten, rounding). ĐỎ: jobs exceed the balance "
     "(paid but not matched to the invoice in Misa, or partial payment not deducted). CAM: balance "
     "exceeds jobs (debt without a job — older than this year, or invoice missing from Bán hàng). "
     "Review lists each colour, plus customers with no balance whose invoices still show unpaid (ĐỎ) "
     "and names on Bán hàng not found in Misa (TÍM). Rows are NOT forced to add up: the gap is what "
     "FIN checks.", "Misa — Tổng hợp 131 vs Bán hàng"),
    ("Invoices", "Open invoices", "Every invoice still owed after payments are applied (see Ageing "
     "buckets). Sum per customer = Total in Receivables.", "Misa — Chi tiết TK131"),
    ("AR Risk", "Reason used for the tier", "Reason for late payment carried from last week's report. "
     "If accounting writes a new reason in this report, the tier is recalculated next week (and by "
     "the dashboard).", "Receivables"),
    ("AR Risk", "Risk Tier", "0 no overdue → X reason mentions offset/payment on behalf → N4 reason "
     "mentions legal event (đang kiện, mất liên lạc, khó đòi, công nợ xấu, phá sản) → N3 any 120+ or "
     "≥ 2 signals (S1-S3) → N2 one signal → N1 otherwise.", "Rules of the chief accountant "
     "(sheet Tham chieu)"),
    ("AR Risk", "CEO review", "Overdue ≥ 100,000,000 and no reason yet.", "Tham chieu"),
    ("Payable", "Current / Total", "Ending credit balance of every vendor that still has one; no filter, "
     "no ageing (same layout as FIN's Payable sheet, FIN 09/10). Both columns hold the same amount.",
     "Misa — Tổng hợp công nợ phải trả (TK331), Số dư cuối kỳ · Có"),
    ("Cash Flow", "Payments to vendors", "Debit movements of TK331 in the week, per vendor.",
     "Misa — Tổng hợp TK331: Phát sinh Nợ"),
    ("Cash Flow", "Collections", "Credit lines of TK131 posted in the week, split by contra account: "
     "112 bank, 111 cash, 331 offset with payables, 515/635/642/811 FX & bank fees, other.",
     "Misa — Chi tiết TK131: TK đối ứng, PS Có"),
    ("Cash Flow", "Net cash", "(bank + cash collected) − paid to vendors. Ledger view, not a bank "
     "statement.", "Cash Flow A + B"),
    ("Summary", "Last week", "Figures stored by the tool when last week's report was produced.",
     "_tool/so-ghi.json"),
]


def ghi_methodology(wb, ngay_chot: dt.date, nguon: List[str]) -> None:
    ws = wb.create_sheet(th.SHEET_METHOD)
    _dau_trang(ws, "How every figure is calculated", ngay_chot, "Methodology")
    hang = bang(ws, 5, [("Sheet", "ten", 14), ("Figure", "ten", 28), ("How it is calculated", "ten", 90),
                        ("Source", "ten", 60)], [list(x) for x in CACH_TINH], tong=False)
    hang += 1
    ws.cell(row=hang, column=1, value="Source files of this report").font = Font(
        name=th.PHONG_TIEU_DE, size=11, bold=True, color=th.CAM)
    for i, n in enumerate(nguon, start=hang + 1):
        ws.cell(row=i, column=2, value=n).font = Font(name=th.PHONG_NOI_DUNG, size=10)


THU_TU_SHEET = [th.SHEET_SUMMARY, th.SHEET_BANG, th.SHEET_RISK, th.SHEET_PAYABLE, th.SHEET_CASH,
                th.SHEET_INVOICES, th.SHEET_METHOD, th.SHEET_REVIEW]


def sap_xep(wb) -> None:
    thu_tu = [n for n in THU_TU_SHEET if n in wb.sheetnames] + \
             [n for n in wb.sheetnames if n not in THU_TU_SHEET]
    wb._sheets = [wb[n] for n in thu_tu]
    wb.active = 0
    for ws in wb.worksheets:      # in / xuất PDF: ngang trang, vừa một trang chiều rộng
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4

# -*- coding: utf-8 -*-
"""Ghi file báo cáo tuần theo mẫu chuẩn, và đọc lại báo cáo cũ để nạp sổ ghi."""
from __future__ import annotations

import datetime as dt
import re
import shutil
from copy import copy
from pathlib import Path
from typing import Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import soghi as sg
from . import thuonghieu as th
from .misa import chuan_hoa_ten
from .tinhtoan import DongBaoCao

DINH_DANG_TIEN = "#,##0;[Red]-#,##0"
THANG_EN = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def ten_file_bao_cao(tu_ngay: dt.date, den_ngay: dt.date) -> str:
    tuan = den_ngay.isocalendar()[1]
    return (f"{den_ngay.year}_W{tuan:02d}_Credit_Report_"
            f"{tu_ngay:%d.%m}-{den_ngay:%d.%m.%Y}.xlsx")


def ghep_ghi_chu(job: str, ghi_chu: str) -> str:
    """Cột Ghi chú của mẫu chuẩn chứa cả số job: 'job: A, B | ghi chú'."""
    phan = []
    if job:
        phan.append(f"job: {job}")
    if ghi_chu:
        phan.append(ghi_chu)
    return " | ".join(phan)


def tach_ghi_chu(o: str):
    """Ngược với ghep_ghi_chu. Trả (job, ghi chú)."""
    o = str(o or "").strip()
    if not o.lower().startswith("job:"):
        return "", o
    dau, _, sau = o.partition(" | ")
    return dau[4:].strip(), sau.strip()


def _ngay_en(d: dt.date) -> str:
    return f"At {d.day:02d} {THANG_EN[d.month - 1]} {d.year}"


def ghi_bao_cao(
    dong: List[DongBaoCao],
    ngay_chot: dt.date,
    mau_chuan: Path,
    dich: Path,
    tong_misa: Optional[float] = None,
    doi_chieu: Optional[dict] = None,
    them_sheet=None,
) -> Path:
    dich.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(mau_chuan, dich)
    wb = openpyxl.load_workbook(dich)
    ws = wb[th.SHEET_BANG] if th.SHEET_BANG in wb.sheetnames else wb.worksheets[0]
    ws.title = th.SHEET_BANG
    # chữ tiêu đề lấy từ thuonghieu.py mỗi lần chạy, sửa ở đó là đủ
    ws["A1"], ws["A2"], ws["A4"] = th.CONG_TY, th.TEN_BAO_CAO, th.TEN_BANG
    for i, (ten_cot, _, rong) in enumerate(th.COT, start=1):
        ws.cell(row=th.DONG_TIEU_DE, column=i).value = ten_cot
        ws.column_dimensions[get_column_letter(i)].width = rong
    ws["A3"] = _ngay_en(ngay_chot)

    vien = Side(style="thin", color="D9D9D9")
    kieu_vien = Border(left=vien, right=vien, top=vien, bottom=vien)
    # v0.5 (WR-06): cột Đối chiếu — dư nợ Misa so với tổng job chưa thanh toán, tô màu cho FIN
    co_doi_chieu = any(d.doi_chieu for d in dong)
    n_cot = len(th.COT) + (1 if co_doi_chieu else 0)
    if co_doi_chieu:
        _cot_doi_chieu(ws)
    hang = th.DONG_DAU_DU_LIEU
    for stt, d in enumerate(dong, start=1):
        gia_tri = [stt, d.ma, d.ten, d.credit_term, d.total] + \
                  [v if abs(v) > 0.4 else None for v in d.nhom] + \
                  [d.salesman, d.job, d.ly_do or d.ly_do_tu_dong]
        for cot, v in enumerate(gia_tri, start=1):
            o = ws.cell(row=hang, column=cot, value=v)
            o.font = Font(name=th.PHONG_NOI_DUNG, size=10)
            o.border = kieu_vien
            if 5 <= cot <= 11:
                o.number_format = DINH_DANG_TIEN
                o.alignment = Alignment(horizontal="right")
            elif cot == 1:
                o.alignment = Alignment(horizontal="center")
            else:
                # Job Number không xuống dòng (danh sách job đầy đủ nằm ở các dòng con)
                o.alignment = Alignment(vertical="top", wrap_text=cot == 14)
        if not d.ly_do and d.ly_do_tu_dong:
            ws.cell(row=hang, column=14).font = Font(name=th.PHONG_NOI_DUNG, size=10, italic=True,
                                                     color="7F7F7F")
        if d.cho_xem_lai:
            mau = (th.DO_CANH_BAO if d.nguon in {"khoi-dong-lanh", "ngoai-le", "uoc-tinh"}
                   else th.VANG_CANH_BAO)
            for cot in range(1, len(th.COT) + 1):
                ws.cell(row=hang, column=cot).fill = PatternFill("solid", fgColor=mau)
        if co_doi_chieu:
            _o_doi_chieu(ws, hang, d.doi_chieu, kieu_vien)
        cha = hang
        hang += 1
        hang = _ghi_dong_con(ws, d, hang, kieu_vien, n_cot)
        if hang > cha + 1:
            ws.cell(row=cha, column=3).font = Font(name=th.PHONG_NOI_DUNG, size=10, bold=True)
            ws.row_dimensions[cha].collapsed = True

    hang_tong = hang
    for cot in range(1, n_cot + 1):
        c = ws.cell(row=hang_tong, column=cot)
        c.fill = PatternFill("solid", fgColor=th.TIM)
        c.font = Font(name=th.PHONG_TIEU_DE, size=10, bold=True, color=th.TRANG)
        if 5 <= cot <= 11:
            # ghi SỐ, không ghi công thức: dòng con (job) nằm xen giữa, SUM cả cột sẽ cộng trùng
            c.value = round(d_tong(dong, cot), 2)
            c.number_format = DINH_DANG_TIEN
            c.alignment = Alignment(horizontal="right")
    ws.cell(row=hang_tong, column=2, value="TOTAL")
    from openpyxl.worksheet.properties import Outline
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=True)

    ws.auto_filter.ref = f"A{th.DONG_TIEU_DE}:{get_column_letter(n_cot)}{hang_tong - 1}"
    _ghi_sheet_xem_lai(wb, dong, ngay_chot, tong_misa, doi_chieu or {})
    if them_sheet:
        them_sheet(wb)
    wb.save(dich)
    return dich


def d_tong(dong: List[DongBaoCao], cot: int) -> float:
    """Tổng một cột tiền của các dòng KHÁCH (cột 5 = Total, 6..11 = sáu nhóm tuổi)."""
    return sum(d.total if cot == 5 else d.nhom[cot - 6] for d in dong)


NHAN_CON = "    └ "


# v0.5 (WR-06): màu đối chiếu — đỏ/vàng/cam là việc FIN cần xem lại, tím là lỗi ghép tên
MAU_DC = {"khop": ("E2F0D9", "375623"), "do": ("F8CBCB", "9C0006"), "vang": ("FFF2CC", "7F6000"),
          "cam": ("FBE0C8", "843C0C"), "tim": ("E4DFEC", "5B3F86"), "kho_doi": ("EDEDED", "595959")}
TEN_DC = {"khop": "Khớp", "do": "ĐỎ", "vang": "VÀNG", "cam": "CAM", "tim": "TÍM", "kho_doi": "Khó đòi"}


def _cot_doi_chieu(ws) -> None:
    c = len(th.COT) + 1
    o = ws.cell(row=th.DONG_TIEU_DE, column=c, value=th.COT_DOI_CHIEU[0])
    goc = ws.cell(row=th.DONG_TIEU_DE, column=len(th.COT))
    o._style = copy(goc._style)
    ws.column_dimensions[get_column_letter(c)].width = th.COT_DOI_CHIEU[1]


def chu_doi_chieu(dc: dict) -> str:
    if not dc:
        return ""
    if dc["mau"] == "khop":
        return "Khớp"
    if dc["mau"] == "kho_doi":
        return f"Khó đòi · {dc['goi_y']}"
    return f"{TEN_DC[dc['mau']]} · lệch {dc['lech']:,.0f} · {dc['goi_y']}"


def _o_doi_chieu(ws, hang: int, dc: dict, kieu_vien) -> None:
    o = ws.cell(row=hang, column=len(th.COT) + 1, value=chu_doi_chieu(dc) or None)
    o.border = kieu_vien
    o.alignment = Alignment(vertical="top", wrap_text=True)
    if dc:
        nen, chu = MAU_DC[dc["mau"]]
        o.fill = PatternFill("solid", fgColor=nen)
        o.font = Font(name=th.PHONG_NOI_DUNG, size=9, bold=dc["mau"] not in ("khop", "kho_doi"), color=chu)


def _ghi_dong_con(ws, d: DongBaoCao, hang: int, kieu_vien, n_cot: Optional[int] = None) -> int:
    """Các dòng con (đang gập) dưới một khách: mỗi job SMS còn nợ một dòng — ai phụ trách,
    còn bao nhiêu, nằm ở nhóm tuổi nào — và một dòng cân đối. Cộng các dòng con theo từng cột
    đúng bằng dòng khách (nhóm tuổi phân bổ: nợ cũ và job cũ nhận nhóm già trước)."""
    n_cot = n_cot or len(th.COT)

    def _nhan(g):
        if g.get("so_hd"):                    # v0.5: job từ Misa — số hoá đơn, ngày, trạng thái
            return f"{NHAN_CON}{g['job']} · HĐ {g['so_hd']} ngày {g['ngay']:%d/%m/%Y}"
        return (f"{NHAN_CON}{g['job']}" + (f" · {g['nguon_ngay']} {g['ngay']:%d/%m/%Y}" if g["ngay"] else "")
                + (" · theo bảng FIN" if g.get("tham_chieu") else ""))
    con = [(_nhan(g), g["tien"], g["sales"], g["job"], g.get("nhom") or [0.0] * 6,
            g.get("trang_thai") or ("Nợ cũ – bảng FIN" if g.get("tham_chieu") else ""))
           for g in d.dong_con]
    if d.can_doi:
        nhan = ("Nợ cũ / chưa có job trên SMS" if d.can_doi > 0 else
                "SMS ghi dư — đã thu, chưa gạch paid (xem Review)")
        con.append((f"{NHAN_CON}{nhan}", d.can_doi, "", "", d.can_doi_nhom or [0.0] * 6, ""))
    for ten, tien, sale, job, nhom, tt in con:
        o_tuoi = [(6 + i, v if abs(v) > 0.4 else None) for i, v in enumerate(nhom)]
        them = [(len(th.COT) + 1, tt or None)] if n_cot > len(th.COT) else []
        for cot, v in ((3, ten), (5, tien), (12, sale), (13, job), *o_tuoi, *them):
            o = ws.cell(row=hang, column=cot, value=v)
            o.font = Font(name=th.PHONG_NOI_DUNG, size=9, color="595959",
                          italic=not job)
            if 5 <= cot <= 11:
                o.number_format = DINH_DANG_TIEN
                o.alignment = Alignment(horizontal="right")
        for cot in range(1, n_cot + 1):
            o = ws.cell(row=hang, column=cot)
            o.fill = PatternFill("solid", fgColor="F7F5FB")
            o.border = kieu_vien
        ws.row_dimensions[hang].outlineLevel = 1
        ws.row_dimensions[hang].hidden = True
        hang += 1
    return hang


def _ghi_sheet_xem_lai(wb, dong: List[DongBaoCao], ngay_chot: dt.date, tong_misa,
                      doi_chieu: dict) -> None:
    ws = wb.create_sheet(th.SHEET_REVIEW)
    ws["A1"] = f"Đối chiếu và những dòng tool thấy không chắc — chốt số {ngay_chot:%d/%m/%Y}"
    ws["A1"].font = Font(name=th.PHONG_TIEU_DE, size=12, bold=True, color=th.TIM)

    tong_bang = sum(d.total for d in dong)
    tong_tuoi_no = doi_chieu.get("tong_tuoi_no")
    lech_khach = doi_chieu.get("khach_lech") or []

    hang = 3
    def _dong_doi_chieu(nhan, gia_tri, ghi_chu="", mau="000000"):
        nonlocal hang
        o = ws.cell(row=hang, column=1, value=nhan)
        o.font = Font(name=th.PHONG_NOI_DUNG, size=10, bold=True, color=mau)
        c = ws.cell(row=hang, column=2, value=gia_tri)
        c.number_format = DINH_DANG_TIEN
        c.font = Font(name=th.PHONG_NOI_DUNG, size=10, bold=True, color=mau)
        o2 = ws.cell(row=hang, column=4, value=ghi_chu)
        o2.font = Font(name=th.PHONG_NOI_DUNG, size=10, color=mau)
        o2.alignment = Alignment(wrap_text=True, vertical="top")
        hang += 1

    nguon = doi_chieu.get("nguon") or []
    for nhan, mo_ta in nguon:
        o = ws.cell(row=hang, column=1, value=nhan)
        o.font = Font(name=th.PHONG_NOI_DUNG, size=9, italic=True)
        o2 = ws.cell(row=hang, column=4, value=mo_ta)
        o2.font = Font(name=th.PHONG_NOI_DUNG, size=9, italic=True)
        hang += 1
    if nguon:
        hang += 1
    for canh_bao in doi_chieu.get("canh_bao") or []:
        o = ws.cell(row=hang, column=1, value=f"CẢNH BÁO: {canh_bao}")
        o.font = Font(name=th.PHONG_NOI_DUNG, size=10, bold=True, color="C00000")
        hang += 1
    if doi_chieu.get("canh_bao"):
        hang += 1

    _dong_doi_chieu("Tổng bảng báo cáo này", tong_bang,
                    "Phải bằng dòng ngay dưới. Lệch là tool có lỗi.")
    if tong_misa is not None:
        khop = abs(tong_bang - tong_misa) < 2
        _dong_doi_chieu("Tổng file Tổng hợp công nợ (dư Nợ)", tong_misa,
                        "KHỚP" if khop else "KHÔNG KHỚP — báo lại người làm tool",
                        "000000" if khop else "C00000")
    if tong_tuoi_no is not None:
        chenh = tong_tuoi_no - (tong_misa if tong_misa is not None else tong_bang)
        _dong_doi_chieu("Tổng file Phân tích tuổi nợ", tong_tuoi_no,
                        "Con số này KHÔNG dùng để đối chiếu tổng báo cáo. "
                        "Nó thường cao hơn vì còn hoá đơn đã thu nhưng chưa gắn "
                        "chứng từ thanh toán.")
        if abs(chenh) >= 2:
            _dong_doi_chieu("Chênh lệch giữa hai file Misa", chenh,
                            "Do các khách liệt kê ngay dưới. Tiền luôn lấy theo "
                            "file Tổng hợp.", "B26B00")
            hang += 1
            o = ws.cell(row=hang, column=1, value="Khách làm lệch hai file")
            o.font = Font(name=th.PHONG_TIEU_DE, size=10, bold=True, color=th.CAM)
            hang += 1
            for ten, v_th, v_tn in lech_khach:
                ws.cell(row=hang, column=1, value=ten).font = Font(name=th.PHONG_NOI_DUNG, size=10)
                for cot, v in ((2, v_th), (3, v_tn)):
                    c = ws.cell(row=hang, column=cot, value=v)
                    c.number_format = DINH_DANG_TIEN
                    c.font = Font(name=th.PHONG_NOI_DUNG, size=10)
                ws.cell(row=hang, column=4,
                        value="Tổng hợp | Tuổi nợ").font = Font(name=th.PHONG_NOI_DUNG, size=9,
                                                                italic=True)
                hang += 1
    if doi_chieu.get("che_do") == "hoa-don":
        hang += 1
        tong_nguon: Dict[str, float] = {}
        for d in dong:
            for k, v in d.phan_nguon.items():
                tong_nguon[k] = tong_nguon.get(k, 0) + v
        _dong_doi_chieu("Tuổi nợ tính từ đâu", tong_bang, "Bằng các dòng dưới cộng lại.")
        for k, nhan, giai_thich, mau in (
            ("hoa-don", "  · hoá đơn trong file Chi tiết 131", "Ngày hoá đơn + credit term. Phiếu thu ghi "
             "số hoá đơn thì trừ đúng hoá đơn đó, còn lại trừ hoá đơn cũ nhất trước.", "000000"),
            ("kho-doi", "  · khó đòi", "Luôn ở nhóm 120+ (bảng ghép khách hoặc lý do ghi 'khó đòi').",
             "000000"),
            ("ngoai-le", "  · ngoại lệ kế toán khai", "", "000000"),
            ("uoc-tinh", "  · Tổng hợp nhiều hơn Chi tiết", "Phần chênh tạm coi là nợ mới (Current). "
             "Thường do hai file xuất khác thời điểm — xem bảng đối chiếu bên dưới.", "C00000"),
        ):
            if tong_nguon.get(k, 0) >= 1:
                _dong_doi_chieu(nhan, tong_nguon[k], giai_thich, mau)
    elif doi_chieu.get("che_do") == "sms":
        hang += 1
        tong_nguon: Dict[str, float] = {}
        for d in dong:
            for k, v in d.phan_nguon.items():
                tong_nguon[k] = tong_nguon.get(k, 0) + v
        _dong_doi_chieu("Tuổi nợ tính từ đâu", tong_bang, "Bằng các dòng dưới cộng lại.")
        for k, nhan, giai_thich, mau in (
            ("sms", "  · ngày job trên SMS", "ETD/ETA của từng job + credit term.", "000000"),
            ("so-ghi", "  · bản final/sổ ghi tuần trước", "Nợ cũ không có trong SMS: giữ nhóm tuổi "
             "kế toán trưởng đã xếp, cộng thêm số ngày đã trôi qua.", "000000"),
            ("kho-doi", "  · khó đòi (FIN khai)", "Luôn ở nhóm 120+.", "000000"),
            ("ngoai-le", "  · ngoại lệ kế toán khai", "", "000000"),
            ("uoc-tinh", "  · chưa rõ tuổi, đang ước tính", "Không có trong SMS lẫn bản final. "
             "Dòng tô đỏ ở sheet Receivable — kế toán xếp lại nhóm trên bản final.", "C00000"),
        ):
            if tong_nguon.get(k, 0) >= 1:
                _dong_doi_chieu(nhan, tong_nguon[k], giai_thich, mau)
    else:
        hang += 1
        _ghi_no_trong_han(ws, dong, _dong_doi_chieu)
        hang = ws.max_row + 1
    hang += 1
    _bang_khach_xem_lai(ws, dong, doi_chieu, hang)


def _ghi_no_trong_han(ws, dong, _dong_doi_chieu) -> None:
    # Nợ trong hạn: vì sao báo cáo khác cột "Nợ trước hạn" của file Tuổi nợ
    cur_tu_misa = sum(d.nhom[0] for d in dong if d.nguon == "misa")
    cur_khac = sum(d.nhom[0] for d in dong if d.nguon != "misa")
    _dong_doi_chieu("Nợ trong hạn trên báo cáo", cur_tu_misa + cur_khac,
                    "Bằng hai dòng dưới cộng lại.")
    _dong_doi_chieu("  · lấy thẳng từ file Tuổi nợ", cur_tu_misa,
                    "Khớp với cột Nợ trước hạn của file Tuổi nợ, phần các khách "
                    "mà hai file Misa khớp nhau.")
    if cur_khac >= 1:
        ten_khac = ", ".join(d.ten for d in dong if d.nguon != "misa" and d.nhom[0] >= 1)
        _dong_doi_chieu("  · chưa xác định được tuổi, tạm để ở đây", cur_khac,
                        f"Hoá đơn không ghi hạn thanh toán, hoặc hai file Misa lệch "
                        f"nhau, nên chưa tính được tuổi. Gồm: {ten_khac}. "
                        "Các dòng này được tô màu ở sheet Receivable.", "B26B00")


def _bang_khach_xem_lai(ws, dong, doi_chieu, hang) -> None:
    tieu_de = ["Khách hàng", "Total", "Nguồn số", "Vì sao cần xem lại"]
    for i, t in enumerate(tieu_de, start=1):
        o = ws.cell(row=hang, column=i, value=t)
        o.fill = PatternFill("solid", fgColor=th.CAM if i != 2 else th.TIM)
        o.font = Font(name=th.PHONG_TIEU_DE, size=10, bold=True, color=th.TRANG)
    hang += 1
    thieu_term = sum(1 for d in dong if any("Credit Term" in x for x in d.thieu_thong_tin))
    thieu_sale = sum(1 for d in dong if any("Salesman" in x for x in d.thieu_thong_tin))
    # vấn đề về số liệu xếp trước, rồi mới tới thiếu thông tin do người điền
    for d in sorted(dong, key=lambda x: (not x.cho_xem_lai, -x.total)):
        ly_do = d.cho_xem_lai + d.thieu_thong_tin
        if not ly_do:
            continue
        ws.cell(row=hang, column=1, value=d.ten).font = Font(name=th.PHONG_NOI_DUNG, size=10)
        c = ws.cell(row=hang, column=2, value=d.total)
        c.number_format = DINH_DANG_TIEN
        c.font = Font(name=th.PHONG_NOI_DUNG, size=10)
        ws.cell(row=hang, column=3, value=d.nguon).font = Font(name=th.PHONG_NOI_DUNG, size=10)
        o = ws.cell(row=hang, column=4, value=" | ".join(ly_do))
        o.font = Font(name=th.PHONG_NOI_DUNG, size=10)
        o.alignment = Alignment(wrap_text=True, vertical="top")
        hang += 1
    da_bo = doi_chieu.get("da_bo") or []
    if da_bo:
        hang += 1
        o = ws.cell(row=hang, column=1, value="Khách bị loại khỏi bảng theo khai báo của kế toán")
        o.font = Font(name=th.PHONG_TIEU_DE, size=10, bold=True, color=th.CAM)
        hang += 1
        for ten, tien, ly_do in da_bo:
            ws.cell(row=hang, column=1, value=ten).font = Font(name=th.PHONG_NOI_DUNG, size=10)
            c = ws.cell(row=hang, column=2, value=tien)
            c.number_format = DINH_DANG_TIEN
            c.font = Font(name=th.PHONG_NOI_DUNG, size=10)
            o = ws.cell(row=hang, column=4, value=ly_do)
            o.font = Font(name=th.PHONG_NOI_DUNG, size=10)
            o.alignment = Alignment(wrap_text=True, vertical="top")
            hang += 1

    for khoa, tieu, cot_gt, ghi in (
        ("sms_cat", "SMS cần gạch paid — SMS ghi nợ nhiều hơn Misa, tool đã bỏ phần này",
         ("Khách hàng", "Số tiền bỏ", "Job", "Ngày ETD/ETA"),
         "FIN kiểm tra và ấn paid trên SMS cho các job này."),
        ("sms_khong_misa", "SMS còn nợ nhưng Misa không còn — không đưa vào báo cáo",
         ("Khách trên SMS", "Số tiền SMS", "Job", "Vì sao"),
         "Nhiều khả năng đã thu tiền mà SMS chưa gạch paid."),
        ("sms_chua_ghep", "Khách SMS chưa có trong bảng ghép — số này CHƯA vào báo cáo",
         ("Mã | tên trên SMS", "Số tiền SMS", "Số dòng", ""),
         "Bổ sung vào file ghép khách (nút 2_Kiem_tra_ghep_khach sẽ liệt kê sẵn)."),
    ):
        ds = doi_chieu.get(khoa) or []
        if not ds:
            continue
        hang += 1
        o = ws.cell(row=hang, column=1, value=f"{tieu} — {ghi}")
        o.font = Font(name=th.PHONG_TIEU_DE, size=10, bold=True, color=th.CAM)
        hang += 1
        for i, t in enumerate(cot_gt, start=1):
            ws.cell(row=hang, column=i, value=t).font = Font(name=th.PHONG_NOI_DUNG, size=9, bold=True)
        hang += 1
        for dong_ds in ds:
            if khoa == "sms_cat":
                vals = (dong_ds[0], dong_ds[2], dong_ds[1],
                        dong_ds[3].strftime("%d/%m/%Y") if dong_ds[3] else "")
            elif khoa == "sms_khong_misa":
                vals = (f"{dong_ds[0]} | {dong_ds[1]}", dong_ds[3], dong_ds[2], dong_ds[4])
            else:
                vals = (f"{dong_ds[0]} | {dong_ds[1]}", dong_ds[2], dong_ds[3], "")
            for i, v in enumerate(vals, start=1):
                c = ws.cell(row=hang, column=i, value=v)
                c.font = Font(name=th.PHONG_NOI_DUNG, size=10)
                if i == 2:
                    c.number_format = DINH_DANG_TIEN
            hang += 1

    for tieu, cot_gt, ds in doi_chieu.get("bang_them") or []:
        if not ds:
            continue
        hang += 1
        o = ws.cell(row=hang, column=1, value=tieu)
        o.font = Font(name=th.PHONG_TIEU_DE, size=10, bold=True, color=th.CAM)
        hang += 1
        for i, t in enumerate(cot_gt, start=1):
            ws.cell(row=hang, column=i, value=t).font = Font(name=th.PHONG_NOI_DUNG, size=9, bold=True)
        hang += 1
        for vals in ds:
            for i, v in enumerate(vals, start=1):
                c = ws.cell(row=hang, column=i, value=v)
                c.font = Font(name=th.PHONG_NOI_DUNG, size=10)
                if isinstance(v, (int, float)) and i > 1:
                    c.number_format = DINH_DANG_TIEN
            hang += 1

    hang += 1
    thieu_term = sum(1 for d in dong if any("Credit Term" in x for x in d.thieu_thong_tin))
    thieu_sale = sum(1 for d in dong if any("Salesman" in x for x in d.thieu_thong_tin))
    o = ws.cell(row=hang, column=1,
                value=f"Thiếu thông tin do người điền: {thieu_term} khách chưa có Credit Term, "
                      f"{thieu_sale} khách chưa có Salesman. Đây không phải lỗi số liệu.")
    o.font = Font(name=th.PHONG_NOI_DUNG, size=10, italic=True)
    for cot, rong in zip("ABCD", (46, 16, 16, 80)):
        ws.column_dimensions[cot].width = rong


# --- nạp trạng thái ban đầu từ một báo cáo đã có -------------------------

ANH_XA_NHOM = {
    "current": 0, "1 - 30": 1, "1-30": 1, "31 - 60": 2, "31-60": 2,
    "61 - 90": 3, "61-90": 3, "60+": 3, "91 - 120": 4, "91 -120": 4, "91-120": 4,
    "120+": 5, "trên 120": 5,
    "not-yet-due": 0, "not yet due": 0, "over 120": 5,
}


def _tim_bang(wb):
    """Tìm sheet và dòng tiêu đề của bảng công nợ: sheet Receivable của mẫu chuẩn,
    hoặc sheet bất kỳ (VD 'data' trong bản final kế toán trưởng)."""
    uu_tien = [n for n in ("Receivables", "Receivable") if n in wb.sheetnames]
    thu_tu = uu_tien + [n for n in wb.sheetnames if n not in uu_tien]
    for ten in thu_tu:
        ws = wb[ten]
        for r, hang in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
            gia_tri = [str(v or "").strip().lower() for v in hang]
            if "code" in gia_tri and any(v in ("total", "total receivables") for v in gia_tri):
                return ws, r
    return None, None


def co_tab_review(duong_dan: Path) -> bool:
    """Báo cáo do tool sinh luôn có tab Review; bản kế toán trưởng tự làm thì không."""
    wb = openpyxl.load_workbook(duong_dan, read_only=True)
    try:
        return th.SHEET_REVIEW in wb.sheetnames or "Cần xem lại" in wb.sheetnames
    finally:
        wb.close()


def ngay_chot_trong_file(duong_dan: Path) -> Optional[dt.date]:
    """Ngày chốt của một báo cáo tuần. Ưu tiên 'At 03 Oct 2026' (mẫu chuẩn, mọi sheet của tool);
    không có thì lấy ngày sau trong 'Reporting period dd/mm/yyyy - dd/mm/yyyy' (bản final cũ).
    Phải ưu tiên 'At': sheet Summary ghi cả hai trên một dòng, mà ngày cuối kỳ (02/10) KHÁC
    ngày chốt (03/10) — đọc nhầm là tạo ra một "tuần trước" giả cách một ngày."""
    wb = openpyxl.load_workbook(duong_dan, read_only=True, data_only=True)
    ky = None
    try:
        for ws in wb.worksheets:
            for hang in ws.iter_rows(min_row=1, max_row=12, values_only=True):
                chu = " ".join(str(v) for v in hang if v is not None)
                m = re.search(r"\bAt (\d{1,2}) (\w{3}) (\d{4})", chu)
                if m and m.group(2) in THANG_EN:
                    return dt.date(int(m.group(3)), THANG_EN.index(m.group(2)) + 1, int(m.group(1)))
                if ky is None and ("period" in chu.lower() or "kỳ" in chu.lower()):
                    ngay = re.findall(r"(\d{1,2})/(\d{1,2})/(\d{4})", chu)
                    if ngay:
                        d, mth, y = ngay[-1]
                        ky = dt.date(int(y), int(mth), int(d))
    finally:
        wb.close()
    return ky


def doc_bao_cao(duong_dan: Path) -> Dict[str, sg.HoSoKhach]:
    """Đọc sheet Receivable của một báo cáo tuần (kể cả bản kế toán làm tay,
    kể cả bản cũ chỉ có 4 cột tuổi nợ) để nạp vào sổ ghi."""
    wb = openpyxl.load_workbook(duong_dan, data_only=True)
    ws, dong_tieu_de = _tim_bang(wb)
    if ws is None:
        raise ValueError(f"{Path(duong_dan).name}: không tìm thấy bảng công nợ (dòng tiêu đề có Code và Total)")

    cot: Dict[str, int] = {}
    nhom_cot: Dict[int, int] = {}
    for c in ws[dong_tieu_de]:
        ten = str(c.value or "").strip().lower()
        if not ten:
            continue
        ten_nhom = re.sub(r"\s+", " ", ten).replace(" days", "").strip()
        if ten_nhom in ANH_XA_NHOM:
            nhom_cot[c.column] = ANH_XA_NHOM[ten_nhom]
        elif "customer" in ten and "ten" not in cot:
            cot["ten"] = c.column
        elif ten == "code":
            cot["ma"] = c.column
        elif "credit term" in ten:
            cot["term"] = c.column
        elif ten in ("total", "total receivables"):
            cot["total"] = c.column
        elif "salesman" in ten and "salesman" not in cot:
            cot["salesman"] = c.column
        elif ten in ("job", "job no", "số job", "job number"):
            cot["job"] = c.column
        elif "ghi chú" in ten or ten in ("comment", "notes", "note"):
            cot["ghi_chu"] = c.column
        elif "lý do" in ten or "reason" in ten:
            cot["ly_do"] = c.column

    ra: Dict[str, sg.HoSoKhach] = {}
    for r in range(dong_tieu_de + 1, ws.max_row + 1):
        ten = ws.cell(row=r, column=cot["ten"]).value if "ten" in cot else None
        ma = ws.cell(row=r, column=cot.get("ma", 2)).value
        if not ten or str(ma or "").strip().upper().startswith(("TỔNG", "TONG")):
            continue
        if str(ten).strip().startswith("└"):
            continue                 # dòng con (job) dưới một khách, không phải khách
        if chuan_hoa_ten(ten) in ("TOTAL", "TỔNG", "TỔNG CỘNG", "GRAND TOTAL"):
            continue
        total = ws.cell(row=r, column=cot["total"]).value
        if not isinstance(total, (int, float)):
            continue
        nhom = [0.0] * 6
        for c, i in nhom_cot.items():
            v = ws.cell(row=r, column=c).value
            if isinstance(v, (int, float)):
                nhom[i] += float(v)
        lay = lambda k: str(ws.cell(row=r, column=cot[k]).value or "").strip() if k in cot else ""
        tc = chuan_hoa_ten(ten)
        job_gc, ghi_chu = tach_ghi_chu(lay("ghi_chu"))
        job = re.sub(r"(?i)^\s*job\s*:\s*", "", lay("job")) or job_gc
        ra[tc] = sg.HoSoKhach(
            ma=str(ma or "").strip(), ten=str(ten).strip(), credit_term=lay("term"),
            salesman=th.ten_sales(lay("salesman")), ghi_chu=ghi_chu,
            # câu tool tự viết không phải lời của Sales: đọc lại thì coi như trống
            ly_do="" if lay("ly_do").startswith(th.NHAN_TU_DONG) else lay("ly_do"),
            tong=round(float(total), 2), nhom=[round(v, 2) for v in nhom], nguon="ban-tay",
            job=job,
        )
    return ra


def nap_vao_so_ghi(so: sg.SoGhi, duong_dan: Path, ngay_chot: dt.date) -> int:
    khach = doc_bao_cao(duong_dan)
    for hs in khach.values():
        hs.lo = sg.lo_tu_nhom(hs.nhom, ngay_chot, uoc_tinh=True)
    so.ghi_trang(ngay_chot, khach)
    return len(khach)

# -*- coding: utf-8 -*-
"""v0.5 (WR-06) — job chưa thanh toán lấy thẳng từ Misa.

Hai file FIN xuất từ đầu năm tới nay (FIN 09/10: xuất từ 2022 sẽ sai vì hoá đơn các năm trước
thanh toán khác loại tiền không gắn được phiếu thu, vẫn hiện "Chưa thanh toán"):
  · Bán hàng            — mỗi hoá đơn một dòng, cột "TT thanh toán"
  · Sổ chi tiết TK131   — mỗi dòng hàng của hoá đơn, cột "Mã đối tượng THCP" = số job

Luật FIN 09/10: chỉ "Đã thanh toán" mới tính là đã trả. "Chưa thanh toán" và "Thanh toán một
phần" đều liệt kê là job chưa thanh toán (số tiền cả hoá đơn). Số tiền theo tỷ giá xuất hoá đơn
nên khớp sổ kế toán, không lấy số tiền job từ SMS nữa.

Khách nhận dạng theo TÊN (hai file không có mã khách) → ghép sang mã Misa bằng tên trong
Tổng hợp 131. Tên không ghép được nằm ở danh sách khong_ghep (màu tím trong Review).
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl

from .misa import LoiDocFile, chuan_hoa_ten

DA_TT = "đã thanh toán"


def _o(x) -> str:
    return str(x).strip() if x is not None else ""


def _ngay(x) -> Optional[dt.date]:
    if isinstance(x, dt.datetime):
        return x.date()
    if isinstance(x, dt.date):
        return x
    return None


def _so(x) -> float:
    return float(x) if isinstance(x, (int, float)) else 0.0


def _tim_tieu_de(ws, can: Tuple[str, ...], toi_da=10):
    for r, hang in enumerate(ws.iter_rows(min_row=1, max_row=toi_da, values_only=True), start=1):
        ten = [_o(v).lower() for v in hang]
        if all(any(c in t for t in ten) for c in can):
            return r, ten
    return None, None


def _cot(ten: List[str], *tu) -> Optional[int]:
    for i, t in enumerate(ten):
        if all(x in t for x in tu):
            return i
    return None


@dataclass
class HoaDon:
    so_hd: str
    ngay: dt.date
    ten: str
    ten_chuan: str
    tien: float
    trang_thai: str                      # nguyên chữ Misa
    jobs: Dict[str, float] = field(default_factory=dict)   # job -> tiền (đã quy về tổng hoá đơn)
    thu_ngay: Optional[dt.date] = None   # ngày phiếu thu gần nhất nhắc tới hoá đơn này (Sổ chi tiết)

    def chua_tt_tai(self, den: dt.date) -> bool:
        """Chưa thanh toán TẠI ngày cuối kỳ: Misa ghi chưa/một phần, hoặc ghi 'Đã thanh toán' nhưng
        phiếu thu nhắc tới hoá đơn này nằm SAU ngày cuối kỳ (file xuất muộn hơn kỳ báo cáo)."""
        return self.chua_tt or (self.thu_ngay is not None and self.thu_ngay > den)

    @property
    def chua_tt(self) -> bool:
        return self.trang_thai.strip().lower() != DA_TT

    @property
    def mot_phan(self) -> bool:
        return "một phần" in self.trang_thai.lower()


@dataclass
class JobMisa:
    hoa_don: List[HoaDon]
    tu_ngay: dt.date                     # ngày hoá đơn sớm nhất trong file: trước mốc này là "nợ cũ"
    den_ngay: dt.date
    co_job: bool                         # có file Sổ chi tiết (biết job) hay chỉ có Bán hàng
    theo_ma: Dict[str, List[HoaDon]] = field(default_factory=dict)
    khong_ghep: List[HoaDon] = field(default_factory=list)   # hoá đơn không tìm ra khách Misa

    def khong_ghep_tai(self, den: dt.date) -> List[Tuple[str, int, float]]:
        """(tên, số hoá đơn chưa TT, tiền) của các khách không ghép được, tính tại ngày cuối kỳ."""
        gom: Dict[str, list] = {}
        for h in self.khong_ghep:
            if h.ngay <= den and h.chua_tt_tai(den):
                g = gom.setdefault(h.ten, [h.ten, 0, 0.0])
                g[1] += 1
                g[2] = round(g[2] + h.tien, 2)
        return [tuple(g) for g in sorted(gom.values(), key=lambda g: -g[2])]

    def chua_tt_cua(self, ma: str, den: dt.date) -> List[HoaDon]:
        return [h for h in self.theo_ma.get(ma, []) if h.ngay <= den and h.chua_tt_tai(den)]


def doc_ban_hang(path) -> List[HoaDon]:
    wb = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    r0, ten = _tim_tieu_de(ws, ("số hóa đơn", "tt thanh toán"))
    if r0 is None:
        wb.close()
        raise LoiDocFile(f"{Path(path).name}: không thấy dòng tiêu đề có 'Số hóa đơn' và 'TT thanh toán'")
    i_ngay, i_hd = _cot(ten, "ngày hạch toán"), _cot(ten, "số hóa đơn")
    i_kh, i_tien, i_tt = _cot(ten, "khách hàng"), _cot(ten, "tổng tiền"), _cot(ten, "tt thanh toán")
    ds = []
    for hang in ws.iter_rows(min_row=r0 + 1, values_only=True):
        n = _ngay(hang[i_ngay]) if i_ngay is not None else None
        so_hd = _o(hang[i_hd])
        if not n or not so_hd:
            continue
        ds.append(HoaDon(so_hd=so_hd, ngay=n, ten=_o(hang[i_kh]), ten_chuan=chuan_hoa_ten(hang[i_kh]),
                         tien=round(_so(hang[i_tien]), 2), trang_thai=_o(hang[i_tt])))
    wb.close()
    return ds


def doc_so_chi_tiet(path):
    """Trả ((số HĐ, tên chuẩn) -> {job: tiền}, (số HĐ int, tên chuẩn) -> ngày thu gần nhất).
    Job: dòng ghi nợ bán hàng có số hoá đơn. Ngày thu: phiếu thu có diễn giải nhắc số hoá đơn."""
    from .chitiet import so_hd_nhac_toi
    wb = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    r0, ten = _tim_tieu_de(ws, ("số hóa đơn", "phát sinh nợ"))
    if r0 is None:
        wb.close()
        raise LoiDocFile(f"{Path(path).name}: không thấy dòng tiêu đề có 'Số hóa đơn' và 'Phát sinh Nợ'")
    i_hd, i_kh = _cot(ten, "số hóa đơn"), _cot(ten, "tên đối tượng")
    i_no, i_co = _cot(ten, "phát sinh nợ"), _cot(ten, "phát sinh có")
    i_job = _cot(ten, "mã đối tượng thcp")
    if i_job is None:
        i_job = _cot(ten, "mã thống kê")
    i_dg = _cot(ten, "diễn giải")
    kq: Dict[Tuple[str, str], Dict[str, float]] = {}
    thu: Dict[Tuple[int, str], dt.date] = {}
    for hang in ws.iter_rows(min_row=r0 + 1, values_only=True):
        n = _ngay(hang[0])
        so_hd = _o(hang[i_hd]) if i_hd is not None else ""
        if n and not so_hd and _so(hang[i_co]) > 0.5 and i_dg is not None:
            tc = chuan_hoa_ten(hang[i_kh])
            for x in so_hd_nhac_toi(_o(hang[i_dg])):
                if thu.get((x, tc)) is None or n > thu[(x, tc)]:
                    thu[(x, tc)] = n
            continue
        if not so_hd or not n:
            continue
        tien = _so(hang[i_no]) - _so(hang[i_co])
        if abs(tien) < 0.5:
            continue
        job = _o(hang[i_job]) if i_job is not None else ""
        g = kq.setdefault((so_hd, chuan_hoa_ten(hang[i_kh])), {})
        g[job] = round(g.get(job, 0.0) + tien, 2)
    wb.close()
    return kq, thu


def ghep_job(hoa_don: List[HoaDon], so_ct: Optional[Dict[Tuple[str, str], Dict[str, float]]]) -> None:
    """Gắn job cho từng hoá đơn. Tiền từng job chia theo tỷ lệ trên Sổ chi tiết để cộng lại
    đúng bằng tổng hoá đơn trên Bán hàng."""
    for h in hoa_don:
        g = (so_ct or {}).get((h.so_hd, h.ten_chuan)) or {}
        g = {j: t for j, t in g.items() if t > 0.5}
        tong = sum(g.values())
        if not g or tong <= 0:
            h.jobs = {"": h.tien}
            continue
        h.jobs = {j: round(h.tien * t / tong, 2) for j, t in g.items()}
        lech = round(h.tien - sum(h.jobs.values()), 2)
        if lech:
            j0 = max(h.jobs, key=h.jobs.get)
            h.jobs[j0] = round(h.jobs[j0] + lech, 2)


def ghep_khach(hoa_don: List[HoaDon], ds_misa: List[Tuple[str, str]],
               bo_sung: Optional[Dict[str, str]] = None) -> Tuple[Dict[str, List[HoaDon]], List[HoaDon]]:
    """ds_misa: (mã, tên) từ Tổng hợp 131. bo_sung: tên chuẩn -> mã (FIN khai tay).
    Trả (mã -> hoá đơn, hoá đơn không ghép được)."""
    from .ghep import diem_giong
    theo_ten = {chuan_hoa_ten(t): m for m, t in ds_misa}
    theo_ten.update(bo_sung or {})
    nho: Dict[str, Optional[str]] = {}
    theo_ma: Dict[str, List[HoaDon]] = {}
    sot: List[HoaDon] = []
    for h in hoa_don:
        if h.ten_chuan not in nho:
            ma = theo_ten.get(h.ten_chuan)
            if ma is None:
                diem = sorted(((diem_giong(h.ten, t), m) for m, t in ds_misa), reverse=True)
                if diem and diem[0][0] >= 0.92 and (len(diem) == 1 or diem[1][0] < 0.92):
                    ma = diem[0][1]
            nho[h.ten_chuan] = ma
        ma = nho[h.ten_chuan]
        if ma is None:
            sot.append(h)
            continue
        theo_ma.setdefault(ma, []).append(h)
    return theo_ma, sot


def doc_job_misa(file_ban_hang, file_so_chi_tiet, ds_misa: List[Tuple[str, str]],
                 bo_sung: Optional[Dict[str, str]] = None) -> Optional[JobMisa]:
    if not file_ban_hang:
        return None
    hd = doc_ban_hang(file_ban_hang)
    if not hd:
        return None
    so_ct, thu = doc_so_chi_tiet(file_so_chi_tiet) if file_so_chi_tiet else (None, {})
    ghep_job(hd, so_ct)
    for h in hd:
        if h.so_hd.isdigit():
            h.thu_ngay = thu.get((int(h.so_hd), h.ten_chuan))
    theo_ma, sot = ghep_khach(hd, ds_misa, bo_sung)
    return JobMisa(hoa_don=hd, tu_ngay=min(h.ngay for h in hd), den_ngay=max(h.ngay for h in hd),
                   co_job=so_ct is not None, theo_ma=theo_ma, khong_ghep=sot)

# -*- coding: utf-8 -*-
"""Sinh lại mẫu chuẩn: một sheet Receivable trắng dữ liệu, kẻ màu thương hiệu.

Chạy lại khi cần đổi bố cục mẫu:
    python tools/make_template.py
Sheet của các báo cáo khác (Payable, General...) thêm vào sau, mỗi thằng một sheet.
"""
import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from wr import thuonghieu as th

DICH = pathlib.Path(__file__).resolve().parents[1] / "template" / "Bang_cong_no_MAU_CHUAN.xlsx"


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Receivable"

    ws["A1"] = "TRUSTANA VIỆT NAM"
    ws["A1"].font = Font(name=th.PHONG_TIEU_DE, size=14, bold=True, color=th.TIM)
    ws["A2"] = "Debtors and Creditors Ageing Reports"
    ws["A2"].font = Font(name=th.PHONG_TIEU_DE, size=11, bold=True, color=th.CAM)
    ws["A3"] = ""  # tool điền: At 12 Sep 2026
    ws["A3"].font = Font(name=th.PHONG_NOI_DUNG, size=10, italic=True)
    ws["A5"] = "Accounts receivable"
    ws["A5"].font = Font(name=th.PHONG_TIEU_DE, size=11, bold=True, color=th.TIM)

    vien = Side(style="thin", color="BFBFBF")
    for i, (ten_cot, loai, rong) in enumerate(th.COT, start=1):
        o = ws.cell(row=th.DONG_TIEU_DE, column=i, value=ten_cot)
        o.fill = PatternFill("solid", fgColor=th.CAM if loai == "ten" else th.TIM)
        o.font = Font(name=th.PHONG_TIEU_DE, size=10, bold=True, color=th.TRANG)
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        o.border = Border(left=vien, right=vien, top=vien, bottom=vien)
        ws.column_dimensions[get_column_letter(i)].width = rong
    ws.row_dimensions[th.DONG_TIEU_DE].height = 30
    ws.freeze_panes = ws.cell(row=th.DONG_DAU_DU_LIEU, column=4)

    DICH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DICH)
    print("Đã ghi mẫu chuẩn:", DICH)


if __name__ == "__main__":
    main()

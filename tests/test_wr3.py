# -*- coding: utf-8 -*-
"""Phép thử WR-03: nguồn SMS, bảng ghép khách, bản final làm trí nhớ nợ cũ.

Toàn bộ dữ liệu ở đây là GIẢ, sinh ngay trong phép thử — không cần dữ liệu thật.
"""
import datetime as dt
import os
import shutil
import sys
import time
from pathlib import Path

import openpyxl
import pytest

GOC = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(GOC))
sys.path.insert(0, str(GOC / "tests"))

from wr import baocao, ghep, misa  # noqa: E402
from wr import soghi as sg  # noqa: E402
from wr import tinhtoan  # noqa: E402
import runner  # noqa: E402
import tao_du_lieu_gia as gia  # noqa: E402

NGAY_CHOT = dt.date(2026, 9, 26)
COT_SMS = ["No.", "Partner Code", "Partner Name", "Salesman", "File/Job No.", "MBL/MAWB No.",
           "HBL/HAWB No.", "POR", "POD", "Vessel", "Flight", "ETD", "ETA", "Charge Name",
           "A/R Amt", "A/R paid", "A/R Paid Date", "A/R remain", "A/P Amt", "A/P Paid",
           "A/P Paid Date", "A/P remain", "ROE Local", "Curr."]


def tao_sms(dich: Path, nhom) -> Path:
    """nhom: [(hướng, [(mã, tên, sale, job, etd, eta, [(phí, còn nợ, roe)])])]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, t in enumerate(COT_SMS, start=1):
        ws.cell(row=2, column=i, value=t)
    r, stt = 3, 1
    for huong, jobs in nhom:
        ws.cell(row=r, column=1, value="+ 1")
        ws.cell(row=r, column=2, value=f"+ {huong}")
        r += 1
        for ma, ten, sale, job, etd, eta, phi in jobs:
            for j, (ten_phi, con, roe) in enumerate(phi):
                ws.cell(row=r, column=1, value=stt)
                if j == 0:
                    for c, v in ((2, ma), (3, ten), (4, sale), (5, job), (12, etd), (13, eta)):
                        ws.cell(row=r, column=c, value=v)
                ws.cell(row=r, column=14, value=ten_phi)
                ws.cell(row=r, column=15, value=con)
                if con:
                    ws.cell(row=r, column=18, value=con)
                ws.cell(row=r, column=23, value=roe)
                r += 1
                stt += 1
    ws.cell(row=r, column=1, value=stt)
    ws.cell(row=r, column=2, value="Sub-Total")
    ws.cell(row=r, column=18, value=999999999)
    dich.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dich)
    return dich


def tao_ghep(dich: Path, dong1, dong2=()) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ghép khách"
    ws.append(["Mã SMS", "Tên khách SMS", "Còn nợ theo SMS (VND)", "Số job còn nợ", "Gợi ý: mã Misa",
               "Gợi ý: tên Misa", "Số dư Misa (VND)", "Mức tin cậy", "Xác nhận",
               "Mã Misa đúng (nếu Sai)", "Tên Misa (tự hiện)", "Ghi chú"])
    for ma, goi_y, muc, xn, dung in dong1:
        ws.append([ma, ma, 0, 1, goi_y, "", 0, muc, xn, dung, None, None])
    ws2 = wb.create_sheet("Misa chưa ghép")
    ws2.append(["Mã khách Misa", "Tên khách Misa", "Số dư Nợ (VND)", "Mã SMS tương ứng", "Hoặc chọn", "Ghi chú"])
    for d in dong2:
        ws2.append(list(d))
    dich.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dich)
    return dich


# --- đọc file SMS ----------------------------------------------------------

def test_doc_sms_thua_huong_ma_va_quy_vnd(tmp_path):
    f = tao_sms(tmp_path / "AR-AP.xlsx", [
        ("Inbound", [("ABC", "ABC LTD", "Hằng", "IMHANA26090001", "01/09/2026", "10/09/2026",
                      [("FREIGHT", 100, 26000), ("THC", 50, 26000), ("PAID", None, 26000)])]),
        ("Outbound", [("XYZ", "XYZ CO", "Hiếu", "EXSANA26080002", "05/08/2026", None,
                       [("FEE", 1_000_000, 1)]),
                      ("NOD", "NO DATE", "", "LOSANA26070003", None, None, [("FEE", 500, 1)])]),
    ])
    assert ghep.la_file_sms(f)
    ds = ghep.doc_sms(f)
    assert len(ds) == 4, "dòng đã trả hết và dòng Sub-Total phải bị bỏ"
    abc = [x for x in ds if x.ma_sms == "ABC"]
    assert len(abc) == 2 and all(x.job == "IMHANA26090001" for x in abc), "dòng phí thừa hưởng mã và job"
    assert abc[0].so_tien == 2_600_000 and abc[0].ngay == dt.date(2026, 9, 10) and abc[0].nguon_ngay == "ETA"
    xyz = next(x for x in ds if x.ma_sms == "XYZ")
    assert xyz.ngay == dt.date(2026, 8, 5) and xyz.nguon_ngay == "ETD"
    nod = next(x for x in ds if x.ma_sms == "NOD")
    assert nod.ngay == dt.date(2026, 7, 15) and nod.nguon_ngay == "job", "thiếu ngày thì đọc năm-tháng trong mã job"


def test_file_misa_khong_bi_nhan_nham_la_sms(tmp_path):
    assert not ghep.la_file_sms(gia.file_tong_hop(tmp_path / "th.xlsx"))


# --- bảng ghép khách -------------------------------------------------------

def test_doc_bang_ghep_va_xung_dot_theo_phia_misa(tmp_path):
    f = tao_ghep(tmp_path / "g.xlsx", [
        ("A1", "M1", "Khớp tên/mã", None, None),        # khớp sẵn, FIN không cần chọn
        ("A2", "M9", "Thấp", "Sai", "M2"),
        ("A3", "M3", "Cao", "Đúng", None),
        ("A4", "M4", "Thấp", "Không có trong Misa", None),
        ("A5", "M5", "Thấp", None, None),                # chưa xác nhận -> lỗi
        ("A6", "M6", "Thấp", "Sai", None),               # Sai mà không điền mã -> lỗi
        ("TT", "M7", "Cao", "Đúng", None),
    ], [
        ("M8", "Viện", 1, "TT", "Nợ cũ", None),          # TT ghép 2 nơi -> phía Misa thắng
        ("HK", "Hoàng Kim", 1, None, "không có SMS", "Số nợ không có khả năng thu hồi"),
    ])
    b = ghep.doc_bang_ghep(f)
    assert b.sms_sang_misa["A1"] == "M1" and b.sms_sang_misa["A2"] == "M2" and b.sms_sang_misa["A3"] == "M3"
    assert "A4" in b.khong_misa and "A5" not in b.sms_sang_misa and "A6" not in b.sms_sang_misa
    assert len(b.loi) == 2
    assert b.sms_sang_misa["TT"] == "M8" and b.xung_dot == [("TT", "M7", "M8")]
    assert "HK" in b.kho_doi


# --- luật ghép tiền và tuổi --------------------------------------------------

def _misa(*khach):
    th = misa.TongHop(tu_ngay=NGAY_CHOT - dt.timedelta(days=7), den_ngay=NGAY_CHOT - dt.timedelta(days=1))
    for ma, ten, tien in khach:
        tc = misa.chuan_hoa_ten(ten)
        th.dong[tc] = misa.DongTongHop(ma, ten, tc, tien, 0, 0)
    return th, misa.TuoiNo(den_ngay=NGAY_CHOT)


def _sms(ma, tien, ngay, job="J", ten=None):
    return ghep.DongSMS(ma, ten or ma, "Sale", job, "Outbound", ngay, "ETD", tien)


def test_misa_quyet_tien_sms_quyet_tuoi(tmp_path):
    th, tn = _misa(("M1", "KHACH MOT", 1_000_000))
    so = sg.SoGhi(tmp_path / "so.json")
    kq = tinhtoan.tinh_bang_sms(th, tn, so, [
        _sms("M1", 600_000, NGAY_CHOT - dt.timedelta(days=10), "J1"),   # hạn +30 -> còn hạn
        _sms("M1", 400_000, NGAY_CHOT - dt.timedelta(days=50), "J2"),   # quá hạn 20 ngày
    ])
    d = kq.dong[0]
    assert d.total == 1_000_000 and d.nhom[0] == 600_000 and d.nhom[1] == 400_000
    assert d.nguon == "sms" and "J1" in d.job and "J2" in d.job and not d.cho_xem_lai


def test_sms_nhieu_hon_misa_thi_cat_job_cu_nhat_va_liet_ke(tmp_path):
    th, tn = _misa(("M1", "KHACH MOT", 500_000))
    so = sg.SoGhi(tmp_path / "so.json")
    kq = tinhtoan.tinh_bang_sms(th, tn, so, [
        _sms("M1", 400_000, NGAY_CHOT - dt.timedelta(days=100), "CU"),
        _sms("M1", 500_000, NGAY_CHOT - dt.timedelta(days=5), "MOI"),
    ])
    d = kq.dong[0]
    assert sum(d.nhom) == pytest.approx(500_000) and d.nhom[0] == 500_000
    assert kq.sms_cat == [("KHACH MOT", "CU", 400_000, NGAY_CHOT - dt.timedelta(days=100))]
    assert d.job == "MOI"


def test_sms_it_hon_misa_thi_phan_du_lay_lo_cu_nhat_trong_so_ghi(tmp_path):
    th, tn = _misa(("M1", "KHACH MOT", 1_000_000))
    so = sg.SoGhi(tmp_path / "so.json")
    truoc = NGAY_CHOT - dt.timedelta(days=7)
    so.ghi_trang(truoc, {misa.chuan_hoa_ten("KHACH MOT"): sg.HoSoKhach(
        ma="M1", ten="KHACH MOT", tong=900_000, nguon="ban-tay", ly_do="đang kiện", job="OLD1",
        lo=[sg.Lo(700_000, (truoc - dt.timedelta(days=200)).isoformat()),      # rất cũ
            sg.Lo(200_000, (truoc - dt.timedelta(days=3)).isoformat())])})
    kq = tinhtoan.tinh_bang_sms(th, tn, so, [_sms("M1", 300_000, NGAY_CHOT - dt.timedelta(days=2), "J")])
    d = kq.dong[0]
    assert d.nhom[0] == 300_000 and d.nhom[5] == 700_000, "phần dư 700k phải lấy lô CŨ NHẤT"
    assert d.phan_nguon == {"sms": 300_000, "so-ghi": 700_000}
    assert d.ly_do == "đang kiện" and "OLD1" in d.job


def test_khong_co_trong_sms_lan_so_ghi_thi_uoc_tinh_va_to_do(tmp_path):
    th, tn = _misa(("M1", "KHACH MOT", 1_000_000))
    so = sg.SoGhi(tmp_path / "so.json")
    kq = tinhtoan.tinh_bang_sms(th, tn, so, [_sms("ZZ", 10, dt.date(2026, 7, 7), "X")])
    d = kq.dong[0]
    assert d.nguon == "uoc-tinh" and d.cho_xem_lai and sum(d.nhom) == 1_000_000
    assert kq.sms_chua_ghep and kq.sms_chua_ghep[0][0] == "ZZ"


def test_kho_doi_luon_120_va_chenh_ty_gia_khong_bao_dong(tmp_path):
    th, tn = _misa(("HK", "HOANG KIM", 470_000_000), ("M2", "KHACH HAI", 10_000_000))
    so = sg.SoGhi(tmp_path / "so.json")
    bang = ghep.BangGhep(kho_doi={"HK": "không có khả năng thu hồi"})
    kq = tinhtoan.tinh_bang_sms(th, tn, so, [_sms("M2", 10_040_000, NGAY_CHOT, "J")], bang)
    hk = next(d for d in kq.dong if d.ma == "HK")
    m2 = next(d for d in kq.dong if d.ma == "M2")
    assert hk.nhom[5] == 470_000_000 and hk.ly_do
    assert m2.total == 10_000_000 and sum(m2.nhom) == pytest.approx(10_000_000)
    assert not m2.cho_xem_lai and not kq.sms_cat, "vênh 40k do tỷ giá không phải nợ thừa"


def test_sms_tro_toi_khach_misa_het_no_thi_bao_gach_paid(tmp_path):
    th, tn = _misa(("M1", "KHACH MOT", 1_000_000))
    th.dong["DA TRA"] = misa.DongTongHop("PAID", "DA TRA", "DA TRA", 0, 0, 0)
    so = sg.SoGhi(tmp_path / "so.json")
    kq = tinhtoan.tinh_bang_sms(th, tn, so, [_sms("PAID", 5_000, NGAY_CHOT, "J9")])
    assert kq.sms_khong_misa and kq.sms_khong_misa[0][2] == "J9"
    assert sum(d.total for d in kq.dong) == 1_000_000


# --- bản final của kế toán trưởng --------------------------------------------

def tao_final(dich: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(["SOURCE DATA"])
    ws.append(["Source", "x.xlsx"])
    ws.append(["Reporting period", "20/09/2026 - 26/09/2026"])
    ws.append(["No.", "Code", "Customer's name", None, "Credit term", "Total receivables", "Not-yet-due",
               "1 - 30 days", "31 - 60 days", "61 - 90  days", "91 - 120  days", "Over 120 days",
               "Total overdue", "Salesman", "Salesman", "Job", "Reason for late payment", "Comment"])
    ws.append([None, None, "Total", None, None, 300, 100, 0, 0, 0, 0, 200, 200])
    ws.append([1, "KHT01", "KHACH MOT", None, "30 days", 300, 100, None, None, None, None, 200, 200,
               "Hằng", "Leena", "job: IMHANA23010011, IMHANA23020022", "đang kiện", "gọi lại"])
    wb.save(dich)
    return dich


def test_doc_ban_final_dang_sheet_data(tmp_path):
    f = tao_final(tmp_path / "w39 final.xlsx")
    assert baocao.ngay_chot_trong_file(f) == NGAY_CHOT
    r = baocao.doc_bao_cao(f)
    assert list(r) == [misa.chuan_hoa_ten("KHACH MOT")], "dòng Total không được đọc thành khách"
    k = r[misa.chuan_hoa_ten("KHACH MOT")]
    assert k.nhom == [100, 0, 0, 0, 0, 200] and k.credit_term == "30 days" and k.salesman == "Lê Thị Hằng", "tên ngắn đổi sang tên đầy đủ (CEO 08/10)"
    assert k.job == "IMHANA23010011, IMHANA23020022" and k.ly_do == "đang kiện" and k.ghi_chu == "gọi lại"


def test_ghi_chu_job_di_mot_vong_khong_nhan_doi():
    o = baocao.ghep_ghi_chu("A, B", "gọi lại")
    assert baocao.tach_ghi_chu(o) == ("A, B", "gọi lại")
    assert baocao.tach_ghi_chu("chỉ ghi chú") == ("", "chỉ ghi chú")


# --- workspace, 3 nút -----------------------------------------------------

@pytest.fixture
def ws(tmp_path, monkeypatch):
    from wr import caidat
    tep = tmp_path / "settings.json"
    monkeypatch.setattr(caidat, "TEP_CAU_HINH", tep)
    monkeypatch.setattr(runner.caidat, "TEP_CAU_HINH", tep)
    monkeypatch.setattr(runner, "NHAT_KY", tmp_path / "log.txt")
    return caidat.dat_workspace(tmp_path / "dich", tao_thu_muc_con=False)


def test_khoi_tao_co_ngan_ghep_khach(ws):
    assert ws.ghep.is_dir() and ws.ghep.name == "sample" and ws.vao.name == "input" and ws.ra == ws.workspace


def _tha_misa(cfg):
    gia.file_tong_hop(cfg.vao / "th.xlsx")
    gia.file_tuoi_no(cfg.vao / "tn.xlsx")


def _tha_sms(cfg, them=()):
    return tao_sms(cfg.vao / "AR-AP.xlsx", [("Outbound", [
        ("KHT01", "BINH MINH", "Hằng", "EXSANA26090001", "20/09/2026", None, [("FEE", 120_000_000, 1)]),
        ("KHT03", "SKYPORT", "Hiếu", "EXSANA26070002", "01/07/2026", None, [("FEE", 300_000_000, 1)]),
        *them,
    ])])


def test_co_2_file_misa_ma_chua_co_sms_thi_cho(ws):
    _tha_misa(ws)
    assert runner.quet_thu_muc(ws) == 0
    assert not list(ws.ra.glob("*.xlsx"))
    assert not any(f.name.startswith("[DONE]") for f in ws.vao.iterdir())


def test_cho_qua_han_thi_chay_luat_cu(ws, monkeypatch):
    _tha_misa(ws)
    monkeypatch.setattr(runner, "CHO_SMS_GIAY", 0)
    assert runner.quet_thu_muc(ws) == 1


def test_tha_3_file_ra_bao_cao_tong_bang_misa(ws):
    _tha_misa(ws)
    _tha_sms(ws)
    tao_ghep(ws.ghep / "ghep.xlsx", [("KHT01", "KHT01", "Khớp tên/mã", "Đúng", None),
                                      ("KHT03", "KHT03", "Khớp tên/mã", "Đúng", None)])
    assert runner.quet_thu_muc(ws) == 1
    ra = next(ws.ra.glob("*.xlsx"))
    may = baocao.doc_bao_cao(ra)
    tong = sum(k[3] for k in gia.KHACH)
    assert sum(v.tong for v in may.values()) == pytest.approx(tong)
    sky = may[misa.chuan_hoa_ten("SKYPORT LOGISTICS PTE. LTD")]
    assert sky.nhom[2] == pytest.approx(240_000_000), "job 01/07 + 30 ngày -> quá hạn 57 ngày"
    assert sorted(f.name for f in ws.vao.iterdir()) == ["[DONE] AR-AP.xlsx", "[DONE] th.xlsx", "[DONE] tn.xlsx"]
    chu = " ".join(str(c) for r in openpyxl.load_workbook(ra)["Review"].iter_rows(values_only=True)
                   for c in r if c)
    assert "SMS cần gạch paid" in chu and "Tuổi nợ tính từ đâu" in chu


def test_ban_final_cua_nguoi_khong_bi_ban_may_ghi_de(ws, monkeypatch):
    tao_final(ws.ban_tay / "w39 final.xlsx")
    _tha_misa(ws)
    monkeypatch.setattr(runner, "CHO_SMS_GIAY", 0)
    runner.quet_thu_muc(ws)
    so = sg.SoGhi(ws.so_ghi)
    assert all(v.nguon == "ban-tay" for v in so.trang["2026-09-26"].values())
    assert "w39 final.xlsx" in so.da_nap


def test_nut_kiem_tra_ghep_khach_sinh_file(ws):
    _tha_misa(ws)
    _tha_sms(ws, [("NEWCO", "PACIFIC ROUTE LTD", "Vi", "EXSANA26090009", "22/09/2026", None,
                    [("FEE", 95_000_000, 1)])])
    tao_ghep(ws.ghep / "ghep.xlsx", [("KHT01", "KHT01", "Khớp tên/mã", "Đúng", None)])
    assert runner.main(["--kiem-tra-ghep-khach", "--khong-mo-file"]) == 0
    f = next(ws.ghep.glob("KIEM_TRA_*.xlsx"))
    wb = openpyxl.load_workbook(f)
    assert {"Tóm tắt", "Cần sửa", "SMS mới chưa ghép", "Misa mới chưa ghép",
            "Đối chiếu từng khách"} <= set(wb.sheetnames)
    moi_sms = [r[0] for r in wb["SMS mới chưa ghép"].iter_rows(min_row=2, values_only=True) if r[0]]
    assert moi_sms == ["NEWCO"], "KHT03 trùng mã Misa thì tự ghép, chỉ NEWCO là khách mới"
    goi_y = next(wb["SMS mới chưa ghép"].iter_rows(min_row=2, values_only=True))
    assert goi_y[4] == "KHT10", "gợi ý đúng khách PACIFIC ROUTE LIMITED"
    assert ghep.tim_file_ghep(ws.ghep).name == "ghep.xlsx", "file KIEM_TRA không được coi là file ghép"


def test_nut_kiem_tra_khi_chua_co_file_ghep(ws):
    assert runner.main(["--kiem-tra-ghep-khach", "--khong-mo-file"]) == 2


def test_tuan_sau_doc_lai_bao_cao_tuan_truoc_o_goc_va_khong_ghi_de(ws, monkeypatch):
    """Bố cục gốc/input/sample: báo cáo nằm ở gốc. Kế toán sửa tay báo cáo tuần
    trước thì tuần sau tool tính tiếp từ bản đã sửa; chạy lại cùng tuần không ghi đè."""
    monkeypatch.setattr(runner, "CHO_SMS_GIAY", 0)
    _tha_misa(ws)
    assert runner.quet_thu_muc(ws) == 1
    bao_cao = next(ws.ra.glob("2026_W*.xlsx"))
    assert bao_cao.parent == ws.workspace
    # kế toán sửa tay: đổi lý do của một khách
    wb = openpyxl.load_workbook(bao_cao)
    sh = wb["Receivables"]
    for r in range(6, sh.max_row + 1):
        if sh.cell(row=r, column=3).value == "EASTWIND CARGO CO., LTD":
            sh.cell(row=r, column=14, value="KTT ghi: đang kiện")
    wb.save(bao_cao)
    time.sleep(1.1)
    os.utime(bao_cao)
    so = sg.SoGhi(ws.so_ghi)
    assert runner.nap_ban_tay(ws.ra, so) == 1
    k = so.trang["2026-09-26"][misa.chuan_hoa_ten("EASTWIND CARGO CO., LTD")]
    assert k.ly_do == "KTT ghi: đang kiện" and k.nguon == "ban-tay"
    # chạy lại cùng tuần: bản đã sửa còn nguyên, bản mới mang hậu tố
    for f in ws.vao.glob("[[]DONE[]] *"):
        f.rename(f.with_name(f.name[7:]))
    so.ghi()
    assert runner.quet_thu_muc(ws) == 1
    ten = sorted(f.name for f in ws.ra.glob("2026_W*.xlsx"))
    assert len(ten) == 2 and any("chay lai" in t for t in ten)
    sh = openpyxl.load_workbook(bao_cao)["Receivables"]
    assert any(sh.cell(row=r, column=14).value == "KTT ghi: đang kiện" for r in range(6, sh.max_row + 1))

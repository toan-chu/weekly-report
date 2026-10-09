# -*- coding: utf-8 -*-
"""Phép thử WR-06: job chưa thanh toán lấy từ Misa (Bán hàng + Sổ chi tiết TK131) và cột Đối chiếu.
Dữ liệu GIẢ, sinh ngay trong phép thử."""
import datetime as dt
from pathlib import Path

import openpyxl

from test_wr4 import (A, B, C, CT40, D, DEN, K, TH40, _tha_4_file, so, tao_chi_tiet,  # noqa: F401
                      tao_tong_hop, ws)
from wr import banhang, chitiet, misa, tindung  # noqa: E402
from wr import thuonghieu as th  # noqa: E402
import runner  # noqa: E402

DT = lambda d: dt.datetime.combine(d, dt.time())  # noqa: E731


def tao_ban_hang(dich, hoa_don):
    """hoa_don: [(ngày, số HĐ, khách, tiền, trạng thái)]"""
    wb = openpyxl.Workbook()
    w = wb.active
    w["A1"] = "Bán hàng"
    w.append([])
    w.append(["STT", "Ngày hạch toán", "Số chứng từ", "Số hóa đơn", "Khách hàng", "Tổng tiền thanh toán",
              "TT lập hóa đơn", "TT thanh toán", "TT xuất hàng"])
    for i, (n, so_hd, kh, tien, tt) in enumerate(hoa_don, start=1):
        w.append([i, DT(n), f"DT{i}", so_hd, kh, tien, "Đã lập", tt, ""])
    wb.save(dich)
    return Path(dich)


def tao_so_chi_tiet(dich, dong):
    """dong: [(ngày, số HĐ, khách, nợ, có, job, diễn giải)]"""
    wb = openpyxl.Workbook()
    w = wb.active
    w["A1"] = "SỔ CHI TIẾT TÀI KHOẢN"
    w["A2"] = "Tài khoản: 131, Loại tiền: <<Tổng hợp>>, Từ ngày 01/01/2026 đến ngày 08/10/2026"
    w.append([])
    w.append(["Ngày hạch toán", "Số chứng từ", "Ngày hóa đơn", "Số hóa đơn", "Diễn giải chung", "Diễn giải",
              "Tên đối tượng", "TK đối ứng", "Phát sinh Nợ", "Phát sinh Có", "Dư Nợ", "Dư Có",
              "Mã khoản mục chi phí", "Tên khoản mục chi phí", "Mã đối tượng THCP", "Mã thống kê"])
    for n, so_hd, kh, no, co, job, dg in dong:
        w.append([DT(n), "CT", DT(n) if so_hd else "", so_hd, dg, dg, kh, "5113" if no else "1121",
                  no, co, 0, 0, "", "", job, ""])
    wb.save(dich)
    return Path(dich)


# A khớp (100 chưa TT, tách 2 job; HĐ 2 đã TT) · B đỏ (130 > 100, có "một phần") · C vàng (98 vs 100)
# K cam (không có hoá đơn năm nay) · D hết nợ trên Misa mà HĐ vẫn chưa TT · LẠ không tìm ra khách
BH = [(D(2026, 8, 1), "00000001", A, 100, "Chưa thanh toán"),
      (D(2026, 9, 1), "00000002", A, 50, "Đã thanh toán"),
      (D(2026, 8, 1), "00000010", B, 30, "Chưa thanh toán"),
      (D(2026, 9, 1), "00000011", B, 40, "Thanh toán một phần"),
      (D(2026, 9, 2), "00000012", B, 60, "Chưa thanh toán"),
      (D(2026, 9, 20), "00000020", C, 98, "Chưa thanh toán"),
      (D(2026, 9, 3), "00000040", "CÔNG TY D", 25, "Chưa thanh toán"),
      (D(2026, 9, 4), "00000050", "CÔNG TY LẠ", 7, "Chưa thanh toán"),
      (D(2026, 10, 5), "00000060", A, 999, "Chưa thanh toán")]       # sau kỳ: không tính
SCT = [(D(2026, 8, 1), "00000001", A, 60, 0, "EXSANA26080001", "Bán hàng"),
       (D(2026, 8, 1), "00000001", A, 40, 0, "EXSANA26080002", "Bán hàng"),
       (D(2026, 9, 2), "00000012", B, 60, 0, "EXHANA26090012", "Bán hàng")]


def _job_misa(tmp_path, bh=BH, sct=SCT):
    t = misa.doc_tong_hop(tao_tong_hop(tmp_path / "th.xlsx", TH40))
    ct = {k: list(v) for k, v in CT40.items()}
    ct["CÔNG TY D"] = [(D(2026, 9, 3), D(2026, 9, 3), "00000040", "Phí", "5113", 25, 0),
                       (D(2026, 9, 20), None, "", "Thu tiền", "1121", 0, 25)]
    c = chitiet.doc_chi_tiet(tao_chi_tiet(tmp_path / "ct.xlsx", ct))
    jm = banhang.doc_job_misa(tao_ban_hang(tmp_path / "bh.xlsx", bh), tao_so_chi_tiet(tmp_path / "sct.xlsx", sct),
                              [(k.ma, k.ten) for k in t.dong.values()])
    return t, c, jm


def test_nhan_dang_hai_file_moi(tmp_path):
    assert misa.nhan_dang(tao_ban_hang(tmp_path / "x.xlsx", BH)) == "ban-hang"
    assert misa.nhan_dang(tao_so_chi_tiet(tmp_path / "y.xlsx", SCT)) == "so-chi-tiet"


def test_dong_con_tu_misa_va_mau_doi_chieu(tmp_path, so, monkeypatch):
    monkeypatch.setattr(tindung, "DUNG_SAI_KHOP", 0.5)     # số giả rất nhỏ; thật là 1.000 VND
    t, c, jm = _job_misa(tmp_path)
    kq = tindung.tinh_full(t, c, so, job_misa=jm)
    d = {x.ma: x for x in kq.dong}
    a = d["KA"]
    assert [(g["job"], g["tien"]) for g in a.dong_con] == [("EXSANA26080001", 60), ("EXSANA26080002", 40)], \
        "chỉ hoá đơn chưa TT, tách theo job ở Sổ chi tiết; HĐ đã TT và HĐ sau kỳ không vào"
    assert a.dong_con[0]["nhom"] == [0, 0, 60, 0, 0, 0], "HĐ 01/08 term 30 → hạn 31/08 → 31-60"
    assert a.doi_chieu["mau"] == "khop"
    b = d["KB"]
    assert b.doi_chieu["mau"] == "do" and b.doi_chieu["lech"] == -30, "luật FIN: 'một phần' vẫn liệt kê đủ"
    assert "một phần" in b.doi_chieu["goi_y"]
    assert any(g["job"].startswith("HĐ 00000010") for g in b.dong_con), "hoá đơn không có job vẫn hiện"
    assert d["KC"].doi_chieu["mau"] == "vang"
    assert d["KK"].doi_chieu["mau"] == "cam" and d["KK"].dong_con == []
    assert kq.het_no_con_hd == [("CÔNG TY D", 1, 25)]
    assert kq.ten_khong_ghep == [("CÔNG TY LẠ", 1, 7)]
    assert sum(x.total for x in kq.dong) == 370, "dư nợ dòng khách giữ nguyên số Misa"


def test_da_thanh_toan_sau_ngay_cuoi_ky_van_tinh_la_no(tmp_path, so):
    sct = SCT + [(D(2026, 10, 5), "", A, 0, 50, "", "Thu tiền khách hàng theo hóa đơn 00000002")]
    t, c, jm = _job_misa(tmp_path, sct=sct)
    a = next(x for x in tindung.tinh_full(t, c, so, job_misa=jm).dong if x.ma == "KA")
    hd2 = [g for g in a.dong_con if g["so_hd"] == "00000002"]
    assert hd2 and "sau kỳ" in hd2[0]["trang_thai"], "file xuất 08/10, phiếu thu 05/10 > cuối kỳ 02/10"


def test_tha_6_file_ra_cot_doi_chieu_va_review(ws, monkeypatch):
    monkeypatch.setattr(tindung, "DUNG_SAI_KHOP", 0.5)
    _tha_4_file(ws)
    tao_ban_hang(ws.vao / "bh.xlsx", BH)
    tao_so_chi_tiet(ws.vao / "sct.xlsx", SCT)
    monkeypatch.setattr(runner, "CHO_SMS_GIAY", 7200)   # đủ 6 file thì không phải chờ
    assert runner.quet_thu_muc(ws) == 1
    assert all(x.name.startswith("[DONE]") for x in ws.vao.glob("*.xlsx"))
    wb = openpyxl.load_workbook(next(ws.ra.glob("2026_W40_*.xlsx")))
    rec = wb[th.SHEET_BANG]
    assert rec.cell(5, 15).value == th.COT_DOI_CHIEU[0]
    hang = {rec.cell(r, 2).value: r for r in range(6, rec.max_row + 1) if rec.cell(r, 2).value}
    assert rec.cell(hang["KA"], 15).value == "Khớp"
    o = rec.cell(hang["KB"], 15)
    assert o.value.startswith("ĐỎ") and o.fill.fgColor.rgb.endswith("F8CBCB")
    assert rec.cell(hang["KA"] + 1, 3).value.strip().startswith("└ EXSANA26080001 · HĐ 00000001")
    assert rec.cell(hang["KA"] + 1, 15).value == "Chưa thanh toán"
    assert rec.cell(hang["KA"] + 1, 12).value == "Lê Thị Hằng", "sales của job lấy từ SMS"
    chu = " ".join(str(c) for r in wb[th.SHEET_REVIEW].iter_rows(values_only=True) for c in r if c)
    assert "ĐỎ — tổng job" in chu and "CÔNG TY D" in chu and "TÍM" in chu and "CÔNG TY LẠ" in chu
    inv = wb[th.SHEET_INVOICES]
    assert [c.value for c in inv[5]][-2:] == ["Misa status", "Job (Misa)"]


def test_thieu_ban_hang_thi_cho(ws, monkeypatch):
    _tha_4_file(ws)
    monkeypatch.setattr(runner, "CHO_SMS_GIAY", 7200)
    assert runner.quet_thu_muc(ws) == 0, "FIN thả 4 file trước: chờ Bán hàng + Sổ chi tiết"


def test_payable_dung_mau_fin(ws):
    """FIN 09/10: Payable y chang mẫu — 5 cột, chỉ vendor còn số dư, chỉ lấy số tổng."""
    _tha_4_file(ws)
    assert runner.quet_thu_muc(ws) == 1
    wb = openpyxl.load_workbook(next(ws.ra.glob("2026_W40_*.xlsx")))
    p = wb[th.SHEET_PAYABLE]
    assert [c.value for c in p[5]] == ["No.", "Code", "Vendor's Name", "Current", "Total"]
    r = [x for x in p.iter_rows(min_row=6, values_only=True) if x[1]]
    assert [x[1] for x in r] == ["V1", "V3", "TOTAL"], "V2 trả xong, V4 mình trả trước: không vào bảng"
    assert r[-1][3] == r[-1][4] == 650
    cash = " ".join(str(c) for x in wb[th.SHEET_CASH].iter_rows(values_only=True) for c in x if c)
    assert "VENDOR TRẢ XONG" in cash, "tiền đã trả vendor vẫn ở Cash Flow"

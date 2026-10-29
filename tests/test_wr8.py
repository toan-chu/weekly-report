# -*- coding: utf-8 -*-
"""Phép thử WR-08: công thức Head of FIN trong AR Risk, credit term trong sample/, khách nhiều Sales,
dư nợ lẻ, bản FIN sửa thẳng trên file tool thắng. Dữ liệu GIẢ."""
import os
import time

import openpyxl

from test_wr4 import A, B, CT40, TH40, _tha_4_file, so, tao_chi_tiet, tao_tong_hop, ws  # noqa: F401
from wr import baocao, chitiet, ghep, misa, tindung  # noqa: E402
from wr import soghi as sg  # noqa: E402
from wr import thuonghieu as th  # noqa: E402
import runner  # noqa: E402


def test_ar_risk_la_cong_thuc_cua_head_of_fin(ws):
    _tha_4_file(ws)
    assert runner.quet_thu_muc(ws) == 1
    r = openpyxl.load_workbook(next(ws.ra.glob("2026_W40_*.xlsx")))[th.SHEET_RISK]
    tieu = {c.value: c.column for c in r[5]}
    h = next(x for x in range(6, r.max_row + 1) if r.cell(x, 2).value == "KA")
    for ten in ("S1", "S2", "S3", "S4", "S5", "S6", "S7", "Signals (S1-S3, S7)", "Risk Tier", "Trend (auto)",
                "Reason for late payment", "Required action", "CEO review", "Total overdue"):
        v = r.cell(h, tieu[ten]).value
        assert isinstance(v, str) and v.startswith("="), f"{ten} phải là công thức"
    assert "Receivables" in r.cell(h, tieu["Reason for late payment"]).value, "lý do lấy từ Receivables"
    assert '"N3·Nợ xấu"' in r.cell(h, tieu["Risk Tier"]).value
    assert isinstance(r.cell(h, tieu["Ending = Total"]).value, (int, float)), "cột tiền vẫn là số"


def test_credit_term_trong_sample_thang(tmp_path, so):
    wb = openpyxl.Workbook()
    w = wb.active
    w.title = ghep.SHEET_CREDIT_TERM
    w.append(ghep.COT_CREDIT_TERM)
    w.append(["KA", A, "thanh toán ngay", "", "hợp đồng", ""])
    w.append(["KB", B, "", "", "", ""])
    wb.save(tmp_path / "SAMPLE_credit_term_khach.xlsx")
    ct_map = ghep.doc_credit_term(tmp_path)
    assert ct_map == {"KA": "thanh toán ngay"}, "ô trống bỏ qua"
    t = misa.doc_tong_hop(tao_tong_hop(tmp_path / "th.xlsx", TH40))
    c = chitiet.doc_chi_tiet(tao_chi_tiet(tmp_path / "ct.xlsx", CT40))
    kq = tindung.tinh_full(t, c, so, credit_terms=ct_map)
    d = {x.ma: x for x in kq.dong}
    assert d["KA"].credit_term == "thanh toán ngay", "bảng trong sample thắng lệnh khai tay (30 days)"
    assert d["KB"].credit_term == "30 days"
    assert ghep.hieu_term("14 ngày của tháng tiếp theo") == "ngày 14 của tháng sau tháng hoá đơn"


def test_khach_nhieu_sales_ghi_ro_job_cua_ai():
    s = tindung.chia_theo_sales([{"sales": "Trần Văn Hiếu", "tien": 45e6, "job": "J1"},
                                 {"sales": "Lê Thị Hằng", "tien": 12e6, "job": "J2"},
                                 {"sales": "Trần Văn Hiếu", "tien": 1e6, "job": "J3"}])
    assert s == "Trần Văn Hiếu (46tr): J1, J3 | Lê Thị Hằng (12tr): J2"
    assert tindung.chia_theo_sales([{"sales": "A", "tien": 1, "job": "J"}, {"sales": "", "tien": 1, "job": "K"}]) == ""


def test_du_no_le_khong_vao_bang(tmp_path, so, monkeypatch):
    monkeypatch.setattr(tindung, "NGUONG_DU_LE", 1000)
    th_ = TH40 + [("KL", "KHÁCH LẺ", 84, 0, 0, 0, 84, 0)]
    t = misa.doc_tong_hop(tao_tong_hop(tmp_path / "th.xlsx", [x if x[0] != "KA" else ("KA", A, 1500, 0, 0, 0, 1500, 0)
                                                              for x in th_]))
    c = chitiet.doc_chi_tiet(tao_chi_tiet(tmp_path / "ct.xlsx", CT40))
    kq = tindung.tinh_full(t, c, so)
    assert "KL" not in {x.ma for x in kq.dong}
    assert any(x[0] == "KHÁCH LẺ" for x in kq.da_bo)


def test_fin_sua_thang_tren_file_tool_thi_thang(ws):
    """W41: FIN sửa lý do ngay trên file tool sinh (còn tab Review) — tuần sau phải lấy lý do của FIN."""
    _tha_4_file(ws)
    assert runner.quet_thu_muc(ws) == 1
    f = next(ws.ra.glob("2026_W40_*.xlsx"))
    so = sg.SoGhi(ws.so_ghi)
    assert f.name in so.file_tool
    time.sleep(1.1)
    wb = openpyxl.load_workbook(f)
    sh = wb[th.SHEET_BANG]
    for r in range(6, sh.max_row + 1):
        if sh.cell(r, 2).value == "KA":
            sh.cell(r, th.COT_LY_DO_NAY, value="FIN: khách hứa trả 15/10")
    wb.save(f)
    os.utime(f, (time.time() + 5, time.time() + 5))
    assert runner.nap_ban_tay(ws.ra, so) == 1
    k = so.trang["2026-10-03"][misa.chuan_hoa_ten(A)]
    assert k.ly_do == "FIN: khách hứa trả 15/10"
    assert so.nguon_trang["2026-10-03"]["ban_nguoi"] is True


def test_client_master_tool_doc_duoc_va_tu_tinh_trang_thai(tmp_path):
    """File Client Master (OneDrive) theo theme: tool đọc đúng cột Mã kh / Thời hạn, cột trạng thái là công thức,
    dòng trống kéo sẵn công thức cho Sales thêm khách."""
    import datetime as dt
    from wr import khachhang as kh
    dong = [dict(ma="KH001", ten="Khách A", term="30 days", het_han=dt.date(2026, 1, 1)),
            dict(ma="KH002", ten="Khách B", term=None),
            dict(ma="KH003", ten="Khách C", term="thanh toán ngay", gia_han=dt.date(2026, 2, 1))]
    f = kh.ghi_master(tmp_path / "SAMPLE_credit_term_khach.xlsx", dong, ["Trần Văn Hiếu"])
    assert ghep.doc_credit_term(tmp_path) == {"KH001": "30 days", "KH003": "thanh toán ngay"}
    w = openpyxl.load_workbook(f)[kh.SHEET_MASTER]
    c = kh.VI_TRI
    assert w[f"{c['trang_thai']}7"].value.startswith("=IF(")
    assert w[f"{c['het_moi']}9"].value == f"=IF({c['gia_han']}9=\"\",\"\",EDATE({c['gia_han']}9,12))"
    assert w[f"{c['stt']}{kh.DONG_CUOI}"].value.startswith("=IF(")      # dòng trống vẫn có công thức


def test_lan_dau_tool_tao_client_master(tmp_path):
    from wr import kiemtra, khachhang as kh
    f = kiemtra.tao_file_credit_term(tmp_path, [("KH001", "Khách A", "30 days"), ("KH002", "Khách B", "")],
                                     lambda m: None)
    assert kh.SHEET_MASTER in openpyxl.load_workbook(f).sheetnames
    assert ghep.doc_credit_term(tmp_path) == {"KH001": "30 days"}
    assert kiemtra.tao_file_credit_term(tmp_path, [], lambda m: None) is None   # đã có thì không tạo lại


def test_lo_hen_s7_day_len_no_xau():
    """user 13/10 (phương án A): lý do hẹn 'tháng 9' mà báo cáo đã sang tháng 10, quá hạn không giảm
    → thêm dấu hiệu S7; cùng S2 thành 2 dấu hiệu → N3 (ca ATH W41)."""
    import datetime as dt
    chot = dt.date(2026, 10, 10)
    assert tindung.ngay_hen("Hẹn thanh toán trong tháng 9", chot) == dt.date(2026, 9, 30)
    assert tindung.ngay_hen("khách hẹn 15/9, đầu tháng 11 trả nốt", chot) == dt.date(2026, 11, 10)
    assert tindung.ngay_hen("trả trong 3 tháng, đã thu 1.500 USD", chot) is None
    assert tindung.ngay_hen("tháng 2", chot) == dt.date(2027, 2, 28)
    def dong(ly_do, qh_tuan_truoc):
        return tindung.DongRuiRo(ma="ATH", ten="ATH", credit_term="thanh toán ngay", dau_ky=100, tang=0, giam=0,
                                 cuoi_ky=100, nhom=[0, 0, 100, 0, 0, 0], salesman="", job="", ly_do=ly_do,
                                 tt_total=100, tt_qua_han=qh_tuan_truoc, tt_61=0, co_tuan_truoc=True,
                                 nhap_bo_sung=0)
    r = tindung.danh_gia_rui_ro(dong("Hẹn thanh toán trong tháng 9", 100), chot)
    assert r.s7 == 1 and r.so_dau_hieu == 2 and r.nhom_rui_ro.startswith("N3"), r.nhom_rui_ro
    assert "S7" in r.ma_dau_hieu and "30/09/2026" in r.xu_huong
    r = tindung.danh_gia_rui_ro(dong("Hẹn thanh toán cuối tháng 10", 100), chot)
    assert r.s7 == 0 and r.nhom_rui_ro.startswith("N2"), "hẹn chưa tới → không phải lỡ hẹn"
    r = tindung.danh_gia_rui_ro(dong("Hẹn thanh toán trong tháng 9", 150), chot)
    assert r.s7 == 0, "quá hạn đã giảm so với tuần trước → không tính lỡ hẹn"


def test_canh_bao_hop_dong_tu_client_master(tmp_path):
    """user 13/10: báo cáo tuần cảnh báo hợp đồng hết hạn còn nợ, sắp hết hạn, chưa có term, term vừa đổi.
    Hợp đồng hết hạn KHÔNG đổi term (vẫn tính theo hợp đồng cũ)."""
    import datetime as dt
    from wr import khachhang as kh
    chot = dt.date(2026, 10, 10)
    kh.ghi_master(tmp_path / "SAMPLE_credit_term_khach.xlsx", [
        dict(ma="A", ten="Khách A", term="30 days", het_han=dt.date(2026, 7, 20)),            # hết hạn, còn nợ
        dict(ma="B", ten="Khách B", term="15 days", gia_han=dt.date(2025, 10, 25)),          # gia hạn → 25/10/2026
        dict(ma="C", ten="Khách C", term="45 days"),                                         # term đổi
        dict(ma="E", ten="Khách E", term="7 days", het_han=dt.date(2026, 1, 1)),             # hết hạn, không nợ
    ], [])
    m = kh.doc_master(tmp_path)
    assert m["B"]["han_dung"] == dt.date(2026, 10, 25)
    def r(ma, term, tang=0.0):
        return tindung.DongRuiRo(ma=ma, ten=ma, credit_term=term, dau_ky=100, tang=tang, giam=0, cuoi_ky=100,
                                 nhom=[100, 0, 0, 0, 0, 0], salesman="S", job="", ly_do="", tt_total=100,
                                 tt_qua_han=0, tt_61=0, co_tuan_truoc=True, nhap_bo_sung=0)
    cb = kh.canh_bao_hop_dong(m, [r("A", "30 days", tang=50), r("C", "45 days"), r("D", "")], chot,
                              {"C": "30 days", "A": "30 days"})
    loai = {(c["ma"], c["loai"]) for c in cb}
    assert ("A", "Hợp đồng đã hết hạn") in loai and "phát sinh 50" in cb[0]["ghi_chu"]
    assert ("B", "Sắp hết hạn ≤30 ngày") in loai, "sắp hết hạn: báo cả khách không còn nợ"
    assert ("C", "Credit term vừa đổi") in loai
    # WR-11 (FIN 27/10): khách đang nợ mà chưa có dòng nào trong Client Master tách thành loại riêng
    assert ("D", "Chưa có trong file credit term") in loai
    assert not any(c["ma"] == "E" for c in cb), "hết hạn mà không còn nợ thì không báo"


def test_client_master_tren_onedrive_la_nguon_duy_nhat(tmp_path):
    """File Client Master chung nằm ngoài sample/ (OneDrive): có file → chỉ đọc file đó; mất file → về sample/."""
    from wr import khachhang as kh
    mau, od = tmp_path / "sample", tmp_path / "OneDrive"
    mau.mkdir(); od.mkdir()
    kh.ghi_master(mau / "SAMPLE_credit_term_khach.xlsx", [dict(ma="A", ten="A", term="7 days"),
                                                         dict(ma="B", ten="B", term="15 days")], [])
    kh.ghi_master(od / "Client Master.xlsx", [dict(ma="A", ten="A", term="30 days")], [])
    ghep.dat_file_master(od / "Client Master.xlsx")
    assert ghep.doc_credit_term(mau) == {"A": "30 days"}, "chỉ một nguồn sự thật"
    assert set(kh.doc_master(mau)) == {"A"}
    ghep.dat_file_master(od / "khong-co.xlsx")
    assert ghep.doc_credit_term(mau) == {"A": "7 days", "B": "15 days"}

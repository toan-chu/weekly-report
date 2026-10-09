# -*- coding: utf-8 -*-
"""Phép thử của tool báo cáo công nợ tuần.

    python -m pytest -q
"""
import datetime as dt
import sys
from pathlib import Path

import pytest

GOC = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(GOC))
FIX = GOC / "handoff" / "docs" / "fixtures"

from wr import baocao, misa  # noqa: E402
from wr import soghi as sg  # noqa: E402
from wr import tinhtoan  # noqa: E402
import runner  # noqa: E402

TH = {w: FIX / f"W{w}_Tong_hop_cong_no_phai_thu_khach_hang.xlsx" for w in (36, 37, 38)}
TN = {w: FIX / f"W{w}_Phan_tich_cong_no_phai_thu_theo_tuoi_no.xlsx" for w in (36, 37, 38)}
TAY = {
    36: FIX / "2026_W36_Bang_cong_no_tuan_29.08-04.09.2026.xlsx",
    37: FIX / "2026_W37_Bang_cong_no_tuan_05.09-11.09.2026.xlsx",
    38: FIX / "2026_W38_Bang_cong_no_tuan_12.09-18.09.2026 (final).xlsx",
}
MAU = GOC / "template" / "Bang_cong_no_MAU_CHUAN.xlsx"

CO_DU_LIEU_THAT = all(f.exists() for f in list(TH.values()) + list(TN.values()) + list(TAY.values()))
can_du_lieu_that = pytest.mark.skipif(
    not CO_DU_LIEU_THAT,
    reason="Không có dữ liệu thật trong handoff/docs/fixtures/ — đây là dữ liệu công nợ "
           "của Trustana, không đi theo repo công khai. Chép vào đó thì các phép thử này chạy lại.",
)


# --- đọc file Misa --------------------------------------------------------

@can_du_lieu_that
def test_doc_dung_ky_va_so_khach():
    th = misa.doc_tong_hop(TH[37])
    assert (th.tu_ngay, th.den_ngay) == (dt.date(2026, 9, 5), dt.date(2026, 9, 11))
    assert len(th.khach_con_no()) == 74
    assert round(th.tong_du_no) == 9_547_752_597


@can_du_lieu_that
def test_nhan_dang_file_khong_dua_vao_ten():
    assert misa.nhan_dang(TH[38]) == "tong-hop"
    assert misa.nhan_dang(TN[38]) == "tuoi-no"


@can_du_lieu_that
def test_hai_file_lech_ky_thi_dung():
    th, tn = misa.doc_tong_hop(TH[37]), misa.doc_tuoi_no(TN[38])
    with pytest.raises(misa.LoiDocFile):
        misa.kiem_cung_ky(th, tn)


@can_du_lieu_that
def test_file_khong_phai_cua_misa_thi_bao_loi():
    with pytest.raises(misa.LoiDocFile):
        misa.doc_tong_hop(TAY[37])


# --- luật tính tuổi -------------------------------------------------------

def test_tuoi_tinh_theo_ngay_that_khong_theo_so_lan_chay():
    D = dt.date(2026, 9, 12)
    lo = [sg.Lo(100.0, (D - dt.timedelta(days=5)).isoformat())]
    assert sg.nhom_tu_lo(lo, D) == [0, 100, 0, 0, 0, 0]
    assert sg.nhom_tu_lo(lo, D + dt.timedelta(days=28)) == [0, 0, 100, 0, 0, 0]
    assert sg.nhom_tu_lo(lo, D + dt.timedelta(days=120)) == [0, 0, 0, 0, 0, 100]


def test_no_tang_vao_current_no_giam_tru_khoan_cu_nhat():
    D = dt.date(2026, 9, 12)
    lo = [sg.Lo(100.0, (D - dt.timedelta(days=40)).isoformat()),
          sg.Lo(50.0, (D + dt.timedelta(days=10)).isoformat())]
    tang = sg.dieu_chinh_lo(lo, 200.0, D, 30)
    assert sg.nhom_tu_lo(tang, D) == [100, 0, 100, 0, 0, 0]
    giam = sg.dieu_chinh_lo(lo, 60.0, D, 30)
    assert sg.nhom_tu_lo(giam, D) == [50, 0, 10, 0, 0, 0]


def test_credit_term():
    assert sg.so_ngay_term("30 days") == 30
    assert sg.so_ngay_term("thanh toán ngay") == 0
    assert sg.so_ngay_term("") == 30


# --- dựng bảng ------------------------------------------------------------

def _chay(tmp_path, tuan):
    so_ghi = tmp_path / "so-ghi.json"
    runner.main(["--nap-bao-cao", str(TAY[tuan - 1]), "--ngay-chot",
                 (misa.doc_tong_hop(TH[tuan - 1]).den_ngay + dt.timedelta(days=1)).isoformat(),
                 "--so-ghi", str(so_ghi)])
    runner.main(["--tong-hop", str(TH[tuan]), "--tuoi-no", str(TN[tuan]),
                 "--ra", str(tmp_path), "--so-ghi", str(so_ghi), "--mau-chuan", str(MAU)])
    ra = list(tmp_path.glob("2026_W*.xlsx"))
    assert len(ra) == 1
    return ra[0], so_ghi


@can_du_lieu_that
def test_chay_w37_ra_dung_file_va_dung_total(tmp_path):
    dich, _ = _chay(tmp_path, 37)
    assert dich.name == "2026_W37_Credit_Report_05.09-11.09.2026.xlsx"
    may = baocao.doc_bao_cao(dich)
    th = misa.doc_tong_hop(TH[37])
    assert len(may) == 74
    for k, v in may.items():
        assert abs(v.tong - th.dong[k].du_no) < 2, k
        assert abs(sum(v.nhom) - v.tong) < 2, f"{k}: tổng nhóm khác Total"
    assert abs(sum(v.tong for v in may.values()) - th.tong_du_no) < 2


@can_du_lieu_that
def test_w37_khop_phan_lon_ban_lam_tay(tmp_path):
    """Ngưỡng canh hồi quy. Lệch còn lại phải giải thích được theo
    handoff/docs/NOTE-doi-chieu-3-tuan.md."""
    dich, _ = _chay(tmp_path, 37)
    may, tay = baocao.doc_bao_cao(dich), baocao.doc_bao_cao(TAY[37])
    gop = lambda n: n[:3] + [n[3] + n[4] + n[5]]
    khop = sum(
        1 for k in tay
        if k in may and abs(may[k].tong - tay[k].tong) < 2
        and all(abs(x - y) < 2 for x, y in zip(gop(may[k].nhom), gop(tay[k].nhom)))
    )
    assert khop >= 62, f"chỉ khớp {khop}/74, luật tính đã đổi so với lúc chốt"


@can_du_lieu_that
def test_khach_tra_truoc_khong_vao_bao_cao(tmp_path):
    dich, _ = _chay(tmp_path, 37)
    may = baocao.doc_bao_cao(dich)
    th = misa.doc_tong_hop(TH[37])
    tra_truoc = [k for k, v in th.dong.items() if v.du_co > 0 and v.du_no <= 0]
    assert tra_truoc and all(k not in may for k in tra_truoc)


@can_du_lieu_that
def test_ngoai_le_ep_dung_nhom(tmp_path):
    so_ghi = tmp_path / "so-ghi.json"
    runner.main(["--ngoai-le", "BSF LLC", "--nhom", "1 - 30", "--vi-sao", "lô tàu chìm",
                 "--so-ghi", str(so_ghi)])
    runner.main(["--tong-hop", str(TH[38]), "--tuoi-no", str(TN[38]),
                 "--ra", str(tmp_path), "--so-ghi", str(so_ghi), "--mau-chuan", str(MAU)])
    may = baocao.doc_bao_cao(next(tmp_path.glob("2026_W38*.xlsx")))
    bsf = may[misa.chuan_hoa_ten("BSF LLC")]
    assert bsf.nhom[1] == pytest.approx(bsf.tong, abs=2)


@can_du_lieu_that
def test_so_ghi_chay_lai_cung_ngay_thi_ghi_de(tmp_path):
    dich, so_ghi = _chay(tmp_path, 37)
    so = sg.SoGhi(so_ghi)
    assert set(so.trang) == {"2026-09-05", "2026-09-12"}
    runner.main(["--tong-hop", str(TH[37]), "--tuoi-no", str(TN[37]),
                 "--ra", str(tmp_path), "--so-ghi", str(so_ghi), "--mau-chuan", str(MAU)])
    assert set(sg.SoGhi(so_ghi).trang) == {"2026-09-05", "2026-09-12"}


@can_du_lieu_that
def test_nghi_mot_tuan_van_ra_bao_cao_va_gia_them_tuoi(tmp_path):
    """Bỏ W37, nhảy thẳng từ W36 sang W38: vẫn ra file, và nợ cũ già thêm đúng
    số ngày thật."""
    so_ghi = tmp_path / "so-ghi.json"
    runner.main(["--nap-bao-cao", str(TAY[36]), "--ngay-chot", "2026-09-05",
                 "--so-ghi", str(so_ghi)])
    runner.main(["--tong-hop", str(TH[38]), "--tuoi-no", str(TN[38]),
                 "--ra", str(tmp_path), "--so-ghi", str(so_ghi), "--mau-chuan", str(MAU)])
    may = baocao.doc_bao_cao(next(tmp_path.glob("2026_W38*.xlsx")))
    th = misa.doc_tong_hop(TH[38])
    assert len(may) == len(th.khach_con_no())
    for v in may.values():
        assert abs(sum(v.nhom) - v.tong) < 2


@can_du_lieu_that
def test_ke_toan_khai_credit_term_moi(tmp_path):
    """Grand Forwarding mới ký hợp đồng 15 ngày — khai một lần, tool dùng mãi."""
    so_ghi = tmp_path / "so-ghi.json"
    runner.main(["--credit-term", "GRAND FORWARDING LIMITED", "--gia-tri", "15 days",
                 "--so-ghi", str(so_ghi)])
    runner.main(["--tong-hop", str(TH[38]), "--tuoi-no", str(TN[38]),
                 "--ra", str(tmp_path), "--so-ghi", str(so_ghi), "--mau-chuan", str(MAU)])
    may = baocao.doc_bao_cao(next(tmp_path.glob("2026_W38*.xlsx")))
    assert may[misa.chuan_hoa_ten("GRAND FORWARDING LIMITED")].credit_term == "15 days"


@can_du_lieu_that
def test_ngoai_le_ahc_ve_current(tmp_path):
    """AHC đang chờ ký biên bản bù trừ nên kế toán để ở nợ trong hạn."""
    so_ghi = tmp_path / "so-ghi.json"
    runner.main(["--ngoai-le", "AHC LOGISTICS(XIAMEN) CO.,LTD", "--nhom", "Current",
                 "--vi-sao", "cho ky bien ban bu tru", "--so-ghi", str(so_ghi)])
    runner.main(["--tong-hop", str(TH[38]), "--tuoi-no", str(TN[38]),
                 "--ra", str(tmp_path), "--so-ghi", str(so_ghi), "--mau-chuan", str(MAU)])
    may = baocao.doc_bao_cao(next(tmp_path.glob("2026_W38*.xlsx")))
    ahc = may[misa.chuan_hoa_ten("AHC LOGISTICS(XIAMEN) CO.,LTD")]
    assert ahc.nhom[0] == pytest.approx(ahc.tong, abs=2) and ahc.nhom[1] == 0


def test_dat_thu_muc_dich_va_tao_cac_ngan(tmp_path, monkeypatch):
    """Người dùng trỏ thư mục đích bất kỳ; tool tạo sẵn các ngăn bên trong và
    không để lại dữ liệu nào trong thư mục mã nguồn."""
    from wr import caidat
    tep = tmp_path / "settings.json"
    monkeypatch.setattr(caidat, "TEP_CAU_HINH", tep)
    monkeypatch.setattr(runner.caidat, "TEP_CAU_HINH", tep)
    dich = tmp_path / "O D" / "Cong no"
    cfg = caidat.dat_workspace(dich)
    assert cfg.workspace == dich / "Cong-No-Workspace"
    for d in cfg.cac_thu_muc():
        assert d.is_dir()
    assert caidat.dat_workspace(dich, tao_thu_muc_con=False).workspace == dich


@can_du_lieu_that
def test_chay_tay_van_dung_thu_muc_dich_da_dat(tmp_path, monkeypatch):
    from wr import caidat
    tep = tmp_path / "settings.json"
    monkeypatch.setattr(caidat, "TEP_CAU_HINH", tep)
    monkeypatch.setattr(runner.caidat, "TEP_CAU_HINH", tep)
    cfg = caidat.dat_workspace(tmp_path / "dich", tao_thu_muc_con=False)
    runner.main(["--nap-bao-cao", str(TAY[37]), "--ngay-chot", "2026-09-12"])
    assert cfg.so_ghi.exists(), "sổ ghi phải nằm trong thư mục đích"
    assert not (GOC / "so-ghi.json").exists(), "không được ghi vào thư mục mã nguồn"


@can_du_lieu_that
def test_sheet_can_xem_lai_giai_thich_chenh_lech_hai_file(tmp_path):
    """Kế toán hay so tổng báo cáo với tổng file Tuổi nợ rồi thấy lệch. Sheet
    Cần xem lại phải nói thẳng vì sao và chỉ đúng khách gây lệch."""
    import openpyxl
    dich, _ = _chay(tmp_path, 38)
    ws = openpyxl.load_workbook(dich)["Review"]
    chu = "\n".join(str(c.value or "") for r in ws.iter_rows(max_row=80) for c in r)
    assert "Tổng file Tổng hợp công nợ" in chu
    assert "Tổng file Phân tích tuổi nợ" in chu
    assert "KHÔNG dùng để đối chiếu" in chu
    assert "Khách làm lệch hai file" in chu


@can_du_lieu_that
def test_giai_thich_no_trong_han_khac_file_tuoi_no(tmp_path):
    import openpyxl
    dich, _ = _chay(tmp_path, 38)
    ws = openpyxl.load_workbook(dich)["Review"]
    chu = "\n".join(str(c.value or "") for r in ws.iter_rows(max_row=80) for c in r)
    assert "Nợ trong hạn trên báo cáo" in chu
    assert "chưa xác định được tuổi" in chu


@can_du_lieu_that
def test_tha_nhieu_tuan_cung_luc_chay_lan_luot_tu_ky_cu_nhat(tmp_path, monkeypatch):
    """Thả W36, W37, W38 vào cùng một lúc: tool ghép đúng cặp cùng kỳ và chạy
    theo thứ tự thời gian, không trộn file của hai tuần khác nhau."""
    import shutil
    from wr import caidat
    tep = tmp_path / "settings.json"
    monkeypatch.setattr(caidat, "TEP_CAU_HINH", tep)
    monkeypatch.setattr(runner.caidat, "TEP_CAU_HINH", tep)
    cfg = caidat.dat_workspace(tmp_path / "dich", tao_thu_muc_con=False)
    shutil.copy(MAU, cfg.mau_chuan) if not cfg.mau_chuan.exists() else None
    for w in (38, 36, 37):                       # cố tình thả lộn xộn
        shutil.copy(TH[w], cfg.vao / TH[w].name)
        shutil.copy(TN[w], cfg.vao / TN[w].name)
    assert runner.quet_thu_muc(cfg) == 3
    ra = sorted(f.name for f in cfg.ra.glob("*.xlsx"))
    assert ra == ["2026_W36_Credit_Report_29.08-04.09.2026.xlsx",
                  "2026_W37_Credit_Report_05.09-11.09.2026.xlsx",
                  "2026_W38_Credit_Report_12.09-18.09.2026.xlsx"]
    assert sorted(sg.SoGhi(cfg.so_ghi).trang) == ["2026-09-05", "2026-09-12", "2026-09-19"]
    assert all(f.name.startswith("[DONE]") for f in cfg.vao.glob("*.xlsx"))


@can_du_lieu_that
def test_thieu_file_di_kem_thi_bo_qua_va_ghi_ro(tmp_path, monkeypatch):
    import shutil
    from wr import caidat
    tep = tmp_path / "settings.json"
    monkeypatch.setattr(caidat, "TEP_CAU_HINH", tep)
    monkeypatch.setattr(runner.caidat, "TEP_CAU_HINH", tep)
    cfg = caidat.dat_workspace(tmp_path / "dich", tao_thu_muc_con=False)
    shutil.copy(TH[37], cfg.vao / TH[37].name)   # chỉ có file Tổng hợp
    assert runner.quet_thu_muc(cfg) == 0
    assert not list(cfg.ra.glob("*.xlsx"))
    assert not any(f.name.startswith("[LOI]") for f in cfg.vao.iterdir())


# --- chạy được ở máy nào, đặt ở đâu cũng được ---------------------------

def test_khong_co_duong_dan_cung_trong_ma_nguon():
    """Không file mã nào được ghi cứng đường dẫn của một máy cụ thể."""
    import re
    xau = re.compile(r"[A-Za-z]:\\\\|/home/|/Users/|/sessions/|OneDrive - ")
    for f in list((GOC / "wr").glob("*.py")) + list((GOC / "tools").glob("*.py")) + [GOC / "runner.py"]:
        assert not xau.search(f.read_text(encoding="utf-8")), f"{f.name} có đường dẫn cứng"


def test_thu_muc_dich_co_dau_cach_dau_tieng_viet_va_o_khac(tmp_path, monkeypatch):
    from wr import caidat
    tep = tmp_path / "settings.json"
    monkeypatch.setattr(caidat, "TEP_CAU_HINH", tep)
    dich = tmp_path / "Ổ D" / "Công nợ Trustana (2026)" / "thư mục có dấu cách"
    cfg = caidat.dat_workspace(dich, tao_thu_muc_con=False)
    assert cfg.workspace == dich and all(d.is_dir() for d in cfg.cac_thu_muc())


def test_moi_may_giu_duong_dan_rieng_trong_cung_file_cau_hinh(tmp_path, monkeypatch):
    """Thư mục mã nguồn có thể nằm trên OneDrive và đồng bộ sang máy khác —
    cấu hình của máy này không được đè lên máy kia."""
    import json
    from wr import caidat
    tep = tmp_path / "settings.json"
    monkeypatch.setattr(caidat, "TEP_CAU_HINH", tep)

    monkeypatch.setattr(caidat, "ten_may", lambda: "MAY-CONG-TY")
    caidat.dat_workspace(tmp_path / "cong-ty", tao_thu_muc_con=False)
    monkeypatch.setattr(caidat, "ten_may", lambda: "MAY-NHA")
    caidat.dat_workspace(tmp_path / "nha", tao_thu_muc_con=False)

    raw = json.loads(tep.read_text(encoding="utf-8"))
    assert set(raw["workspace_theo_may"]) == {"MAY-CONG-TY", "MAY-NHA"}
    monkeypatch.setattr(caidat, "ten_may", lambda: "MAY-CONG-TY")
    assert caidat.doc().workspace == tmp_path / "cong-ty"
    monkeypatch.setattr(caidat, "ten_may", lambda: "MAY-NHA")
    assert caidat.doc().workspace == tmp_path / "nha"
    monkeypatch.setattr(caidat, "ten_may", lambda: "MAY-LA")
    with pytest.raises(caidat.ChuaDatThuMuc):
        caidat.doc()


@can_du_lieu_that
def test_file_ket_qua_dang_mo_trong_excel_thi_hoan_lai(tmp_path, monkeypatch):
    """Windows khoá file đang mở. Tool phải hoãn kỳ đó, không được đánh dấu
    [LOI] rồi bỏ luôn."""
    import shutil
    from wr import caidat
    tep = tmp_path / "settings.json"
    monkeypatch.setattr(caidat, "TEP_CAU_HINH", tep)
    monkeypatch.setattr(runner.caidat, "TEP_CAU_HINH", tep)
    cfg = caidat.dat_workspace(tmp_path / "dich", tao_thu_muc_con=False)
    shutil.copy(TH[37], cfg.vao / TH[37].name)
    shutil.copy(TN[37], cfg.vao / TN[37].name)

    def _khoa(*args, **kwargs):
        raise PermissionError("[Errno 13] file đang được mở bởi chương trình khác")
    monkeypatch.setattr(runner.baocao, "ghi_bao_cao", _khoa)

    assert runner.quet_thu_muc(cfg) == 0
    con_lai = sorted(f.name for f in cfg.vao.glob("*.xlsx"))
    assert con_lai == sorted([TH[37].name, TN[37].name]), "file đầu vào phải giữ nguyên để chạy lại"


@can_du_lieu_that
def test_loai_khach_khoi_bao_cao_va_van_ghi_lai(tmp_path):
    """Nợ ảo do lỗi xuất hoá đơn: kế toán khai loại một lần, nhưng vẫn phải
    nhìn thấy khoản đó ở sheet Cần xem lại để không ai quên."""
    import openpyxl
    so_ghi = tmp_path / "so-ghi.json"
    runner.main(["--bo-qua", "OPTIMALOG LLC", "--vi-sao",
                 "no ao do loi xuat hoa don tien coc, se thu trong lo sau",
                 "--so-ghi", str(so_ghi)])
    runner.main(["--tong-hop", str(TH[38]), "--tuoi-no", str(TN[38]),
                 "--ra", str(tmp_path), "--so-ghi", str(so_ghi), "--mau-chuan", str(MAU)])
    f = next(tmp_path.glob("2026_W38*.xlsx"))
    assert misa.chuan_hoa_ten("OPTIMALOG LLC") not in baocao.doc_bao_cao(f)
    ws = openpyxl.load_workbook(f)["Review"]
    chu = "\n".join(str(c.value or "") for r in ws.iter_rows() for c in r)
    assert "Khách bị loại khỏi bảng" in chu and "OPTIMALOG" in chu
    runner.main(["--nhan-lai", "OPTIMALOG LLC", "--so-ghi", str(so_ghi)])
    runner.main(["--tong-hop", str(TH[38]), "--tuoi-no", str(TN[38]),
                 "--ra", str(tmp_path), "--so-ghi", str(so_ghi), "--mau-chuan", str(MAU)])
    assert misa.chuan_hoa_ten("OPTIMALOG LLC") in baocao.doc_bao_cao(next(tmp_path.glob("2026_W38*.xlsx")))


@can_du_lieu_that
def test_bao_cao_ghi_lai_nguon_du_lieu(tmp_path):
    """Số trên Misa đổi theo thời điểm xuất, nên báo cáo phải ghi lại nó dựa
    trên lần xuất nào."""
    import openpyxl
    dich, _ = _chay(tmp_path, 37)
    ws = openpyxl.load_workbook(dich)["Review"]
    chu = "\n".join(str(c.value or "") for r in ws.iter_rows(max_row=10) for c in r)
    assert "Nguồn dữ liệu" in chu and "file sửa lần cuối" in chu
    assert "cùng một lần xuất Misa" in chu


def test_lenh_chi_doc_khong_de_ra_thu_muc_la(tmp_path, monkeypatch):
    """--xem-cau-hinh chỉ đọc, không được tạo thư mục nào."""
    from wr import caidat
    tep = tmp_path / "settings.json"
    monkeypatch.setattr(caidat, "TEP_CAU_HINH", tep)
    tep.write_text('{"workspace_theo_may": {"%s": "%s"}}'
                   % (caidat.ten_may(), (tmp_path / "chua-ton-tai").as_posix()), encoding="utf-8")
    cfg = caidat.doc()
    assert [d.name for d in cfg.cac_thu_muc()]
    assert not (tmp_path / "chua-ton-tai").exists(), "đọc cấu hình không được tạo thư mục"


def test_duong_dan_cua_may_khac_thi_bao_ro_khong_tao_thu_muc(tmp_path, monkeypatch):
    """Cấu hình của máy Windows mà đem sang máy khác thì phải báo, không được
    tạo một thư mục tên 'C:\\Users\\...' ngay trong thư mục mã nguồn."""
    import os
    from wr import caidat
    if os.name == "nt":
        pytest.skip("chỉ kiểm trên máy không phải Windows")
    tep = tmp_path / "settings.json"
    monkeypatch.setattr(caidat, "TEP_CAU_HINH", tep)
    tep.write_text('{"workspace": "C:\\\\Users\\\\ai do\\\\Cong-No-Workspace"}', encoding="utf-8")
    with pytest.raises(caidat.ChuaDatThuMuc):
        caidat.doc()
    assert not list(GOC.glob("C:*"))

# -*- coding: utf-8 -*-
"""Phép thử WR-04: full credit report từ 4 file (Tổng hợp 131, Chi tiết 131 từ 2022,
Tổng hợp 331, SMS). Dữ liệu GIẢ, sinh ngay trong phép thử."""
import datetime as dt
import os
import sys
import time
from pathlib import Path

import openpyxl
import pytest

GOC = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(GOC))
sys.path.insert(0, str(GOC / "tests"))

from wr import baocao, chitiet, misa, tindung  # noqa: E402
from wr import soghi as sg  # noqa: E402
from wr import thuonghieu as th  # noqa: E402
import runner  # noqa: E402

D = dt.date
TU, DEN = D(2026, 9, 26), D(2026, 10, 2)


# --- sinh file giả --------------------------------------------------------

def _tieu_de_tong_hop(ws, ten, tu, den):
    ws["A1"] = ten
    ws["A2"] = f"Tài khoản: 131, Loại tiền: <<Tổng hợp>>, Từ ngày {tu:%d/%m/%Y} đến ngày {den:%d/%m/%Y}"
    for i, t in enumerate(["Mã", "Tên", "Mã số thuế ", "TK công nợ", "Số dư đầu kỳ", "", "Phát sinh",
                           "", "Số dư cuối kỳ", ""], start=1):
        ws.cell(row=4, column=i, value=t)
    for i, t in enumerate(["Nợ", "Có"] * 3, start=5):
        ws.cell(row=5, column=i, value=t)


def tao_tong_hop(dich, khach, tu=TU, den=DEN, ten="TỔNG HỢP CÔNG NỢ PHẢI THU KHÁCH HÀNG"):
    """khach: [(mã, tên, dk_nợ, dk_có, ps_nợ, ps_có, ck_nợ, ck_có)]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    _tieu_de_tong_hop(ws, ten, tu, den)
    r = 6
    for ma, ten_kh, *so in khach:
        for i, v in enumerate([ma, ten_kh, "", "131", *so], start=1):
            ws.cell(row=r, column=i, value=v)
        r += 1
    ws.cell(row=r, column=1, value="Tổng cộng")
    Path(dich).parent.mkdir(parents=True, exist_ok=True)
    wb.save(dich)
    return Path(dich)


def tao_phai_tra(dich, vendor, tu=TU, den=DEN):
    return tao_tong_hop(dich, vendor, tu, den, ten="TỔNG HỢP CÔNG NỢ PHẢI TRẢ NHÀ CUNG CẤP")


def tao_chi_tiet(dich, khach, den=DEN, tu=D(2022, 1, 1)):
    """khach: {tên: [(ngày hạch toán, ngày hoá đơn, số hoá đơn, diễn giải, TK đối ứng, nợ, có)]}"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "CHI TIẾT CÔNG NỢ PHẢI THU KHÁCH HÀNG"
    ws["A2"] = f"Tài khoản: 131, Loại tiền: <<Tổng hợp>>, Từ ngày {tu:%d/%m/%Y} đến ngày {den:%d/%m/%Y}"
    for i, t in enumerate(["Ngày hạch toán", "Ngày chứng từ", "Số chứng từ", "Ngày hóa đơn",
                           "Số hóa đơn", "Diễn giải", "TK công nợ", "TK đối ứng", "Phát sinh", "",
                           "Số dư", ""], start=1):
        ws.cell(row=4, column=i, value=t)
    r = 6
    for ten, dong in khach.items():
        ws.cell(row=r, column=1, value=f"Tên khách hàng: {ten}")
        r += 1
        for j, (ht, hd, so, dg, tk, no, co) in enumerate(dong):
            for i, v in enumerate([dt.datetime.combine(ht, dt.time()), dt.datetime.combine(ht, dt.time()),
                                   f"CT{r}", dt.datetime.combine(hd, dt.time()) if hd else "", so or "",
                                   dg, "131", tk, no, co, 0, 0], start=1):
                ws.cell(row=r, column=i, value=v)
            r += 1
        ws.cell(row=r, column=6, value="Cộng")
        r += 1
    ws.cell(row=r, column=1, value="Tổng cộng")
    Path(dich).parent.mkdir(parents=True, exist_ok=True)
    wb.save(dich)
    return Path(dich)


# Bộ dữ liệu tuần 40 giả
A, B, C, K = "CÔNG TY A", "CÔNG TY B", "CÔNG TY C", "CÔNG TY KHÓ ĐÒI"
CT40 = {
    # A: 2 hoá đơn; phiếu thu ghi rõ trả HĐ 00000002 (mới) chứ không phải HĐ cũ
    A: [(D(2026, 8, 1), D(2026, 8, 1), "00000001", "Phí", "5113", 100, 0),
        (D(2026, 9, 1), D(2026, 9, 1), "00000002", "Phí", "5113", 50, 0),
        (D(2026, 9, 28), None, "", "Thu tiền theo HD 00000002", "1121", 0, 50)],
    # B: hạn đúng ngày cuối kỳ (Current) và hạn hôm trước (1 - 30); thu tiền không ghi HĐ -> cũ nhất
    B: [(D(2026, 8, 1), D(2026, 8, 1), "00000010", "Phí", "5113", 30, 0),
        (D(2026, 9, 1), D(2026, 9, 1), "00000011", "Phí", "5113", 40, 0),
        (D(2026, 9, 2), D(2026, 9, 2), "00000012", "Phí", "5113", 60, 0),
        (D(2026, 9, 30), None, "", "Thu tiền khách hàng", "1122", 0, 25),
        (D(2026, 9, 30), None, "", "Phí ngân hàng", "6427", 0, 5)],
    # C: Chi tiết thiếu 20 so với Tổng hợp (hai file xuất lệch giờ)
    C: [(D(2026, 9, 20), D(2026, 9, 20), "00000020", "Phí", "5113", 80, 0)],
    K: [(D(2026, 9, 25), D(2026, 9, 25), "00000030", "Phí", "5113", 70, 0)],
}
TH40 = [  # mã, tên, dk_nợ, dk_có, ps_nợ, ps_có, ck_nợ, ck_có
    ("KA", A, 150, 0, 0, 50, 100, 0),
    ("KB", B, 130, 0, 0, 30, 100, 0),
    ("KC", C, 0, 0, 100, 0, 100, 0),
    ("KK", K, 70, 0, 0, 0, 70, 0),
    ("KP", "KHÁCH TRẢ TRƯỚC", 0, 0, 0, 0, 0, 9),
]
TRA40 = [  # vendor
    ("V1", "VENDOR CÒN NỢ", 0, 500, 100, 50, 0, 450),
    ("V2", "VENDOR TRẢ XONG", 0, 300, 300, 0, 0, 0),
    ("V3", "VENDOR MỚI", 0, 0, 0, 200, 0, 200),
    ("V4", "VENDOR MÌNH TRẢ TRƯỚC", 0, 0, 10, 0, 10, 0),
]


@pytest.fixture
def so(tmp_path):
    s = sg.SoGhi(tmp_path / "so.json")
    s.dieu_chinh = {misa.chuan_hoa_ten(A): {"credit_term": "30 days"},
                    misa.chuan_hoa_ten(B): {"credit_term": "30 days"},
                    misa.chuan_hoa_ten(C): {"credit_term": "30 days"}}
    return s


def _kq(tmp_path, so, ct=CT40, tong_hop=TH40, tra=TRA40):
    t = misa.doc_tong_hop(tao_tong_hop(tmp_path / "th.xlsx", tong_hop))
    c = chitiet.doc_chi_tiet(tao_chi_tiet(tmp_path / "ct.xlsx", ct))
    p = misa.doc_phai_tra(tao_phai_tra(tmp_path / "tra.xlsx", tra)) if tra is not None else None
    return tindung.tinh_full(t, c, so, pt=p), t, c


# --- đọc file -------------------------------------------------------------

def test_nhan_dang_4_loai_file_theo_noi_dung(tmp_path):
    assert misa.nhan_dang(tao_tong_hop(tmp_path / "x1.xlsx", TH40)) == "tong-hop"
    assert misa.nhan_dang(tao_phai_tra(tmp_path / "x2.xlsx", TRA40)) == "phai-tra"
    assert misa.nhan_dang(tao_chi_tiet(tmp_path / "x3.xlsx", CT40)) == "chi-tiet"
    assert chitiet.ky_cua(tmp_path / "x3.xlsx") == (D(2022, 1, 1), DEN)


def test_so_hoa_don_trong_dien_giai():
    f = chitiet.so_hd_nhac_toi
    assert f("Thu tiền của LCC UPP theo HD 00001574, 00001419") == [1574, 1419]
    assert f("thanh toán job: EXSANA24100118 - HD 1142") == [1142]
    assert f("Thu tiền khách hàng theo hóa đơn 00001600, 00001602 ngày 23/09/2026") == [1600, 1602]
    assert f("Thu tiền job EXSANA24100118") == []


def test_term_14_ngay_cua_thang_tiep_theo():
    assert sg.han_tu_term(D(2026, 9, 18), "14 ngày của tháng tiếp theo") == D(2026, 10, 14)
    assert sg.han_tu_term(D(2026, 12, 5), "14 ngày của tháng tiếp theo") == D(2027, 1, 14)
    assert sg.han_tu_term(D(2026, 9, 18), "15 days") == D(2026, 10, 3)
    assert sg.han_tu_term(D(2026, 9, 18), "") == D(2026, 10, 18)


def test_phieu_thu_ghi_so_hoa_don_tru_dung_hoa_don(tmp_path):
    c = chitiet.doc_chi_tiet(tao_chi_tiet(tmp_path / "ct.xlsx", CT40))
    con, du = c.khach[misa.chuan_hoa_ten(A)].hoa_don_con_no(DEN)
    assert con == [(D(2026, 8, 1), "00000001", 100)] and du == 0, "HĐ cũ vẫn nợ vì phiếu thu ghi HĐ 2"


def test_so_hoa_don_trung_nam_truoc_chon_hoa_don_gan_nhat(tmp_path):
    ct = {A: [(D(2024, 3, 1), D(2024, 3, 1), "00000005", "Phí", "5113", 10, 0),
              (D(2024, 3, 9), None, "", "Thu HD 5", "1121", 0, 10),
              (D(2026, 9, 1), D(2026, 9, 1), "00000005", "Phí", "5113", 40, 0),
              (D(2026, 9, 2), D(2026, 9, 2), "00000006", "Phí", "5113", 60, 0),
              (D(2026, 9, 20), None, "", "Thu theo HD 00000006", "1121", 0, 60)]}
    c = chitiet.doc_chi_tiet(tao_chi_tiet(tmp_path / "ct.xlsx", ct))
    con, _ = c.khach[misa.chuan_hoa_ten(A)].hoa_don_con_no(DEN)
    assert con == [(D(2026, 9, 1), "00000005", 40)]


# --- tuổi nợ --------------------------------------------------------------

def test_tuoi_no_theo_hoa_don_va_luat_qua_han_tu_hom_sau(tmp_path, so):
    kq, t, _ = _kq(tmp_path, so)
    d = {x.ma: x for x in kq.dong}
    assert d["KA"].nhom == [0, 0, 100, 0, 0, 0], "HĐ 01/08 + 30 = 31/08, đến 02/10 quá 32 ngày"
    # B: thu 30 trừ HĐ 01/08 (30) — còn HĐ 01/09 (hạn 01/10, quá 1 ngày) và HĐ 02/09 (hạn 02/10 = Current)
    assert d["KB"].nhom == [60, 40, 0, 0, 0, 0]
    assert sum(x.total for x in kq.dong) == pytest.approx(t.tong_du_no)
    assert "KP" not in d, "khách trả trước không vào Receivables"


def test_chi_tiet_thieu_so_voi_tong_hop_thi_tien_theo_tong_hop(tmp_path, so):
    kq, _, _ = _kq(tmp_path, so)
    c = next(x for x in kq.dong if x.ma == "KC")
    assert c.total == 100 and c.nhom[0] == 100 and c.nguon == "uoc-tinh"
    assert (C, 100, 80) in kq.lech_chi_tiet


def test_kho_doi_luon_120(tmp_path, so):
    so.trang["2026-09-26"] = {misa.chuan_hoa_ten(K): sg.HoSoKhach(
        ma="KK", ten=K, ly_do="Công nợ khó đòi, đang kiện", tong=70, nhom=[70, 0, 0, 0, 0, 0])}
    kq, _, _ = _kq(tmp_path, so)
    k = next(x for x in kq.dong if x.ma == "KK")
    assert k.nhom == [0, 0, 0, 0, 0, 70]
    r = next(x for x in kq.rui_ro if x.ma == "KK")
    assert r.nhom_rui_ro == tindung.N4


# --- rủi ro theo luật của chị KTT -------------------------------------------

def _r(**kw):
    base = dict(ma="X", ten="X", credit_term="30 days", dau_ky=0, tang=0, giam=0, cuoi_ky=100,
                nhom=[0, 100, 0, 0, 0, 0], salesman="", job="", ly_do="đã hứa trả",
                tt_total=100, tt_qua_han=100, tt_61=0, co_tuan_truoc=True, nhap_bo_sung=0)
    base.update(kw)
    return tindung.danh_gia_rui_ro(tindung.DongRuiRo(**base))


def test_luat_rui_ro():
    assert _r(nhom=[100, 0, 0, 0, 0, 0]).nhom_rui_ro == tindung.N0
    r = _r()                                           # quá hạn đứng yên = S2 -> N2
    assert r.s[:3] == [0, 1, 0] and r.nhom_rui_ro == tindung.N2
    r = _r(cuoi_ky=200, nhom=[0, 100, 0, 100, 0, 0], tt_qua_han=100)   # S1 + S3 -> N3
    assert r.s[:3] == [1, 0, 1] and r.nhom_rui_ro == tindung.N3
    assert _r(nhom=[0, 0, 0, 0, 0, 100]).nhom_rui_ro == tindung.N3, "có 120+ là N3"
    assert _r(ly_do="chờ cấn trừ với vendor").nhom_rui_ro == tindung.NX
    assert _r(ly_do="khách mất liên lạc").nhom_rui_ro == tindung.N4
    assert _r(tt_qua_han=150).nhom_rui_ro == tindung.N1, "quá hạn giảm, không dấu hiệu"
    assert _r(nhom=[0, 0, 0, 0, 0, 0], cuoi_ky=0).nhom_rui_ro == tindung.N0
    assert _r(nhom=[0, 150_000_000, 0, 0, 0, 0], cuoi_ky=150_000_000, ly_do="").ceo


# --- phải trả, dòng tiền ---------------------------------------------------

def test_payable_trang_thai_va_khong_loc(tmp_path, so):
    kq, _, _ = _kq(tmp_path, so)
    tt = {v.ma: v.trang_thai for v in kq.phai_tra}
    assert tt == {"V1": "Partly paid", "V2": "Paid off this week", "V3": "New this week"}
    assert kq.vendor_tra_truoc == [("V4", "VENDOR MÌNH TRẢ TRƯỚC", 10)]
    assert kq.chi_so["tong_phai_tra"] == 650 and kq.chi_so["da_tra_vendor"] == 410


def test_tien_thu_tach_theo_tk_doi_ung(tmp_path, so):
    kq, _, _ = _kq(tmp_path, so)
    cs = kq.chi_so
    assert cs["thu_ngan_hang"] == 75 and cs["thu_ty_gia_phi"] == 5 and cs["da_thu_tien"] == 75
    assert cs["dong_tien_rong"] == 75 - 410


# --- chạy cả vòng trong workspace -----------------------------------------

@pytest.fixture
def ws(tmp_path, monkeypatch):
    from wr import caidat
    tep = tmp_path / "settings.json"
    monkeypatch.setattr(caidat, "TEP_CAU_HINH", tep)
    monkeypatch.setattr(runner.caidat, "TEP_CAU_HINH", tep)
    monkeypatch.setattr(runner, "NHAT_KY", tmp_path / "log.txt")
    # v0.5: tool chờ thêm file Bán hàng + Sổ chi tiết; các phép thử 4 file ở đây chạy kiểu cũ ngay
    monkeypatch.setattr(runner, "CHO_SMS_GIAY", 0)
    return caidat.dat_workspace(tmp_path / "dich", tao_thu_muc_con=False)


def _tha_4_file(cfg, tuan=40):
    if tuan == 40:
        tao_tong_hop(cfg.vao / "a.xlsx", TH40)
        tao_chi_tiet(cfg.vao / "b.xlsx", CT40)
        tao_phai_tra(cfg.vao / "c.xlsx", TRA40)
    else:
        tu, den = D(2026, 10, 3), D(2026, 10, 9)
        ct = {k: list(v) for k, v in CT40.items()}
        ct[A].append((D(2026, 10, 5), D(2026, 10, 5), "00000003", "Phí", "5113", 30, 0))
        # hoá đơn nhập lùi ngày cho kỳ trước: B thêm 15 mang ngày 20/09
        ct[B].append((D(2026, 9, 20), D(2026, 9, 20), "00000013", "Phí", "5113", 15, 0))
        tao_tong_hop(cfg.vao / "a2.xlsx", [("KA", A, 100, 0, 30, 0, 130, 0), ("KB", B, 115, 0, 0, 0, 115, 0),
                                          ("KC", C, 100, 0, 0, 0, 100, 0), ("KK", K, 70, 0, 0, 0, 70, 0)],
                     tu, den)
        tao_chi_tiet(cfg.vao / "b2.xlsx", ct, den)
        tao_phai_tra(cfg.vao / "c2.xlsx", [("V1", "VENDOR CÒN NỢ", 0, 470, 0, 0, 0, 470)], tu, den)
    from test_wr3 import tao_sms
    tao_sms(cfg.vao / f"sms{tuan}.xlsx", [("Outbound", [
        ("KA", A, "Hằng", "EXSANA26080001", "01/08/2026", None, [("FEE", 100, 1)])])])


def test_tha_4_file_ra_full_report(ws):
    _tha_4_file(ws)
    assert runner.quet_thu_muc(ws) == 1
    f = next(ws.ra.glob("2026_W40_Credit_Report_*.xlsx"))
    wb = openpyxl.load_workbook(f)
    assert wb.sheetnames == [th.SHEET_SUMMARY, th.SHEET_BANG, th.SHEET_RISK, th.SHEET_PAYABLE,
                             th.SHEET_CASH, th.SHEET_INVOICES, th.SHEET_METHOD, th.SHEET_REVIEW]
    rec = wb[th.SHEET_BANG]
    assert [c.value for c in rec[5]][2] == "Customer's Name" and rec.cell(5, 13).value == "Job Number"
    hang_a = next(r for r in rec.iter_rows(min_row=6, values_only=True) if r[1] == "KA")
    assert hang_a[12] == "EXSANA26080001", "Job Number không có chữ 'job:'"
    assert all(x.name.startswith("[DONE]") for x in ws.vao.glob("*.xlsx"))
    chu = " ".join(str(c) for r in wb[th.SHEET_REVIEW].iter_rows(values_only=True) for c in r if c)
    assert "KHỚP" in chu and "Chi tiết 131 và Tổng hợp 131 lệch" in chu
    # đọc lại được thành sổ ghi (tuần sau dùng)
    may = baocao.doc_bao_cao(f)
    assert may[misa.chuan_hoa_ten(A)].job == "EXSANA26080001"


def test_thieu_file_phai_tra_thi_cho_roi_moi_chay(ws, monkeypatch):
    _tha_4_file(ws)
    next(ws.vao.glob("c.xlsx")).unlink()
    monkeypatch.setattr(runner, "CHO_SMS_GIAY", 7200)
    assert runner.quet_thu_muc(ws) == 0, "chưa đủ 2 giờ thì chờ file phải trả"
    monkeypatch.setattr(runner, "CHO_SMS_GIAY", 0)
    assert runner.quet_thu_muc(ws) == 1
    f = next(ws.ra.glob("*.xlsx"))
    assert "chưa có file Tổng hợp công nợ phải trả" in str(openpyxl.load_workbook(f)[th.SHEET_PAYABLE]["A6"].value)


def test_tuan_sau_co_lich_su_va_tach_nhap_bo_sung(ws, monkeypatch):
    _tha_4_file(ws, 40)
    assert runner.quet_thu_muc(ws) == 1
    _tha_4_file(ws, 41)
    assert runner.quet_thu_muc(ws) == 1
    f = next(ws.ra.glob("2026_W41_*.xlsx"))
    wb = openpyxl.load_workbook(f, data_only=True)
    sm = wb[th.SHEET_SUMMARY]
    tuan = [r[0] for r in sm.iter_rows(values_only=True) if r[0] in ("W40", "W41")]
    assert tuan == ["W40", "W41"]
    dong = {r[5]: r for r in sm.iter_rows(min_row=20, values_only=True) if r and r[5]}
    assert dong["tong_phai_thu"][1] == 370 and dong["tong_phai_thu"][2] == 415
    risk = {r[1]: r for r in wb[th.SHEET_RISK].iter_rows(min_row=6, values_only=True) if r[1]}
    assert risk["KB"][9] == 15, "B: dư đầu kỳ Misa 115 − báo cáo tuần trước 100 = nhập bổ sung 15"
    pay = {r[1]: r for r in wb[th.SHEET_PAYABLE].iter_rows(min_row=6, values_only=True) if r[1]}
    assert pay["V1"][3] == pay["V1"][4] == 470, "Payable theo mẫu FIN: Current = Total = số dư cuối"


def test_ban_nguoi_lam_thang_ban_tool_cung_ngay_chot(ws):
    _tha_4_file(ws)
    assert runner.quet_thu_muc(ws) == 1
    tool = next(ws.ra.glob("*.xlsx"))
    # bản người làm: chép bảng Receivables sang file không có tab Review, sửa lý do
    wb = openpyxl.load_workbook(tool)
    for n in list(wb.sheetnames):
        if n != th.SHEET_BANG:
            del wb[n]
    sh = wb[th.SHEET_BANG]
    for r in range(6, sh.max_row + 1):
        if sh.cell(r, 2).value == "KA":
            sh.cell(r, 14, value="KTT: khách hứa trả 10/10")
    nguoi = ws.ra / "W40 final cua KTT.xlsx"
    wb.save(nguoi)
    so = sg.SoGhi(ws.so_ghi)
    runner.nap_ban_tay(ws.ra, so)
    time.sleep(1.1)
    os.utime(tool)                       # bản tool được mở/lưu lại SAU bản người
    runner.nap_ban_tay(ws.ra, so)
    k = so.trang["2026-10-03"][misa.chuan_hoa_ten(A)]
    assert k.ly_do == "KTT: khách hứa trả 10/10", "bản tool sửa sau vẫn không đè bản người làm"


def test_sample_nhan_file_ghep_theo_noi_dung(tmp_path):
    from test_wr3 import tao_ghep
    from wr import ghep
    g = tao_ghep(tmp_path / "ghep cua FIN.xlsx", [("KA", "KA", "Khớp tên/mã", "Đúng", None)])
    time.sleep(1.1)
    wb = openpyxl.Workbook()
    wb.active.title = "Customer Master List"
    wb.save(tmp_path / "Customer_Master_List.xlsx")     # mới hơn nhưng không phải file ghép
    assert ghep.tim_file_ghep(tmp_path) == g


def test_ten_sales_day_du():
    assert th.ten_sales("Hiếu") == "Trần Văn Hiếu"
    assert th.ten_sales("Hiếu 1") == "Phạm Trần Hiếu"
    assert th.ten_sales("Admin,Admin") == "Admin"
    assert th.ten_sales("Phạm Thị Thương Hoài") == "Phạm Thị Thương Hoài"


def test_ly_do_tu_dong_khi_sales_chua_ghi_va_khong_mang_sang_tuan_sau(ws):
    _tha_4_file(ws)
    assert runner.quet_thu_muc(ws) == 1
    f = next(ws.ra.glob("*.xlsx"))
    rec = openpyxl.load_workbook(f)[th.SHEET_BANG]
    hang_a = next(r for r in rec.iter_rows(min_row=6, values_only=True) if r[1] == "KA")
    assert hang_a[13].startswith(th.NHAN_TU_DONG) and "CHƯA CÓ GIẢI TRÌNH" in hang_a[13]
    may = baocao.doc_bao_cao(f)
    assert may[misa.chuan_hoa_ten(A)].ly_do == "", "câu tự động không được coi là lời Sales"



def test_dong_con_theo_job_va_hoa_don_cong_dung_total(ws):
    _tha_4_file(ws)
    assert runner.quet_thu_muc(ws) == 1
    wb = openpyxl.load_workbook(next(ws.ra.glob("*.xlsx")))
    rec = wb[th.SHEET_BANG]
    r_a = next(r for r in range(6, rec.max_row + 1) if rec.cell(r, 2).value == "KA")
    con = rec.cell(r_a + 1, 3).value
    assert con.strip().startswith("└ EXSANA26080001") and rec.row_dimensions[r_a + 1].outlineLevel == 1
    assert rec.row_dimensions[r_a + 1].hidden, "dòng con mặc định đang gập"
    assert rec.cell(r_a + 1, 12).value == "Lê Thị Hằng", "tên ngắn trên SMS đổi sang tên đầy đủ"
    assert rec.cell(r_a + 1, 6).value is None and rec.cell(r_a + 1, 8).value == 100, \
        "dòng con nằm đúng cột tuổi của khách (KA: 100 ở 31 - 60)"
    tong = rec.cell(rec.max_row, 5).value
    assert rec.cell(rec.max_row, 2).value == "TOTAL" and tong == 370
    inv = wb[th.SHEET_INVOICES]
    tien = [r[9] for r in inv.iter_rows(min_row=6, values_only=True) if isinstance(r[0], int)]
    assert sum(tien) == tong
    may = baocao.doc_bao_cao(next(ws.ra.glob("*.xlsx")))
    assert len(may) == 4, "dòng con không bị đọc thành khách"



def test_dashboard_nam_o_thu_muc_goc(ws):
    runner.quet_thu_muc(ws)
    f = ws.ra / "Credit_Dashboard.html"
    assert f.exists() and f.read_bytes() == runner.DASHBOARD.read_bytes()
    chu = f.read_text(encoding="utf-8")
    assert "/*__" not in chu, "build.py đã cắm đủ CSS, thư viện đọc Excel và mã"
    for sheet in ("Summary", "Receivables", "AR Risk", "Payable", "Cash Flow", "Invoices"):
        assert f'"{sheet}"' in chu, f"dashboard đọc sheet {sheet}"



def test_ngay_chot_bao_cao_tool_la_hom_sau_ngay_cuoi_ky(ws):
    """Summary ghi cả 'At 03 Oct' lẫn 'Reporting period … 02/10' — phải đọc 03/10, không thì
    tuần sau thấy một 'tuần trước' giả cách một ngày."""
    _tha_4_file(ws)
    runner.quet_thu_muc(ws)
    assert baocao.ngay_chot_trong_file(next(ws.ra.glob("*.xlsx"))) == D(2026, 10, 3)


def test_fin_khai_job_cho_no_cu_va_dong_con_dung_cot_tuoi(tmp_path, so):
    from wr import ghep
    wb = openpyxl.Workbook()
    w = wb.active
    w.title = ghep.SHEET_JOB_NO_CU
    w.append(ghep.COT_JOB_NO_CU)
    w.append(["KB", B, "00000011", "01/09/2026", 40, "EXSANA26090099", "Hiếu", ""])
    wb.save(tmp_path / "Tham_chieu_job_no_cu.xlsx")
    ref = ghep.doc_job_no_cu(tmp_path)
    assert ref == {"KB": [(11, "EXSANA26090099", "Hiếu")]}
    t = misa.doc_tong_hop(tao_tong_hop(tmp_path / "th.xlsx", TH40))
    c = chitiet.doc_chi_tiet(tao_chi_tiet(tmp_path / "ct.xlsx", CT40))
    kq = tindung.tinh_full(t, c, so, job_no_cu=ref)
    b = next(x for x in kq.dong if x.ma == "KB")          # B: [60 Current, 40 1-30]
    g = b.dong_con[0]
    assert g["job"] == "EXSANA26090099" and g["sales"] == "Trần Văn Hiếu"
    assert g["nhom"] == [0, 40, 0, 0, 0, 0], "HĐ 00000011 hạn 01/10 → quá 1 ngày"
    assert b.can_doi == 60 and b.can_doi_nhom == [60, 0, 0, 0, 0, 0]

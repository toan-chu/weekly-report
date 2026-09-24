# -*- coding: utf-8 -*-
"""Sinh 2 file Misa GIẢ để thử tool mà không cần dữ liệu thật.

    python tests/tao_du_lieu_gia.py "<thư mục 01_Input>"

Số liệu bịa hoàn toàn, tên khách không có thật. Bộ dữ liệu cố tình gài đủ các
tình huống: nợ trong hạn, quá hạn từng nhóm, hoá đơn không ghi hạn, hai file
Misa lệch nhau, khách trả trước, khách đã trả hết nợ.
"""
import sys
from pathlib import Path

from openpyxl import Workbook

TU_NGAY, DEN_NGAY, NGAY_CHOT = "19/09/2026", "25/09/2026", "26/09/2026"

# ma, ten, mst, du_no, [khong_han, truoc_han, 1-30, 31-60, 61-90, 91-120, >120], du_co
KHACH = [
    ("KHT01", "CÔNG TY TNHH GIAO NHẬN BÌNH MINH", "0101234567", 120_000_000,
     [0, 120_000_000, 0, 0, 0, 0, 0], 0),
    ("KHT02", "CÔNG TY TNHH THỰC PHẨM AN KHANG", "0102345678", 85_500_000,
     [0, 30_000_000, 55_500_000, 0, 0, 0, 0], 0),
    ("KHT03", "SKYPORT LOGISTICS PTE. LTD", "", 240_000_000,
     [0, 0, 0, 240_000_000, 0, 0, 0], 0),
    ("KHT04", "CÔNG TY CỔ PHẦN DỆT MAY HỒNG PHÁT", "0103456789", 63_250_000,
     [0, 0, 0, 0, 40_000_000, 23_250_000, 0], 0),
    ("KHT05", "EASTWIND CARGO CO., LTD", "", 310_000_000,
     [0, 0, 0, 0, 0, 0, 310_000_000], 0),
    ("KHT06", "CÔNG TY TNHH CƠ KHÍ TRƯỜNG AN", "0104567890", 47_800_000,
     [47_800_000, 0, 0, 0, 0, 0, 0], 0),           # hoá đơn không ghi hạn
    ("KHT07", "NORDSTAR FREIGHT OY", "", 150_000_000,
     [0, 0, 420_000_000, 0, 0, 0, 0], 0),          # hai file lệch nhau
    ("KHT08", "CÔNG TY TNHH NỘI THẤT MINH QUÂN", "0105678901", 0,
     None, 12_000_000),                            # trả trước, không vào báo cáo
    ("KHT09", "CÔNG TY TNHH XNK HẢI ĐĂNG", "0106789012", 0,
     None, 0),                                     # đã trả hết, không vào báo cáo
    ("KHT10", "PACIFIC ROUTE LIMITED", "", 95_000_000,
     [0, 60_000_000, 35_000_000, 0, 0, 0, 0], 0),
    ("KHT11", "CÔNG TY TNHH DƯỢC PHẨM TÂM AN", "0107890123", 18_600_000,
     [0, 18_600_000, 0, 0, 0, 0, 0], 0),
    ("KHT12", "GULF LINE SHIPPING LLC", "", 505_000_000,
     [0, 0, 0, 0, 205_000_000, 0, 300_000_000], 0),
]


def file_tong_hop(dich: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "TỔNG HỢP CÔNG NỢ PHẢI THU KHÁC"
    ws["A1"] = "TỔNG HỢP CÔNG NỢ PHẢI THU KHÁCH HÀNG"
    ws["A2"] = f"Tài khoản: 131, Loại tiền: <<Tổng hợp>>, Từ ngày {TU_NGAY} đến ngày {DEN_NGAY}"
    for cot, ten in zip("ABCDEGI", ["Mã khách hàng", "Tên khách hàng", "Mã số thuế ",
                                    "TK công nợ", "Số dư đầu kỳ", "Phát sinh", "Số dư cuối kỳ"]):
        ws[f"{cot}4"] = ten
    for cot, ten in zip("EFGHIJ", ["Nợ", "Có", "Nợ", "Có", "Nợ", "Có"]):
        ws[f"{cot}5"] = ten
    hang = 6
    tong = [0] * 6
    for ma, ten, mst, du_no, _, du_co in KHACH:
        dau_no = round(du_no * 0.6)
        ps_no = du_no - dau_no
        gia_tri = [ma, ten, mst, 131, dau_no, 0, ps_no, 0, du_no, du_co]
        for i, v in enumerate(gia_tri, start=1):
            ws.cell(row=hang, column=i, value=v)
        for i, v in enumerate([dau_no, 0, ps_no, 0, du_no, du_co]):
            tong[i] += v
        hang += 1
    ws.cell(row=hang, column=1, value="Tổng cộng")
    for i, v in enumerate(tong, start=5):
        ws.cell(row=hang, column=i, value=v)
    dich.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dich)
    return dich


def file_tuoi_no(dich: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "Phan_tich_cong_no_phai_thu_the"
    ws["A1"] = "CÔNG TY TNHH TRUSTANA VIỆT NAM"
    ws["A2"] = "Lô S1, Tầng 6, toà nhà Viwaseen Tower, 48 Tố Hữu, Hà Nội"
    ws["A3"] = "PHÂN TÍCH CÔNG NỢ PHẢI THU THEO TUỔI NỢ"
    ws["A4"] = f"Tài khoản: 131; Đến ngày {NGAY_CHOT}"
    for cot, ten in zip(["A", "B", "D", "F", "H", "I", "P"],
                        ["Mã khách hàng", "Tên khách hàng", "Địa chỉ", "Tổng nợ",
                         "Không có hạn nợ", "Nợ trước hạn", "Nợ quá hạn"]):
        ws[f"{cot}6"] = ten
    for cot, ten in zip(["I", "J", "K", "L", "M", "N", "P", "Q", "R", "S", "U", "V"],
                        ["0-30 ngày", "31-60 ngày", "61-90 ngày", "91-120 ngày",
                         "Trên 120 ngày", "Tổng", "1-30 ngày", "31-60 ngày", "61-90 ngày",
                         "91-120 ngày", "Trên 120 ngày", "Tổng"]):
        ws[f"{cot}8"] = ten
    ws["A9"] = "Mã nhóm khách hàng: <<Khác>>"
    ws["C9"] = "Tên nhóm khách hàng: <<Khác>>"
    hang = 10
    cong = [0] * 8
    for ma, ten, _, _, tuoi, _ in KHACH:
        if not tuoi:
            continue
        khong_han, truoc_han, b1, b2, b3, b4, b5 = tuoi
        tong_no = sum(tuoi)
        qua_han = b1 + b2 + b3 + b4 + b5
        ws.cell(row=hang, column=1, value=ma)
        ws.cell(row=hang, column=2, value=ten)
        ws.cell(row=hang, column=4, value="Địa chỉ thử nghiệm")
        ws.cell(row=hang, column=6, value=tong_no)
        if khong_han:
            ws.cell(row=hang, column=8, value=khong_han)
        if truoc_han:
            ws.cell(row=hang, column=9, value=truoc_han)    # 0-30 ngày tới hạn
            ws.cell(row=hang, column=14, value=truoc_han)   # tổng nợ trước hạn
        for cot, v in zip((16, 17, 18, 19, 21), (b1, b2, b3, b4, b5)):
            if v:
                ws.cell(row=hang, column=cot, value=v)
        if qua_han:
            ws.cell(row=hang, column=22, value=qua_han)
        for i, v in enumerate((tong_no, khong_han, truoc_han, b1, b2, b3, b4, b5)):
            cong[i] += v if i < 3 else v
        hang += 1
    ws.cell(row=hang, column=1, value="Cộng nhóm: <<Khác>>")
    ws.cell(row=hang, column=6, value=cong[0])
    ws.cell(row=hang + 1, column=1, value="Tổng cộng")
    ws.cell(row=hang + 1, column=6, value=cong[0])
    dich.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dich)
    return dich


def main() -> int:
    thu_muc = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(".")
    a = file_tong_hop(thu_muc / "THU_NGHIEM_Tong_hop_cong_no_phai_thu_khach_hang.xlsx")
    b = file_tuoi_no(thu_muc / "THU_NGHIEM_Phan_tich_cong_no_phai_thu_theo_tuoi_no.xlsx")
    tong = sum(k[3] for k in KHACH)
    print(f"Đã sinh 2 file dữ liệu giả, kỳ {TU_NGAY}–{DEN_NGAY}, chốt {NGAY_CHOT}")
    print(f"  {a.name}\n  {b.name}")
    print(f"  {sum(1 for k in KHACH if k[3] > 0)} khách còn nợ, tổng {tong:,.0f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

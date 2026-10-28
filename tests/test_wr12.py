# -*- coding: utf-8 -*-
"""WR-12 (27/10): khách mới tự vào file credit term (FIN không phải làm gì), link hợp đồng bấm mở từ báo cáo,
khoá lời Sales đã ghi (tuần đã chốt không đọc lại), file Excel gọi là "file credit term" (sheet 'Credit term')."""
import hashlib
import sys
import time
from pathlib import Path

import openpyxl

GOC = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(GOC))
sys.path.insert(0, str(GOC / "tests"))

from wr import khachhang as kh  # noqa: E402
from wr import thuonghieu as th  # noqa: E402
import runner  # noqa: E402
from test_wr4 import A, _tha_4_file, ws  # noqa: E402,F401
from test_wr11 import _cot, _sua_ly_do  # noqa: E402

TEN = kh.TEN_FILE_MAC_DINH


def _file_cu(ws_):
    """File credit term theo mẫu cũ: sheet tên 'Client Master', có một cột người dùng tự chèn."""
    ws_.ghep.mkdir(parents=True, exist_ok=True)
    f = ws_.ghep / TEN
    kh.ghi_master(f, [dict(ma="KA", ten=A, term="30 days", co_hd="✓", han_muc=50,
                           link_hd="https://trustana.sharepoint.com/hd/KA.pdf")], ["Lê Thị Hằng"])
    wb = openpyxl.load_workbook(f)
    sh = wb[kh.SHEET_MASTER]
    sh.title = "Client Master"
    c = sh.max_column + 1
    sh.cell(kh.DONG_TIEU_DE, c, "Ghi chú FIN")
    sh.cell(kh.DONG_DAU, c, "khách VIP")
    wb.save(f)
    return f


def _bang(f):
    w = openpyxl.load_workbook(f)
    sh = kh.sheet_credit_term(w)
    tieu = [sh.cell(kh.DONG_TIEU_DE, i).value for i in range(1, sh.max_column + 1)]
    rows = {}
    for r in range(kh.DONG_DAU, sh.max_row + 1):
        ma = sh.cell(r, tieu.index("Mã kh") + 1).value
        if ma:
            rows[ma] = {t: sh.cell(r, i + 1).value for i, t in enumerate(tieu) if t}
    return sh.title, rows


def test_khach_moi_tu_vao_file_credit_term_chi_them_dong_va_sao_luu(ws):
    f = _file_cu(ws)
    _tha_4_file(ws, 40)
    assert runner.quet_thu_muc(ws) == 1
    ten_sheet, rows = _bang(f)
    assert ten_sheet == kh.SHEET_MASTER == "Credit term", "file cũ được ghi lại theo mẫu mới"
    assert {"KB", "KC", "KK"} <= set(rows), "khách mới đang nợ đã vào file"
    assert rows["KB"]["Ghi chú"].startswith("Tool thêm") and rows["KB"]["Thời hạn"] is None, "ô term để FIN điền"
    ka = rows["KA"]
    assert (ka["Thời hạn"], ka["Hạn mức nợ (VND)"], ka["Đã có hợp đồng"], ka["Ghi chú FIN"]) == \
        ("30 days", 50, "✓", "khách VIP"), "dòng đã có giữ nguyên, kể cả cột FIN tự chèn"
    luu = sorted((ws.so_ghi.parent / "luu-tru").glob("SAMPLE_credit_term_khach_truoc-khi-tool-ghi_*.xlsx"))
    # file mẫu cũ + có khách mới: lần 1 nâng mẫu (đầu lần chạy), lần 2 thêm khách — mỗi lần ghi đều có bản sao lưu
    assert len(luu) == 2 and _bang(luu[0])[0] == "Client Master", "bản sao lưu đầu tiên là file gốc chưa đụng"
    ct = openpyxl.load_workbook(next(ws.ra.glob("2026_W40_*.xlsx")))[th.SHEET_CONTRACT]
    assert any(str(o.value or "").startswith("New customers — added") for r in ct.iter_rows() for o in r)


def test_file_credit_term_dang_mo_thi_khong_ghi_bao_cao_van_ra(ws):
    f = _file_cu(ws)
    (f.parent / f"~${f.name}").write_text("khoá")          # Excel đang mở file trên máy này
    truoc = hashlib.md5(f.read_bytes()).hexdigest()
    _tha_4_file(ws, 40)
    assert runner.quet_thu_muc(ws) == 1, "báo cáo vẫn ra"
    assert hashlib.md5(f.read_bytes()).hexdigest() == truoc, "file đang mở thì không đụng vào"
    wb = openpyxl.load_workbook(next(ws.ra.glob("2026_W40_*.xlsx")))
    ct = wb[th.SHEET_CONTRACT]
    o = next(o for r in ct.iter_rows() for o in r if str(o.value or "").startswith("New customers — NOT yet"))
    assert "đang mở" in str(ct.cell(o.row + 1, 3).value)
    assert any("đang mở" in str(v) for r in wb[th.SHEET_REVIEW].iter_rows(values_only=True) for v in r if v)


def test_link_hop_dong_bam_mo_tu_ar_risk_va_contracts(ws):
    f = _file_cu(ws)
    wb = openpyxl.load_workbook(f)
    sh = wb["Client Master"]
    tieu = [sh.cell(kh.DONG_TIEU_DE, i).value for i in range(1, sh.max_column + 1)]
    sh.append([None] * len(tieu))
    r = kh.DONG_DAU + 1
    for t, v in (("Mã kh", "KB"), ("Khách hàng", "CÔNG TY B"), ("Thời hạn", "30 days"),
                 ("Link hợp đồng", "C:\\\\hop dong\\\\KB.pdf")):
        sh.cell(r, tieu.index(t) + 1, v)
    wb.save(f)
    _tha_4_file(ws, 40)
    assert runner.quet_thu_muc(ws) == 1
    risk = openpyxl.load_workbook(next(ws.ra.glob("2026_W40_*.xlsx")))[th.SHEET_RISK]
    c, h = _cot(risk, "Contract (PDF)")
    o = {risk.cell(i, 2).value: risk.cell(i, c) for i in range(h + 1, risk.max_row + 1)}
    assert o["KA"].value == "Mở hợp đồng" and o["KA"].hyperlink.target.endswith("/hd/KA.pdf")
    assert o["KB"].value is None, "link không phải http(s) thì không tạo ô bấm"


def test_tuan_da_chot_khong_doc_lai_chinh_sua_va_o_tuan_truoc_chan_go(ws):
    _tha_4_file(ws, 40)
    assert runner.quet_thu_muc(ws) == 1
    w40 = next(ws.ra.glob("2026_W40_*.xlsx"))
    time.sleep(1.1)
    _sua_ly_do(w40, "KA", "Khách hẹn trả 15/10")
    _tha_4_file(ws, 41)
    assert runner.quet_thu_muc(ws) == 1
    time.sleep(1.1)
    _sua_ly_do(w40, "KA", "Sales sửa lại lời cũ")                       # sửa báo cáo W40 SAU khi đã có W41
    for f in ws.vao.glob("[[]DONE[]] *"):
        if "2." in f.name or "41" in f.name:
            f.rename(f.with_name(f.name.replace("[DONE] ", "")))
    assert runner.quet_thu_muc(ws) == 1
    lai = next(ws.ra.glob("2026_W41_*chay lai*.xlsx"))
    sh = openpyxl.load_workbook(lai)[th.SHEET_BANG]
    c_lw, _ = _cot(sh, th.COT_LY_DO_TRUOC[0])
    ka = next(r for r in range(6, sh.max_row + 1) if sh.cell(r, 2).value == "KA")
    assert sh.cell(ka, c_lw).value == "Khách hẹn trả 15/10", "lời đã ghi khi tuần còn mở, không phải bản sửa sau"
    dv = [d for d in sh.data_validations.dataValidation if d.type == "custom"]
    cl = openpyxl.utils.get_column_letter(c_lw)
    assert dv and dv[0].formula1 == "FALSE" and str(dv[0].sqref).startswith(f"{cl}6"), "ô tuần trước chặn gõ"


def _file_mau_cu_du_khach(ws_):
    """File credit term mẫu CŨ (sheet 'Client Master', chỉ vài cột) đã có đủ 4 khách → tuần này không có khách mới."""
    ws_.ghep.mkdir(parents=True, exist_ok=True)
    f = ws_.ghep / TEN
    wb = openpyxl.Workbook()
    sh = wb.active
    sh.title = "Client Master"
    tieu = ["STT", "Mã kh", "Khách hàng", "Sales phụ trách", "Đã có hợp đồng", "Thời hạn", "Ghi chú"]
    for i, t in enumerate(tieu, start=1):
        sh.cell(kh.DONG_TIEU_DE, i, t)
    for j, (ma, ten) in enumerate((("KA", A), ("KB", "CÔNG TY B"), ("KC", "CÔNG TY C"), ("KK", "CÔNG TY KHÓ ĐÒI"))):
        for i, v in enumerate([None, ma, ten, "Lê Thị Hằng", "✓" if ma == "KA" else None, "30 days", "cũ"], start=1):
            sh.cell(kh.DONG_DAU + j, i, v)
    wb.save(f)
    return f


def test_file_mau_cu_tu_nang_mau_khi_khong_co_khach_moi(ws):
    f = _file_mau_cu_du_khach(ws)
    assert kh.thieu_mau(f), "file mẫu cũ thiếu cột"
    _tha_4_file(ws, 40)
    assert runner.quet_thu_muc(ws) == 1
    ten_sheet, rows = _bang(f)
    assert ten_sheet == "Credit term" and kh.thieu_mau(f) == [], "đã nâng lên mẫu mới dù không có khách mới"
    assert set(rows) == {"KA", "KB", "KC", "KK"}, "không thêm, không bớt khách"
    assert rows["KA"]["Thời hạn"] == "30 days" and rows["KA"]["Đã có hợp đồng"] == "✓" and rows["KB"]["Ghi chú"] == "cũ"
    assert "Hạn mức nợ (VND)" in rows["KA"] and "Link hợp đồng" in rows["KA"]
    luu = list((ws.so_ghi.parent / "luu-tru").glob("*_truoc-khi-tool-ghi_*.xlsx"))
    assert len(luu) == 1
    wb = openpyxl.load_workbook(next(ws.ra.glob("2026_W40_*.xlsx")))
    assert any("Đã nâng" in str(v) for r in wb[th.SHEET_REVIEW].iter_rows(values_only=True) for v in r if v)
    assert not any("chưa có cột 'Đã có hợp đồng'" in str(v)
                   for r in wb[th.SHEET_REVIEW].iter_rows(values_only=True) for v in r if v), \
        "nâng xong mới đọc file → không còn cảnh báo thiếu cột"
    # tuần sau: file đã đúng mẫu, không có khách mới → không ghi lại, không thêm bản sao lưu
    truoc = f.stat().st_mtime
    _tha_4_file(ws, 41)
    assert runner.quet_thu_muc(ws) == 1
    assert f.stat().st_mtime == truoc and len(list((ws.so_ghi.parent / "luu-tru").glob("*_truoc-khi-tool-ghi_*"))) == 1


def test_file_mau_cu_dang_mo_thi_khong_nang(ws):
    f = _file_mau_cu_du_khach(ws)
    (f.parent / f"~${f.name}").write_text("khoá")
    truoc = hashlib.md5(f.read_bytes()).hexdigest()
    _tha_4_file(ws, 40)
    assert runner.quet_thu_muc(ws) == 1, "báo cáo vẫn ra"
    assert hashlib.md5(f.read_bytes()).hexdigest() == truoc
    wb = openpyxl.load_workbook(next(ws.ra.glob("2026_W40_*.xlsx")))
    assert any("Chưa nâng" in str(v) and "đang mở" in str(v)
               for r in wb[th.SHEET_REVIEW].iter_rows(values_only=True) for v in r if v)

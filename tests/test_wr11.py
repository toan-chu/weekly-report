# -*- coding: utf-8 -*-
"""WR-11 (27/10): màu theo tuổi nợ + nâng hạng (CEO), Client Master thêm Hạn mức nợ / Không nâng hạng màu /
Link hợp đồng, lý do tuần này + tuần trước (PA1), khách mới chưa có trong Client Master (FIN), sửa lỗi WR-09."""
import datetime as dt
import shutil
import subprocess
import sys
import time
from pathlib import Path

import openpyxl
import pytest

GOC = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(GOC))
sys.path.insert(0, str(GOC / "tests"))

from wr import chitiet as ctm, khachhang as kh, tindung  # noqa: E402
from wr import thuonghieu as th  # noqa: E402
import runner  # noqa: E402
from test_wr4 import A, _tha_4_file, ws  # noqa: E402,F401

D = dt.date


# --- Client Master ----------------------------------------------------------

def test_o_tich_chi_nhan_dau_tich_va_thieu_cot_la_khong_biet(tmp_path):
    # lỗi WR-09 tự soi: gõ "Chưa" vào ô Đã có hợp đồng từng bị coi là ĐÃ có hợp đồng
    kh.ghi_master(tmp_path / "SAMPLE_credit_term_khach.xlsx", [
        dict(ma="A", ten="A", term="30 days", co_hd="✓", han_muc=500_000_000, khong_nang="✓",
             link_hd="https://sharepoint/hd-a.pdf"),
        dict(ma="B", ten="B", term="30 days", co_hd="Chưa", han_muc="200tr"),
        dict(ma="C", ten="C", term="30 days", co_hd="x")], [])
    m = kh.doc_master(tmp_path, D(2026, 10, 24))
    assert (m["A"]["co_hd"], m["B"]["co_hd"], m["C"]["co_hd"]) == (True, False, True)
    assert m["A"]["han_muc"] == 500_000_000 and m["B"]["han_muc"] == 200_000_000 and m["C"]["han_muc"] is None
    assert m["A"]["khong_nang"] and not m["B"]["khong_nang"] and m["A"]["link_hd"].endswith("hd-a.pdf")
    # file mẫu cũ (KTT) không có cột "Đã có hợp đồng" → None, không phải False
    cu = tmp_path / "cu"
    cu.mkdir()
    wb = openpyxl.Workbook()
    wb.active.append(["Mã kh", "Khách hàng", "Thời hạn"])
    wb.active.append(["A", "A", "30 days"])
    wb.save(cu / "credit-term.xlsx")
    assert kh.doc_master(cu)["A"]["co_hd"] is None


def test_nang_cap_giu_cot_tu_chen_va_khong_tu_tich(tmp_path):
    sys.path.insert(0, str(GOC / "tools"))
    import nang_cap_credit_term as nc
    goc = tmp_path / "cu.xlsx"
    kh.ghi_master(goc, [dict(ma="A", ten="A", term="30 days", so_hd="HD-01")], ["Trần Văn Hiếu"])
    wb = openpyxl.load_workbook(goc)
    ws_ = wb[kh.SHEET_MASTER]
    c = ws_.max_column + 1
    ws_.cell(kh.DONG_TIEU_DE, c, "Ghi chú của Sales")
    ws_.cell(kh.DONG_DAU, c, "khách quen")
    wb.save(goc)
    moi = nc.nang_cap(goc)
    assert goc.exists() and moi != goc, "không ghi đè file cũ"
    w = openpyxl.load_workbook(moi)[kh.SHEET_MASTER]
    tieu = [w.cell(kh.DONG_TIEU_DE, i).value for i in range(1, w.max_column + 1)]
    i = tieu.index("Ghi chú của Sales") + 1
    assert w.cell(kh.DONG_DAU, i).value == "khách quen", "cột người dùng tự chèn vẫn còn"
    j = tieu.index("Đã có hợp đồng") + 1
    assert w.cell(kh.DONG_DAU, j).value is None, "có Số hợp đồng cũng không tự tích ✓"
    assert "Hạn mức nợ (VND)" in tieu and "Không nâng hạng màu" in tieu and "Link hợp đồng" in tieu


# --- màu ------------------------------------------------------------------

def _r(nhom, **kw):
    T = sum(nhom)
    r = tindung.DongRuiRo(ma="X", ten="X", credit_term="30 days", dau_ky=T, tang=0, giam=0, cuoi_ky=T,
                          nhom=list(nhom), salesman="S", job="", ly_do=kw.pop("ly_do", "có lý do"),
                          tt_total=kw.pop("tt_total", T), tt_qua_han=0, tt_61=0, co_tuan_truoc=True,
                          nhap_bo_sung=0, **kw)
    return tindung.danh_gia_rui_ro(r, D(2026, 10, 24))


def test_mau_goc_theo_nhom_tuoi_cu_nhat_bo_so_le(monkeypatch):
    monkeypatch.setattr(tindung, "NGUONG_DU_LE", 1_000)     # conftest hạ ngưỡng cho số liệu giả; ở đây cần số thật
    assert _r([10, 0, 0, 0, 0, 0]).muc_goc == 0
    assert _r([0, 5_000, 0, 0, 0, 0]).muc_goc == 1
    assert _r([0, 5_000, 5_000, 0, 0, 0]).muc_goc == 2
    assert _r([0, 0, 0, 5_000, 0, 0]).muc_goc == 3
    assert _r([0, 0, 0, 0, 5_000, 0]).muc_goc == 4 and _r([0, 0, 0, 0, 0, 5_000]).muc_goc == 4
    assert _r([0, 5_000, 0, 0, 0, 500]).muc_goc == 1, "dưới 1.000đ ở nhóm cũ là số lẻ, không kéo màu"


def test_nang_hang_mot_muc_khong_cong_don():
    X = tindung.NGUONG_X
    r = _r([0, X, 0, 0, 0, 0])                                         # T1
    assert r.tieu_chi[0] == 1 and r.muc == 2
    r = _r([0, X, 0, 0, 0, 0], han_muc=1, so_hd_tre=5, co_hd=False)     # T1..T4 đều đúng → vẫn +1
    assert sum(r.tieu_chi) == 4 and r.muc == 2
    assert _r([0, 5_000, 0, 0, 0, 0], han_muc=1_000).muc == 2           # T2 vượt hạn mức
    assert _r([0, 5_000, 0, 0, 0, 0], so_hd_tre=2).muc == 2             # T3
    assert _r([0, 5_000, 0, 0, 0, 0], so_hd_tre=1).muc == 1
    assert _r([0, 5_000, 0, 0, 0, 0], co_hd=None).muc == 1, "Client Master chưa có cột → không nâng"
    Y = tindung.NGUONG_Y
    assert _r([Y, 0, 0, 0, 0, 0], tt_total=0).muc == 1, "khách mới nợ ≥ Y: từ Xanh lên Vàng"
    assert _r([Y - 1, 0, 0, 0, 0, 0], tt_total=0).muc == 0
    assert _r([0, 0, 0, 0, 0, X], han_muc=1).muc == 4, "tối đa Đỏ đậm"


def test_khong_nang_hang_va_nang_thang_len_do():
    X = tindung.NGUONG_X
    assert _r([0, X, 0, 0, 0, 0], khong_nang=True).muc == 1, "khách tốt: bỏ +1"
    assert _r([0, 5_000, 0, 0, 0, 0], ly_do="Khách hứa không trả").muc == 3
    assert _r([0, 5_000, 0, 0, 0, 0], ly_do="mất liên lạc với kế toán", khong_nang=True).muc == 3, \
        "không nâng hạng KHÔNG chặn nâng thẳng"
    assert _r([0, 0, 0, 0, 5_000, 0], ly_do="mất liên lạc").muc == 4, "đã Đỏ đậm thì giữ Đỏ đậm"
    assert _r([0, 5_000, 0, 0, 0, 0], ly_do="khách không trả kịp vì bank Ấn Độ").muc == 1, \
        "'không trả' trơn không đủ để lên Đỏ (bản thử W43 bắt nhầm)"
    r = _r([0, 5_000, 0, 0, 0, 0], ly_do="", ly_do_truoc="mất liên lạc từ tuần trước")
    assert r.muc == 3 and r.s[3] == 1, "tuần này chưa ghi: xét theo lý do tuần trước, vẫn tính chưa giải trình"


def test_dem_hoa_don_tra_tre_trong_6_thang():
    def bt(ngay, so_hd, no=0.0, co=0.0, dg=""):
        return ctm.ButToan(ngay_ht=ngay, ngay_hd=ngay if no else None, so_hd=so_hd, so_ct="X", dien_giai=dg,
                           tk_doi_ung="1122", no=no, co=co)
    k = ctm.KhachChiTiet("K", "K", [
        bt(D(2026, 1, 5), "00000001", no=100),                     # trả 01/03: trễ nhưng ngoài 6 tháng
        bt(D(2026, 3, 1), "", co=100, dg="theo HD 00000001"),
        bt(D(2026, 6, 1), "00000002", no=100),                     # hạn 01/07, trả 10/08 → trễ 40 ngày
        bt(D(2026, 8, 10), "", co=100, dg="theo HD 00000002"),
        bt(D(2026, 7, 1), "00000003", no=100),                     # hạn 31/07, trả 10/08 → trễ 10 ngày
        bt(D(2026, 8, 10), "", co=100, dg="theo HD 00000003"),
        bt(D(2026, 8, 1), "00000004", no=100),                     # hạn 31/08, trả 30/09 (FIFO) → trễ 30
        bt(D(2026, 9, 30), "", co=100),
        bt(D(2026, 9, 1), "00000005", no=100),                     # còn nợ → không đếm ở đây
    ])
    ra = k.hoa_don_tra_tre(D(2026, 10, 23), lambda n: n + dt.timedelta(days=30))
    assert [(s, tre) for s, _n, _d, tre in ra] == [("00000002", 40), ("00000004", 30)]


# --- chạy thật từ input (PA1, khách mới, công thức) ----------------------

def _sua_ly_do(f: Path, ma: str, chu: str) -> None:
    wb = openpyxl.load_workbook(f)
    sh = wb[th.SHEET_BANG]
    for r in range(6, sh.max_row + 1):
        if sh.cell(r, 2).value == ma:
            sh.cell(r, th.COT_LY_DO_NAY, value=chu)
    wb.save(f)


def _cot(sh, ten):
    for r in sh.iter_rows(min_row=1, max_row=6):
        for o in r:
            if o.value == ten:
                return o.column, o.row
    raise KeyError(ten)


def test_pa1_ly_do_tuan_nay_trong_tuan_truoc_chi_doc_va_chay_lai_giu_ly_do(ws):
    _tha_4_file(ws, 40)
    assert runner.quet_thu_muc(ws) == 1
    w40 = next(ws.ra.glob("2026_W40_*.xlsx"))
    time.sleep(1.1)
    _sua_ly_do(w40, "KA", "Khách hẹn trả 15/10")                       # Sales ghi vào báo cáo W40
    _tha_4_file(ws, 41)
    assert runner.quet_thu_muc(ws) == 1
    w41 = next(ws.ra.glob("2026_W41_*.xlsx"))
    sh = openpyxl.load_workbook(w41)[th.SHEET_BANG]
    c_lw, _ = _cot(sh, th.COT_LY_DO_TRUOC[0])
    ka = next(r for r in range(6, sh.max_row + 1) if sh.cell(r, 2).value == "KA")
    assert sh.cell(ka, th.COT_LY_DO_NAY).value is None, "tuần mới: ô lý do tuần này để trống cho Sales ghi"
    assert sh.cell(ka, c_lw).value == "Khách hẹn trả 15/10", "lý do tuần trước hiện ở cột chỉ đọc"
    risk = openpyxl.load_workbook(w41)[th.SHEET_RISK]
    c_pd, h = _cot(risk, "Promise date")
    rk = next(r for r in range(h + 1, risk.max_row + 1) if risk.cell(r, 2).value == "KA")
    assert risk.cell(rk, c_pd).value is not None, "ngày hẹn vẫn đọc từ lý do tuần trước (S7)"
    # Sales ghi lý do tuần này rồi chạy lại cùng tuần → giữ lý do tuần này
    time.sleep(1.1)
    _sua_ly_do(w41, "KA", "Đã gửi công văn, khách xin trả 30/10")
    for f in ws.vao.glob("[[]DONE[]] *"):
        if "2." in f.name or "41" in f.name:
            f.rename(f.with_name(f.name.replace("[DONE] ", "")))
    assert runner.quet_thu_muc(ws) == 1
    lai = next(ws.ra.glob("2026_W41_*chay lai*.xlsx"))
    sh = openpyxl.load_workbook(lai)[th.SHEET_BANG]
    ka = next(r for r in range(6, sh.max_row + 1) if sh.cell(r, 2).value == "KA")
    assert sh.cell(ka, th.COT_LY_DO_NAY).value == "Đã gửi công văn, khách xin trả 30/10"
    assert sh.cell(ka, c_lw).value == "Khách hẹn trả 15/10"


def test_khach_moi_chua_co_trong_client_master_va_sales_chua_giai_trinh(ws):
    ws.ghep.mkdir(parents=True, exist_ok=True)
    kh.ghi_master(ws.ghep / "SAMPLE_credit_term_khach.xlsx",
                  [dict(ma="KA", ten=A, term="30 days", co_hd="✓", han_muc=50)], [])
    _tha_4_file(ws, 40)
    assert runner.quet_thu_muc(ws) == 1
    wb = openpyxl.load_workbook(next(ws.ra.glob("2026_W40_*.xlsx")))
    ct = wb[th.SHEET_CONTRACT]
    o = next(o for r in ct.iter_rows() for o in r if str(o.value or "").startswith("New customers"))
    dau = o.row + 2
    assert [ct.cell(dau, c).value for c in (3, 4, 5)] == ["Mã kh", "Khách hàng", "Sales phụ trách"], \
        "3 cột đầu đúng thứ tự Client Master để FIN dán"
    moi = {ct.cell(r, 3).value for r in range(dau + 1, ct.max_row + 1) if ct.cell(r, 3).value}
    assert {"KB", "KC", "KK"} <= moi and "KA" not in moi
    sm = {r[5]: r for r in wb[th.SHEET_SUMMARY].iter_rows(min_row=20, values_only=True) if r and r[5]}
    assert sm["khach_chua_master"][2] == len(moi)
    assert sm["khach_chua_hd"][2] == 3, "KB, KC, KK không có trong Client Master → chưa có hợp đồng"
    assert any(str(r[0] or "").startswith("Reasons missing this week") for r in
               wb[th.SHEET_SUMMARY].iter_rows(values_only=True))


@pytest.mark.skipif(not shutil.which("soffice"), reason="Cần LibreOffice để tính lại công thức")
def test_cong_thuc_mau_trong_excel_khop_python(ws, tmp_path):
    ws.ghep.mkdir(parents=True, exist_ok=True)
    kh.ghi_master(ws.ghep / "SAMPLE_credit_term_khach.xlsx",
                  [dict(ma="KA", ten=A, term="30 days", co_hd="✓", han_muc=50),
                   dict(ma="KB", ten="CÔNG TY B", term="30 days", co_hd="✓", khong_nang="✓")], [])
    _tha_4_file(ws, 40)
    assert runner.quet_thu_muc(ws) == 1
    f = next(ws.ra.glob("2026_W40_*.xlsx"))
    out = tmp_path / "lo"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(f)],
                   check=True, capture_output=True, timeout=120)
    wb = openpyxl.load_workbook(out / f.name, data_only=True)
    risk = wb[th.SHEET_RISK]
    c_fin, h = _cot(risk, "Final level")
    muc = [risk.cell(r, c_fin).value for r in range(h + 1, risk.max_row + 1)
           if isinstance(risk.cell(r, 1).value, int)]
    assert muc and all(isinstance(m, (int, float)) for m in muc), "công thức tính ra số, không lỗi"
    sm = {r[5]: r for r in wb[th.SHEET_SUMMARY].iter_rows(min_row=20, values_only=True) if r and r[5]}
    for i in range(5):
        assert muc.count(i) == (sm[f"mau_{i}"][2] or 0), f"mức {i}: công thức Excel = tool"


def test_to_mau_co_dieu_kien_ghi_bgcolor_cho_excel(ws):
    """Lỗi chạy lại W43 trên máy FIN (29/10): CF chỉ ghi fgColor → Excel không tô màu. Excel đọc nền CF ở bgColor."""
    import re
    import zipfile
    _tha_4_file(ws, 40)
    assert runner.quet_thu_muc(ws) == 1
    f = next(ws.ra.glob("2026_W40_*.xlsx"))
    st = zipfile.ZipFile(f).read("xl/styles.xml").decode("utf-8")
    dxf = re.findall(r"<dxf>.*?</dxf>", re.search(r"<dxfs.*?</dxfs>", st, re.S).group(0), re.S)
    nen = [d for d in dxf if "<fill>" in d]
    assert nen and all("bgColor" in d for d in nen), "mọi nền định dạng có điều kiện phải có bgColor"
